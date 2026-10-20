"""Build a disposable aggregate from current split sheets, preserving formula caches."""
from copy import copy
from pathlib import Path
import io
import json
import shutil
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
_MASTER_CACHE = {}

def build_current_master(game: Path, destination: Path, sections):
    signature = tuple((str(p), p.stat().st_mtime_ns, p.stat().st_size) for p in sorted((game/'config_excel').glob('*.xlsx')))
    if signature in _MASTER_CACHE:
        destination.write_bytes(_MASTER_CACHE[signature]); return
    book = openpyxl.Workbook(); book.remove(book.active)
    caches = {}
    for name in sections:
        source = game / 'config_excel' / (name + '.xlsx')
        if not source.exists():
            continue
        incoming = openpyxl.load_workbook(source, data_only=False)
        sheet = incoming[name]; target = book.create_sheet(name)
        for row in sheet:
            for cell in row:
                out = target.cell(cell.row, cell.column, cell.value)
                if cell.has_style:
                    for attr in ('font','fill','border','alignment','protection'):
                        setattr(out, attr, copy(getattr(cell, attr)))
                out.number_format = cell.number_format
        for merged in sheet.merged_cells.ranges: target.merge_cells(str(merged))
        incoming.close()
        # Each split source has one sheet. Keep cached values rather than
        # recalculating formulas with a different evaluator in the test.
        with zipfile.ZipFile(source) as archive:
            part = next(n for n in archive.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'))
            tree = ET.fromstring(archive.read(part))
            caches[name] = {c.attrib['r']: c.find(NS+'v') for c in tree.iter(NS+'c') if c.find(NS+'f') is not None}
    book.save(destination); names = book.sheetnames; book.close()
    raw = destination.read_bytes(); output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                index = int(Path(item.filename).stem.removeprefix('sheet'))-1
                tree = ET.fromstring(data)
                for cell in tree.iter(NS+'c'):
                    if cell.find(NS+'f') is None: continue
                    old = cell.find(NS+'v')
                    if old is not None: cell.remove(old)
                    cached = caches[names[index]].get(cell.attrib['r'])
                    if cached is not None: cell.append(copy(cached))
                data = ET.tostring(tree)
            target.writestr(item, data)
    _MASTER_CACHE[signature] = output.getvalue()
    destination.write_bytes(output.getvalue())

def seed_current_data(game: Path, directory: Path):
    directory.mkdir(parents=True, exist_ok=True)
    for source in (game/'data').glob('*.json'):
        shutil.copy2(source, directory/source.name)
