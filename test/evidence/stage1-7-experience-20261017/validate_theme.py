import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[3]
data = json.loads((root / 'space-battleship/data/game_data.json').read_text(encoding='utf-8'))
levels, groups, enemies = data['levels'], data['groups'], data['enemies']
expected = [5, 6, 7, 8, 9, 9, 9]
assert [len(levels[i]['groups']) for i in range(7)] == expected

for stage in range(1, 8):
    print(f'STAGE {stage}: {expected[stage-1]} points, length={levels[stage-1]["length"]}')
    for point, group in enumerate(levels[stage-1]['groups'], 1):
        units = [enemies[str(i)] for i in groups[str(group['id'])]['slots'] if i is not None]
        hull = {unit['armourType'] for unit in units}
        shield = {unit['shieldType'] for unit in units if unit['shield'] > 0}
        attacks = {item['name'].replace('-', '_') for unit in units for item in unit['equipment']}
        recharge = sum(unit['shield'] > 0 and unit['shieldRecovery'] > 0 for unit in units)
        assert len(hull) == 1, (stage, point, 'mixed hull', hull)
        assert len(attacks) == 1, (stage, point, 'mixed attack', attacks)
        if stage in (3, 4):
            assert not shield or shield == hull, (stage, point, 'mixed resistance')
        if stage in (6, 7):
            assert hull == {2} and attacks == {'laser_mon'}
            assert recharge >= (0 if stage == 6 and point == 1 else 1)
        print(f'  P{point}: n={len(units)}, hull={next(iter(hull))}, shield={next(iter(shield)) if shield else 0}, attack={next(iter(attacks))}, recharge={recharge}')

stage5 = [groups[str(group['id'])]['slots'] for group in levels[4]['groups']]
assert [sum(i is not None for i in slots) for slots in stage5] == [3,3,10,10,1,1,3,10,1]
assert [enemies[str(next(i for i in slots if i is not None))]['armourType'] for slots in stage5] == [2,2,1,1,1,1,2,1,1]
for point, expected_visual in enumerate(('sentry','sentry','swarm','swarm','reactive','reactive','sentry','swarm','reactive'), 1):
    units = [enemies[str(i)] for i in stage5[point-1] if i is not None]
    assert all(unit.get('archetype_id') == expected_visual for unit in units)
    if expected_visual == 'reactive':
        assert all(unit.get('reactiveReduction') == 1 for unit in units)
assert enemies['400004']['archetype_id'] == 'interceptor'
print('STATIC PASS: point counts, single-resistance and attack family, stage5 composition, stage6/7 recharge')

config = root / 'space-battleship/config_excel'
level_book = load_workbook(config / 'level.xlsx', read_only=True, data_only=True)
assert level_book['level']['D64'].value is not None, 'level formula cache lost'
group_book = load_workbook(config / 'monGroup.xlsx', read_only=True)
assert group_book['monGroup']['C2295'].value.count('271017') == 2
used_groups = {group['id'] for level in levels[:7] for group in level['groups']}
group_rows = {}
for row in group_book['monGroup'].iter_rows(min_row=4, values_only=True):
    if row[0] in used_groups:
        group_rows[row[0]] = [None if x == 'null' else int(x) for x in row[2].strip('{}').split(',')]
assert used_groups == set(group_rows)
for gid, slots in group_rows.items():
    assert slots == groups[str(gid)]['slots'], ('group workbook/JSON mismatch', gid)

mon_book = load_workbook(config / 'mon.xlsx', read_only=True)
mon_sheet = mon_book['mon']
headers = {cell.value: cell.column - 1 for cell in next(mon_sheet.iter_rows(max_row=1))}
used_enemies = {i for slots in group_rows.values() for i in slots if i is not None}
checked = set()
fields = ('armourType', 'shield', 'shieldType', 'shieldRecovery', 'shieldDelay', 'dmgMultiple', 'reactiveReduction', 'archetype_id', 'intelVisualRef')
for row in mon_sheet.iter_rows(min_row=4, values_only=True):
    if row[headers['id']] not in used_enemies:
        continue
    eid = int(row[headers['id']])
    checked.add(eid)
    for field in fields:
        source = row[headers[field]] if field in headers else None
        projected = enemies[str(eid)].get(field)
        if source in (None, 0) and projected in (None, 0):
            continue
        assert source == projected, ('enemy workbook/JSON mismatch', eid, field, source, projected)
assert used_enemies == checked
visual_resistance = {'scout': 2, 'sentry': 2, 'interceptor': 1, 'swarm': 1, 'anchor': 1, 'reactive': 1, 'shielded': 1, 'command': 1}
for eid in used_enemies:
    enemy = enemies[str(eid)]
    archetype = enemy.get('archetype_id')
    if archetype in visual_resistance:
        assert visual_resistance[archetype] == enemy['armourType'], ('visual identity resistance mismatch', eid)
hyperspace = json.loads((root / 'space-battleship/data/hyperspace_config.json').read_text(encoding='utf-8'))
assert hyperspace['unlock_stage'] == 8, 'highestLevel becomes 8 after stage 7 clear'
print('SOURCE PASS: level formula cache, all 1–7 group/enemy workbook projections, stage7 hyperspace gate semantics')
