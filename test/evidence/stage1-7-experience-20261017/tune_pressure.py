"""Apply one bounded stage 6/7 enemy attack-pressure pass to source and projection."""
import json
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[3] / 'space-battleship'
book_path = root / 'config_excel/mon.xlsx'
json_path = root / 'data/game_data.json'
book = load_workbook(book_path)
sheet = book['mon']
headers = {cell.value: cell.column for cell in sheet[1]}
factors = {
    **{eid: 2.0 for eid in range(600001, 600004)},
    **{eid: 2.5 for eid in range(600004, 600007)},
    **{eid: 2.8 for eid in range(600007, 600010)},
    960003: 2.5,
    **{eid: 2.3 for eid in range(700001, 700004)},
    **{eid: 2.7 for eid in range(700004, 700007)},
    **{eid: 3.2 for eid in range(700007, 700010)},
    970004: 2.7,
}
original = {
    600001: 1.1006, 600002: 1.4, 600003: 0.8503,
    600004: 1.5, 600005: 1.6, 600006: 0.9694,
    600007: 1.7, 600008: 1.8, 600009: 1.8, 960003: 0.4859,
    700001: 1.3, 700002: 1.25, 700003: 1.5,
    700004: 0.9671, 700005: 1.7, 700006: 1.8,
    700007: 1.1236, 700008: 1.9, 700009: 2.0, 970004: 0.5526,
}
data = json.loads(json_path.read_text(encoding='utf-8'))
seen = set()
changed = False
for row in range(4, sheet.max_row + 1):
    value = sheet.cell(row, headers['id']).value
    if not isinstance(value, (int, float)) or int(value) not in factors:
        continue
    eid = int(value)
    old = float(sheet.cell(row, headers['dmgMultiple']).value)
    assert old == float(data['enemies'][str(eid)]['dmgMultiple']), eid
    new = round(original[eid] * factors[eid], 4)
    assert old in (original[eid], new), (eid, old)
    if old != new:
        sheet.cell(row, headers['dmgMultiple']).value = new
        data['enemies'][str(eid)]['dmgMultiple'] = new
        changed = True
    seen.add(eid)
assert seen == set(factors)
if changed:
    book.save(book_path)
    json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print('PRESSURE_ROWS', len(seen), 'changed', changed, 'IDs', sorted(seen))
