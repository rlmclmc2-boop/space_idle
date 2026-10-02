"""Excel -> official incremental projection -> normal-speed weapon consumer test."""
import argparse,copy,json,os,shutil,subprocess,sys
from pathlib import Path
import openpyxl
p=argparse.ArgumentParser();p.add_argument('--project',type=Path,default=Path.cwd());args=p.parse_args()
game=args.project.resolve()
if 'work' not in game.parts:raise SystemExit('Use test/run.py or an isolated test/work project')
sys.path.insert(0,str(game/'tools'))
from config_workbooks import incremental_import
from import_workbook import validate_projection
config=game/'config_excel';target=game/'data/game_data.json'
original={name:(config/(name+'.xlsx')).read_bytes() for name in ['equipment','weapon_motion']}
incremental_import(config,target);base=json.loads(target.read_text(encoding='utf8'))
rejects=0
for section,key,field,value in [('weapon_motion','missile_launch_speed','value',0),('weapon_motion','missile_cruise_at','value',.1),('weapon_motion','missile_turn_rate','value',-1),('weapon_motion','missile_launch_forward_y','value',.1),('equipment','missile','para1',2.5)]:
 bad=copy.deepcopy(base)
 row=bad[section][key][0] if section=='equipment' else bad[section][key]
 row[field]=value
 try:validate_projection(bad)
 except ValueError:rejects+=1
 else:raise AssertionError(f'{section}.{key}.{field} invalid value accepted')
fields={'dmg':180,'dmgMulti':.15,'cd':1.8,'para1':3,'para2':22.5}
changes=[('equipment','missile',field,value) for field,value in fields.items()]
changes.append(('equipment','cannon','para1',50))
values={'player_projectile_pixels_per_unit':35,'enemy_projectile_pixels_per_unit':42,'player_cannon_speed_multiplier':6,'chain_carrier_speed':850,'missile_ejection_gap':.14,'missile_launch_speed':150,'missile_turn_rate':3,'missile_orphan_lifetime':2,'missile_reacquire_interval':.24,'missile_departure_angle':8,'missile_ignition_at':.1,'missile_seek_start':.2,'missile_cruise_at':.8,'missile_lifetime':2,'missile_brake_range':90,'missile_brake_angle':.2,'missile_min_guided_speed':45,'missile_brake_factor':.6,'missile_hit_radius':8,'missile_launch_edge_margin':70,'missile_launch_forward_y':-.2}
changes.extend(('weapon_motion',key,'value',value) for key,value in values.items())
cases=[];parsed=[]
try:
 for name,key,field,value in changes:
  for restore,payload in original.items():(config/(restore+'.xlsx')).write_bytes(payload)
  file=config/(name+'.xlsx');w=openpyxl.load_workbook(file);s=w[name];heads={c.value:c.column for c in s[1] if c.value}
  row=next(r for r in range(4,s.max_row+1) if s.cell(r,heads['name' if name=='equipment' else 'id']).value==key)
  s.cell(row,heads[field]).value=value;w.save(file)
  result=incremental_import(config,target);data=json.loads(target.read_text(encoding='utf8'))
  cases.append({'section':name,'id':key,'field':field,'value':value,'equipment':data['equipment'],'weapon_motion':data['weapon_motion'],'enemy_weapon_base':data['enemy_weapon_base'],'defaults':data['defaults']})
  parsed.append({'case':name+'.'+key+'.'+field,'parsed':result['parsed']})
finally:
 for name,payload in original.items():(config/(name+'.xlsx')).write_bytes(payload)
 incremental_import(config,target)
assert json.loads(target.read_text(encoding='utf8'))['equipment']==base['equipment']
(game/'weapon-config-variants.json').write_text(json.dumps(cases,ensure_ascii=False),encoding='utf8')
(game/'excel-mutation-evidence.json').write_text(json.dumps(parsed,indent=2),encoding='utf8')
engine=os.environ.get('SPACE_BATTLESHIP_GODOT') or shutil.which('godot') or str(game/'engine/Godot_v4.7.2-stable_win64.exe')
script=Path(__file__).with_suffix('.gd')
env=os.environ.copy()
for name in ['XDG_DATA_HOME','XDG_CONFIG_HOME','XDG_CACHE_HOME','APPDATA','LOCALAPPDATA']:
 env[name]=str(game.parent/'weapon-config-userdata'/name);Path(env[name]).mkdir(parents=True,exist_ok=True)
r=subprocess.run([engine,'--headless','--path',str(game),'--script',str(script)],check=False,timeout=30,env=env)
print('CONFIG VALIDATION:',rejects,'invalid inputs rejected')
print('EXCEL MUTATIONS:',len(cases),'official incremental exports; restored original source and projection')
raise SystemExit(r.returncode)
