"""Read-only source checks and isolated projection compatibility."""
from pathlib import Path
import json
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[1] / 'space-battleship'
sys.path.insert(0, str(ROOT / 'tools'))
from import_workbook import read_rows, convert_sheet
from config_workbooks import incremental_import

rows = read_rows(openpyxl.load_workbook(ROOT / 'config_excel/jewel.xlsx', data_only=True, read_only=True)['jewel'])
gems = convert_sheet('jewel', rows)
assert len(gems) == 10 and gems['7']['func'] == rows[6]['func']
assert gems['1']['para_3'] == 1000 and gems['1']['maxLevel'] == 10
try:
    convert_sheet('jewel', rows + [rows[0]])
    raise AssertionError('duplicate id accepted')
except ValueError:
    pass
projected = json.loads((ROOT / 'data/game_data.json').read_text(encoding='utf-8'))
assert projected['jewel'] == gems
config = convert_sheet('config', read_rows(openpyxl.load_workbook(ROOT / 'config_excel/config.xlsx', data_only=True, read_only=True)['config']))
assert projected['config'] == config
assert config['jewelCreat'] == 1000
assert all((ROOT / g['image'].removeprefix('res://')).is_file() for g in gems.values())
levels = convert_sheet('level', read_rows(openpyxl.load_workbook(ROOT / 'config_excel/level.xlsx', data_only=True, read_only=True)['level']))
assert [r['jewelRatio'] for r in projected['levels']] == [r['jewelRatio'] for r in levels]
for invalid in (-1, None, 'bad'):
    try:
        convert_sheet('level', [{'id': 1, 'jewelRatio': invalid, 'monGroup': '1|0.1'}])
        raise AssertionError('invalid jewelRatio accepted')
    except ValueError:
        pass
assert convert_sheet('level', [{'id': 1, 'monGroup': '1|0.1'}])[0]['jewelRatio'] == 1
for invalid in (0, -1, 1.5, None):
    try:
        convert_sheet('config', [{'name': 'jewelCreat', 'para_1': invalid}])
        raise AssertionError('invalid jewelCreat accepted')
    except ValueError:
        pass
legacy_gem = dict(rows[0]); legacy_gem.pop('para_3', None)
assert convert_sheet('jewel', [legacy_gem])['1']['image'].endswith('/1.svg')
# run.py already isolates source/target and user dirs. Simulate a pre-jewel projection.
target = ROOT / '.runtime/jewel-import.json'
projected.pop('jewel')
target.write_text(json.dumps(projected, ensure_ascii=False), encoding='utf-8')
result = incremental_import(ROOT / 'config_excel', target)
assert result['ok']
assert json.loads(target.read_text(encoding='utf-8'))['jewel'] == gems
assert incremental_import(ROOT / 'config_excel', target)['changed'] == []
print('Jewel import: source preservation, duplicate validation, discovery and unchanged cache passed')
