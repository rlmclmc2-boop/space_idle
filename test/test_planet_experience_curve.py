"""Boundary and authoritative-import checks; no game simulation or saves."""
import json
import math
from pathlib import Path
import sys
import unittest

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'test/progression'))
sys.path.insert(0, str(ROOT / 'space-battleship/tools'))
from planet_experience import populate_planet_experience
from import_workbook import convert_sheet, read_rows


class PlanetExperienceTests(unittest.TestCase):
    def test_generation_covers_existing_tail_without_touching_combat(self):
        sheet = openpyxl.Workbook().active
        sheet.append(['id', 'atkRatio', 'planetExpRatio'])
        sheet.append(['description']); sheet.append(['type'])
        for stage in (1, 29, 30, 60, 61, 220):
            sheet.append([stage, '=ROUND(A4*1.2,2)', 12345])
        protected = [row[1].value for row in sheet.iter_rows(min_row=4)]
        populate_planet_experience(sheet)
        self.assertEqual(protected, [row[1].value for row in sheet.iter_rows(min_row=4)])
        self.assertEqual([0, 0, 1], [sheet.cell(row, 3).value for row in range(4, 7)])
        self.assertAlmostEqual(1.2 ** .2, sheet.cell(8, 3).value / sheet.cell(7, 3).value)
        self.assertAlmostEqual(1.2 ** 38, sheet.cell(9, 3).value)

    def test_all_authored_stages_import_with_continuous_curve(self):
        book = openpyxl.load_workbook(ROOT / 'space-battleship/config_excel/level.xlsx', data_only=True)
        projected = convert_sheet('level', read_rows(book.active))
        runtime = json.loads((ROOT / 'space-battleship/data/game_data.json').read_text())['levels']
        self.assertEqual([row['id'] for row in projected], [row['id'] for row in runtime])
        for row, imported in zip(projected, runtime):
            self.assertEqual(row['planetExpRatio'], imported['planetExpRatio'])
            stage = row['id']
            expected = 0 if stage < 30 else 1.2 ** ((stage - 30) / 5)
            self.assertTrue(math.isclose(expected, row['planetExpRatio'], rel_tol=1e-14, abs_tol=1e-14), stage)
        by_id = {row['id']: row for row in runtime}
        self.assertAlmostEqual(1.2 ** .2, by_id[61]['planetExpRatio'] / by_id[60]['planetExpRatio'])
        book.close()


if __name__ == '__main__':
    unittest.main()
