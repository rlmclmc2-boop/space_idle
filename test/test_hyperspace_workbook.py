"""Actual authored-table integration and transactional failure checks."""
import importlib.util
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import openpyxl
ROOT = Path(__file__).resolve().parents[1] / 'space-battleship'
sys.path.insert(0, str(ROOT / 'tools'))
import config_workbooks as cw
import hyperspace_workbook as hw

class HyperspaceExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.folder = self.root / 'config_excel'
        shutil.copytree(ROOT / 'config_excel', self.folder)
        self.target = self.root / 'data/game_data.json'
        self.target.parent.mkdir()
        for name in ('game_data.json', 'hyperspace_config.json'):
            shutil.copy2(ROOT / 'data' / name, self.target.with_name(name))
        cw.incremental_import(self.folder, self.target)
        self.book = self.folder / 'hyperspace_config.xlsx'
    def edit(self, name, value):
        book = openpyxl.load_workbook(self.book)
        for row in book.active.iter_rows(min_row=2):
            if row[0].value == name:
                row[2].value = value
                break
        else: raise AssertionError(name)
        book.save(self.book);book.close()
    def outputs(self):
        return {p.name:p.read_bytes() for p in self.target.parent.iterdir() if p.is_file()}
    def test_whole_authored_config_roundtrip(self):
        self.assertEqual(json.loads((ROOT/'data/hyperspace_config.json').read_text()),hw.read_config(self.book.read_bytes()))
    def test_real_numeric_cell_import_preserves_main_tables_and_noop(self):
        before=self.target.read_bytes();self.edit('/energy_rate',20)
        self.assertEqual(['hyperspace_config'],cw.incremental_import(self.folder,self.target)['changed'])
        self.assertEqual(20,json.loads(self.target.with_name('hyperspace_config.json').read_text())['energy_rate'])
        self.assertEqual(before,self.target.read_bytes())
        state=self.outputs();mtime=self.target.with_name('hyperspace_config.json').stat().st_mtime_ns
        self.assertEqual([],cw.incremental_import(self.folder,self.target)['changed'])
        self.assertEqual(state,self.outputs());self.assertEqual(mtime,self.target.with_name('hyperspace_config.json').stat().st_mtime_ns)
    def test_invalid_literal_leaves_all_outputs_unchanged(self):
        before=self.outputs();self.edit('/energy_rate',-1)
        with self.assertRaises(ValueError):cw.incremental_import(self.folder,self.target)
        self.assertEqual(before,self.outputs())
    def test_formula_rejected_and_missing_authority_rejected(self):
        before=self.outputs();self.edit('/energy_rate','=10*2')
        with self.assertRaises(ValueError):cw.incremental_import(self.folder,self.target)
        self.assertEqual(before,self.outputs());self.book.unlink()
        with self.assertRaises(ValueError):cw.incremental_import(self.folder,self.target)
        self.assertEqual(before,self.outputs())
    def test_duplicate_and_incomplete_array_rejected(self):
        book=openpyxl.load_workbook(self.book);sheet=book.active;sheet.append(('/energy_rate','float',10));book.save(self.book);book.close()
        with self.assertRaises(ValueError):hw.read_config(self.book.read_bytes())
        shutil.copy2(ROOT/'config_excel/hyperspace_config.xlsx',self.book)
        book=openpyxl.load_workbook(self.book);sheet=book.active
        row=next(row for row in sheet.iter_rows(min_row=2) if row[1].value=='array');row[2].value+=1;book.save(self.book);book.close()
        with self.assertRaises(ValueError):hw.read_config(self.book.read_bytes())
    def test_combined_import_write_failure_rolls_back_every_output(self):
        self.edit('/energy_rate',20)
        # Stale state forces a real main-table/state update alongside hyperspace.
        state=self.target.with_name('.import_state.json');state.write_text('{}')
        before=self.outputs();original=cw.os.replace;failed=False
        def replace(src,dest):
            nonlocal failed
            if Path(dest).name=='hyperspace_config.json' and not failed:
                failed=True;raise OSError('Directed final output failure')
            return original(src,dest)
        with patch.object(cw.os,'replace',side_effect=replace):
            with self.assertRaises(OSError):cw.incremental_import(self.folder,self.target)
        self.assertTrue(failed);self.assertEqual(before,self.outputs())

if __name__=='__main__':unittest.main()
