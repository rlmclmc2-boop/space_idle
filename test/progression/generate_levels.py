"""Versioned Excel-source progression candidates. Writes existing fields, imports formally."""
from pathlib import Path
import argparse, json, math, sys, subprocess
import openpyxl
ROOT=Path(__file__).resolve().parents[2]; SRC=ROOT/'space-battleship'; CFG=SRC/'config_excel'
sys.path.insert(0,str(SRC/'tools'))
from config_workbooks import incremental_import
from excel_cache import recache_level
p=argparse.ArgumentParser();p.add_argument('--version',default='progression-v4');p.add_argument('--through',type=int,default=20);p.add_argument('--late-income',type=float,default=.1);p.add_argument('--late-income-step',type=float,default=13);p.add_argument('--roster-through',type=int,default=20);p.add_argument('--smooth-income-floor',action='store_true');p.add_argument('--themed-beam-bosses',action='store_true');p.add_argument('--teaching-fifth-income',type=float,default=1);p.add_argument('--first-reforge-steps',default='8,16,32,48,68');p.add_argument('--future-growth-step',type=float,default=8);p.add_argument('--future-income-step',type=float,default=8);p.add_argument('--future-income-coefficient',type=float,default=4);p.add_argument('--later-cycle-steps',default='6,12,18,30,50');p.add_argument("--boss-health-factor",type=float,default=1.0);p.add_argument("--boss-damage-factor",type=float,default=1.0);p.add_argument("--boss-factor-from",type=int,default=11);p.add_argument("--boss-factor-through",type=int,default=20);p.add_argument("--later-boss-health-factor",type=float,default=1.0);p.add_argument("--later-ultimate-health-factor",type=float);p.add_argument("--later-ultimate-damage-factor",type=float,default=1.0);p.add_argument('--elite-damage-stage',type=int,default=0);p.add_argument('--elite-damage-factor',type=float,default=1.0);p.add_argument('--elite-health-factor',type=float,default=1.0);p.add_argument('--stage-boss-damage-stage',type=int,default=0);p.add_argument('--stage-boss-damage-factor',type=float,default=1.0);p.add_argument('--stage-boss-health-factor',type=float,default=1.0);p.add_argument('--wave-health-factors',default='');p.add_argument('--attack-steps-from30',default='');a=p.parse_args()
wave_health_factors={(int(stage),tier):float(value) for stage,tier,value in [part.split(':') for part in a.wave_health_factors.split(',') if part]}
attack_steps_from30={int(stage):float(value) for stage,value in [part.split(':') for part in a.attack_steps_from30.split(',') if part]}
from source_lock import acquire
_source_lock=acquire(ROOT)
data=json.loads((SRC/'data/game_data.json').read_text())
books={n:openpyxl.load_workbook(CFG/(n+'.xlsx')) for n in ['level','mon','monGroup']}
groups_sheet=books['monGroup'].active
if 'combatTier' not in [cell.value for cell in groups_sheet[1]]:
 groups_sheet.cell(1,groups_sheet.max_column+1,'combatTier')
 groups_sheet.cell(2,groups_sheet.max_column,'Runtime wave tier, separate from final-wave completion')
 groups_sheet.cell(3,groups_sheet.max_column,'string')
sheets={n:b.active for n,b in books.items()};headers={n:{c.value:c.column for c in s[1] if c.value is not None} for n,s in sheets.items()}
indices={n:{s.cell(r,headers[n]["id"]).value:r for r in range(4,s.max_row+1)} for n,s in sheets.items()}
next_rows={n:s.max_row+1 for n,s in sheets.items()}
def replace(n,key,row):
 s=sheets[n]; h=headers[n]
 if key in indices[n]:ri=indices[n][key]
 else:ri=next_rows[n];next_rows[n]+=1;indices[n][key]=ri
 for k,v in row.items():
  if k in h:s.cell(ri,h[k],v)
def source_row(n,key):
 s=sheets[n];h=headers[n];ri=indices[n][key]
 return {k:s.cell(ri,c).value for k,c in h.items()}
for design in data['battle_design'].values():
 replace('monGroup',int(design['group_id']),{'combatTier':design['tier']})
themes=['laser','missile','cannon','neutral1','neutral2','longLaser','laser','missile','cannon','neutral3']
themes += ['neutral1','missile','neutral2','laser','cannon','neutral4','longLaser','neutral3','missile','laser']
# Later levels normally mix themes; beam has no fixed tutorial slot after10.
cycle=['neutral2','cannon','neutral4','missile','laser','neutral1','longLaser','neutral3','cannon','neutral4','missile','neutral2','laser','neutral3','longLaser','neutral1']
while len(themes)<a.roster_through:themes.extend(cycle)
ratios=[1,1.5,2,3,4.5,1.2**15]+[1.2**(30+10*i) for i in range(4)]
ratios += [1.2**(60+13*i) for i in range(1,11)]
resources=[1.4**i for i in range(5)]+[12*1.2**(10*i) for i in range(5)]
resources[4]*=a.teaching_fifth_income
resources += [resources[9]*1.2**13*a.late_income*1.2**(a.late_income_step*(i-1)) for i in range(1,11)]
if a.smooth_income_floor:
 for index in range(10,20):resources[index]=max(resources[index],resources[index-1]*1.4)
if a.through>20:
 # Provisional future budget. Every value is subject to segment regression.
 ratios += [1.2**(190+a.future_growth_step*i) for i in range(1,11)]
 resources += [resources[19]*1.2**(a.future_income_step*i)*a.future_income_coefficient for i in range(1,11)]
 first_steps=[float(value) for value in a.first_reforge_steps.split(",")]
 if len(first_steps)!=5:raise ValueError("First reforge requires five stage growth steps")
 first_income=[0,0,4,4,10]
 ratios += [ratios[29]*1.2**i for i in first_steps]
 resources += [resources[29]*1.2**i for i in first_income]
 later_steps=[float(value) for value in a.later_cycle_steps.split(",")]
 if len(later_steps)!=5:raise ValueError("Later cycles require five stage growth steps")
 for planet_cycle in range(5):
  base_ratio=ratios[-1];base_income=resources[-1]
  ratios += [base_ratio*1.2**i for i in later_steps]
  resources += [base_income*1.2**i for i in [4,8,12,20,32]]
manifest=[]
for stage,theme in enumerate(themes[:max(a.through,a.roster_through)],1):
 tiers=['normal']*4+['elite'] if stage<=5 else ['normal']*5+['elite']*3+['boss']
 if stage>=20 and ((stage<=70 and stage%5==0) or (stage>70 and stage%10==0)):tiers=['elite']*5+['boss']*3+['ultimate']
 groups=[]
 for node,tier in enumerate(tiers,1):
  variant='physical_attack' if (stage+node)%2==0 else 'energy_attack'
  if tier in ['boss','ultimate']:key=tier+('_energy' if theme in ['missile','cannon'] else '_physical')
  else:key=tier+'_'+theme
  # Neutral elites carry numeric suffix; their source IDs exist.
  gid=int(data['battle_design'][key]['group_id'])+(1000 if variant=='physical_attack' else 0)
  original=data['groups'][str(gid)];newgid=30000+stage*10+node;slots=[]
  for slot,eid in enumerate(original['slots']):
   if eid is None:slots.append('null');continue
   newid=40000+stage*1000+node*20+slot;row=source_row('mon',eid);row['id']=newid
   row['des']=str(row['des'])
   if a.themed_beam_bosses and theme=='longLaser' and tier in ['boss','ultimate']:
    # Same authored recovery rule as the level's ordinary/elite enemies.
    # Split total base defence; keep every attack variant and formation intact.
    total=float(row['health'])+float(row['shield'] or 0)
    row.update(health=round(total*.5),shield=round(total*.5),armourType=2,shieldType=2,shieldRecovery=.2,shieldDelay=.3)
   row['health']=max(1,round(float(row['health'])*.1454));row['shield']=max(0,round(float(row['shield'] or 0)*.1454))
   loot=(8 if tier=='normal' else 20) if stage<=5 else max(1,round((30 if tier=='normal' else 60 if tier=='elite' else 120)/len([e for e in original['slots'] if e is not None])))
   row['res']='{1,%d,1}'%loot
   row['dmgMultiple']=float(row['dmgMultiple'])*.1454
   if stage==a.elite_damage_stage and tier=='elite':
    row['dmgMultiple']*=a.elite_damage_factor
    row['health']=max(1,round(row['health']*a.elite_health_factor));row['shield']=max(0,round(row['shield']*a.elite_health_factor))
   if stage==a.stage_boss_damage_stage and tier=='boss':
    row['dmgMultiple']*=a.stage_boss_damage_factor
    row['health']=max(1,round(row['health']*a.stage_boss_health_factor));row['shield']=max(0,round(row['shield']*a.stage_boss_health_factor))
   if (stage,tier) in wave_health_factors:
    factor=wave_health_factors[(stage,tier)];row['health']=max(1,round(row['health']*factor));row['shield']=max(0,round(row['shield']*factor))
   if tier in ["boss","ultimate"] and a.boss_factor_from<=stage<=a.boss_factor_through:
    row["health"]=max(1,round(float(row["health"])*a.boss_health_factor));row["shield"]=max(0,round(float(row["shield"])*a.boss_health_factor))
    row["dmgMultiple"]*=a.boss_damage_factor
   if tier in ["boss","ultimate"] and stage>=21:
    health_factor=a.later_ultimate_health_factor if tier=="ultimate" and a.later_ultimate_health_factor is not None else a.later_boss_health_factor
    row["health"]=max(1,round(float(row["health"])*health_factor));row["shield"]=max(0,round(float(row["shield"])*health_factor))
    if tier=="ultimate":row["dmgMultiple"]*=a.later_ultimate_damage_factor
   # Normalize attacks with HP for original level10 -> early-game scaling; keep actual damage types.
   replace('mon',newid,row);slots.append(str(newid))
  replace('monGroup',newgid,{'id':newgid,'des':f'第{stage}关第{node}战点','combatTier':tier,'mon':'{'+','.join(slots)+'}'})
  groups.append((newgid,node*.9/len(tiers)))
  manifest.append({'stage':stage,'node':node,'tier':tier,'theme':theme,'source_group':gid,'group':newgid,'variant':variant,'suggested':stage in [4,5],'themed_recovery_shield':a.themed_beam_bosses and theme=='longLaser' and tier in ['boss','ultimate']})
 row=source_row('level',stage)
 row.update(length=4000 if stage<=5 else 1000,monGroup='{'+','.join(f'{g}|{pos:.6f}' for g,pos in groups)+'}')
 if stage<=a.through:
  row.update(atkRatio=ratios[stage-1],lifeRatio=ratios[stage-1],resRatio=resources[stage-1])
  row['planetExpRatio']=0 if stage<30 else 1.2**((stage-30)/5)
  if stage in attack_steps_from30:row['atkRatio']=ratios[29]*1.2**attack_steps_from30[stage]
 # Uncalibrated future numeric cells/formulas retain their identity, not old cached values.
 replace('level',stage,row)
for n,b in books.items():b.save(CFG/(n+'.xlsx'))
recache_level(CFG/"level.xlsx",sheets["level"],subprocess.check_output(["git","show","04a5a307bcef9325efa9026e1ca94affa577e10d:space-battleship/config_excel/level.xlsx"],cwd=ROOT))
result=incremental_import(CFG,SRC/'data/game_data.json')
evidence=ROOT/'test/progression/candidates';evidence.mkdir(exist_ok=True)
meta={'version':a.version,'stage_range':[1,a.through],'roster_range':[1,max(a.through,a.roster_through)],
 'reason':'Segmented numerical candidate. Old basic timings are diagnostic only. Explicit income/defence changes require fresh formal-scene tests; stage4/5 remain suggestions. Original40 identities retained, with numerical pair compensation recorded separately. Future formula dependencies are recalculated and unaccepted.',
 'health_normalization':.1454,'ratios':ratios,'resource_ratios':resources,
 'late_income_coefficient':a.late_income,'late_income_step':a.late_income_step,
 'smooth_income_floor':a.smooth_income_floor,'parameters':vars(a),
 'affected':'Generated numeric range and roster, plus later formula dependencies. Future candidate is provisional until segmented and single-version full fresh tests.',
 'waves':manifest,'import':result}
(evidence/(a.version+'.json')).write_text(json.dumps(meta,indent=2,ensure_ascii=False))
print(json.dumps(result))
