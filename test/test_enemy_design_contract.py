"""Candidate formation/source contract; run from repository root."""
import copy
import io
import json
import pathlib
import subprocess
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
PROJECT = ROOT / "space-battleship"
BASE = "a22fd7d4c12d8f573c332cfc9e5ef85ffc19337d"
sys.path.insert(0, str(PROJECT / "tools"))
from import_workbook import validate_projection

data = json.loads((PROJECT / "data/game_data.json").read_text())
validate_projection(data)
records = data["battle_design"]
assert len(records) == 20
assert {tier: sum(r["tier"] == tier for r in records.values())
        for tier in ("normal", "elite", "boss", "ultimate")} == {
            "normal": 8, "elite": 8, "boss": 2, "ultimate": 2}
signatures = set()
for record in records.values():
    slots = data["groups"][str(record["group_id"])]["slots"]
    assert len(slots) == 15
    occupied = [data["enemies"][str(e)] for e in slots if e is not None]
    sizes = [int(e["size"]) for e in occupied]
    tier = record["tier"]
    if tier == "normal":
        assert max(sizes) <= 3
    elif tier == "elite":
        assert sizes.count(4) == 2 and all(s in (1, 2, 4) for s in sizes)
    elif tier == "boss":
        assert sizes.count(5) <= 1 and sizes.count(4) <= 2
        assert all(s in (1, 2, 4, 5) for s in sizes)
    else:
        assert sizes == [6]
    for row in range(3):
        assert slots[row*5:row*5+5] == slots[row*5:row*5+5][::-1]
    signature = tuple(sorted((e["health"], e["size"], e["armourType"],
                              e["dmgMultiple"], tuple(m["name"] for m in e["equipment"]))
                             for e in occupied))
    assert signature not in signatures, record["id"]
    signatures.add(signature)

for name in ("equipment", "mon", "monGroup", "level", "weapon_motion"):
    relative = f"space-battleship/config_excel/{name}.xlsx"
    old_bytes = subprocess.check_output(["git", "show", f"{BASE}:{relative}"], cwd=ROOT)
    old = openpyxl.load_workbook(io.BytesIO(old_bytes)).active
    new = openpyxl.load_workbook(ROOT / relative).active
    allowed = {"E8", "G8"} if name == "equipment" else set()
    for row in old:
        for cell in row:
            assert cell.value == new[cell.coordinate].value or cell.coordinate in allowed, (name, cell.coordinate)
    if name in ("level", "weapon_motion"):
        assert old_bytes == (ROOT / relative).read_bytes(), name
    # Growth and cost fields must remain authored values even in the allowed row.
    if name == "equipment":
        growth_columns = [c.column for c in old[1] if c.value and "Multi" in str(c.value)] + [16,18]
        for row in range(4, old.max_row+1):
            for col in growth_columns:
                assert old.cell(row,col).value == new.cell(row,col).value

def rejects(changed):
    try:
        validate_projection(changed)
    except (ValueError, TypeError):
        return
    raise AssertionError("invalid draft accepted")

invalid = copy.deepcopy(data)
invalid["groups"]["1001"]["slots"].append(None)
rejects(invalid)
for value in (True, 1.5, -1, 4):
    invalid = copy.deepcopy(data)
    invalid["battle_design"]["normal_laser"]["min_upgrade"] = value
    rejects(invalid)
print("Enemy design contract: 20 unique symmetric groups, size rules, legacy rows, growth, level/motion bytes, malformed drafts passed")
