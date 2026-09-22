"""Unified gates, source projection, and import transactions; run via run.py."""
import copy
import json
from pathlib import Path
import sys
import unittest
import openpyxl

ROOT = Path(__file__).resolve().parents[1] / 'space-battleship'
sys.path.insert(0, str(ROOT / 'tools'))
from config_workbooks import incremental_import, read_changed_file
from import_workbook import convert_sheet, validate_projection
from level_editor_store import Store


class UnlockImportTests(unittest.TestCase):
    def setUp(self):
        self.data = json.loads((ROOT / 'data/game_data.json').read_text(encoding='utf-8'))

    def test_source_and_unique_complete_targets(self):
        path = ROOT / 'config_excel/unlock.xlsx'
        rows = read_changed_file(path, 'unlock', path.read_bytes())
        self.assertEqual(rows, self.data['unlock'])
        self.assertEqual(len(rows), 25)
        self.assertEqual(len({(r['type'], r['target']) for r in rows.values()}), 25)
        validate_projection(self.data)
        for kind in ('equipment', 'ship', 'charge', 'hightech'):
            for row in self.data[kind].values():
                for item in row if isinstance(row, list) else [row]:
                    self.assertNotIn('unlock', item)
        self.assertNotIn('jewelDropLevel', self.data['config'])

    def test_invalid_rows_rejected(self):
        for field, value in [('level', -1), ('level', 1.5), ('level', 999), ('level', True),
                             ('title', ''), ('desc', None), ('mode', 'unknown'),
                             ('target', 'missing'), ('type', 'unknown')]:
            with self.subTest(field=field, value=value):
                data = copy.deepcopy(self.data)
                data['unlock']['shield'][field] = value
                with self.assertRaises(ValueError): validate_projection(data)
        data = copy.deepcopy(self.data)
        del data['unlock']['shield']
        with self.assertRaises(ValueError): validate_projection(data)
        data = copy.deepcopy(self.data)
        data['unlock']['duplicate'] = {**data['unlock']['shield'], 'name':'duplicate'}
        with self.assertRaises(ValueError): validate_projection(data)
        row = self.data['unlock']['shield']
        with self.assertRaises(ValueError): convert_sheet('unlock', [row, row])

    def test_legacy_gates_cannot_override(self):
        expected = copy.deepcopy(self.data['unlock'])
        self.data['ship']['Destroyer']['unlock'] = 1
        self.data['equipment']['shield'][0]['unlock'] = 1
        for kind in ('charge','hightech'):
            for row in self.data[kind].values(): row['unlock'] = 1
        self.data['config']['jewelDropLevel'] = 1
        validate_projection(self.data)
        self.assertEqual(self.data['unlock'], expected)
        self.assertNotIn('unlock', self.data['ship']['Destroyer'])

    def test_incremental_and_failure_atomicity(self):
        directory = ROOT / 'config_excel'
        target = ROOT / 'data/game_data.json'
        path = directory / 'unlock.xlsx'
        original = path.read_bytes()
        try:
            incremental_import(directory, target)
            book = openpyxl.load_workbook(path)
            sheet = book['unlock']
            shield = next(row[0].row for row in sheet if row[0].value == 'shield')
            sheet.cell(shield, 4, 6)
            sheet.cell(shield, 6, 'Changed title')
            book.save(path); book.close()
            result = incremental_import(directory, target)
            self.assertEqual(result['changed'], ['unlock'])
            current = json.loads(target.read_text(encoding='utf-8'))
            self.assertEqual(current['unlock']['shield']['level'], 6)
            self.assertEqual(current['unlock']['shield']['title'], 'Changed title')
            before = target.read_bytes()
            book = openpyxl.load_workbook(path)
            book['unlock'].cell(shield, 4, -1)
            book.save(path); book.close()
            with self.assertRaises(ValueError): incremental_import(directory, target)
            self.assertEqual(target.read_bytes(), before)
        finally:
            path.write_bytes(original)
            incremental_import(directory, target)

    def test_level_editor_discovers_table(self):
        store = Store(ROOT)
        self.assertIn('unlock', store.paths)
        result = store.execute(store.load(), False)
        self.assertTrue(result)


if __name__ == '__main__': unittest.main()
