"""Run via run.py. Crew participates in existing workbook projection and transactions."""
import copy,json,sys,unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
ROOT=Path(__file__).resolve().parents[1]/'space-battleship'
sys.path.insert(0,str(ROOT/'tools'))
from import_workbook import convert_sheet,validate_crew,full_import,SECTIONS
from config_workbooks import incremental_import,read_changed_file

class CrewImportTests(unittest.TestCase):
    def setUp(self):
        self.data=json.loads((ROOT/'data/game_data.json').read_text(encoding='utf-8'))

    def test_sources(self):
        self.assertFalse((ROOT/'config_excel/crew_level.xlsx').exists())
        self.assertNotIn('crew_level',self.data)
        self.assertEqual(self.data['crew_config']['base_exp']['value'],100)
        self.assertEqual(self.data['crew_config']['exp_multiplier']['value'],1.2)
        for name in ('crew','crew_assignment','crew_config'):
            path=ROOT/'config_excel'/f'{name}.xlsx'
            self.assertEqual(read_changed_file(path,name,path.read_bytes()),self.data[name])

    def test_invalid_rows(self):
        row=self.data['crew']['navigator']
        with self.assertRaises(ValueError):convert_sheet('crew',[row,row])
        for section,id,field,value in [('crew','navigator','basePower',float('nan')),('crew','navigator','unlockId','missing'),('crew_assignment','equipment_upgrade','interval',0),('crew_assignment','jewel_auto','interval',0),('crew_assignment','reactor_upgrade','interval',0),('crew_assignment','reactor_upgrade','baseValue',0),('crew_assignment','equipment_upgrade','maxCrew',1.5)]:
            data=copy.deepcopy(self.data);data[section][id][field]=value
            with self.assertRaises(ValueError):validate_crew(data)
        for field,value in [('base_exp',0),('base_exp',float('nan')),('exp_multiplier',0.9),('exp_multiplier',float('inf'))]:
            data=copy.deepcopy(self.data);data['crew_config'][field]['value']=value
            with self.assertRaises(ValueError):validate_crew(data)
        legacy=dict(row,baseLevel=1,maxLevel=3,expGroup='normal',baseExp=999,expGrowth=9)
        self.assertEqual(convert_sheet('crew',[legacy]),{row['id']:row})

    def test_incremental_and_optional_discovery(self):
        target=ROOT/'.runtime/crew-import.json'
        previous=copy.deepcopy(self.data);previous['crew_level']={'normal':{}}
        target.write_text(json.dumps(previous,ensure_ascii=False),encoding='utf-8')
        result=incremental_import(ROOT/'config_excel',target)
        self.assertTrue(set(('crew','crew_assignment','crew_config')).issubset(result['parsed']))
        projected=json.loads(target.read_text(encoding='utf-8'))
        self.assertNotIn('crew_level',projected)
        self.assertNotIn('crew_level',SECTIONS)
        for name in ('crew','crew_assignment','crew_config'):self.assertEqual(projected[name],self.data[name])
        raw=target.read_bytes()
        self.assertEqual(incremental_import(ROOT/'config_excel',target)['changed'],[])
        self.assertEqual(target.read_bytes(),raw)

    def test_full_import_preserves_optional_crew_from_older_source(self):
        # Read real current split sheets as a single workbook view; the old master
        # has known unrelated missing fields and is not a valid rule baseline.
        books={name:openpyxl.load_workbook(ROOT/'config_excel'/f'{name}.xlsx',read_only=True,data_only=True) for name in SECTIONS if not name.startswith('crew')}
        class View:
            sheetnames=list(books)+["crew_level"]
            def __getitem__(self,name):return books[name][name]
            def close(self):pass
        previous=copy.deepcopy(self.data);previous['crew_level']={'normal':{}}
        target=ROOT/'.runtime/crew-full.json';target.write_text(json.dumps(previous),encoding='utf-8')
        try:
            with patch('import_workbook.openpyxl.load_workbook',return_value=View()):
                result=full_import(Path('current-legacy-view.xlsx'),target)
            self.assertNotIn('crew_level',result)
            for name in ('crew','crew_assignment','crew_config'):self.assertEqual(result[name],self.data[name])
        finally:
            for book in books.values():book.close()

if __name__=='__main__':unittest.main()
