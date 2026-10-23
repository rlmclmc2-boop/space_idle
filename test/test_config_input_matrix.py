"""Published configuration boundaries: isolated inputs and real public consumers.

Run only through test/run.py. No production defaults or validators are patched.
Fault injection verifies current publication/recovery, not retired direct writers.
"""
from contextlib import ExitStack
import io
import json
from pathlib import Path
import shutil
import sys
from types import SimpleNamespace
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1] / 'space-battleship'
assert ROOT.parent.parent.name == 'work', 'Use test/run.py; never run on formal inputs'
sys.path.insert(0, str(ROOT / 'tools'))
import config_workbooks as cw
import level_editor_store as le
import import_workbook as full

AREA = ROOT.parent / 'phase8-matrix'
AREA.mkdir(exist_ok=True)
originals = {p.relative_to(ROOT).as_posix():p.read_bytes()
             for folder in ('data','config_excel') for p in (ROOT/folder).rglob('*') if p.is_file()}
originals['configuration_manifest.json']=(ROOT/'configuration_manifest.json').read_bytes()
originals['.configuration-lock']=b''

checks = 0


def check(condition, label):
    global checks
    checks += 1
    if not condition:
        raise AssertionError(label)


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(full.encode(data))


def fresh(name):
    root = AREA / name
    assert root.resolve().is_relative_to(AREA.resolve()) and root != AREA
    if root.exists():
        shutil.rmtree(root)
    for name, raw in originals.items():
        path = root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(raw)
    return root


def snapshot(root):
    return {str(p.relative_to(root)).replace('\\', '/'): cw.sha(p.read_bytes())
            for p in sorted(root.rglob('*')) if p.is_file()}


def delta(before, after):
    return {key: ('created' if key not in before else 'deleted' if key not in after else 'modified')
            for key in sorted(before.keys() | after.keys()) if before.get(key) != after.get(key)}


def xml_cell(root, sheet, field, value=None, missing_cache=False):
    path = root / 'config_excel' / (sheet + '.xlsx')
    table = le.Store(root).load()['tables'][sheet]
    address = le.get_column_letter(table['headers'].index(field) + 1) + '4'
    with zipfile.ZipFile(path) as book:
        part = dict(cw.sheet_parts(book))[sheet]
        tree = cw.ET.fromstring(book.read(part))
        cell = next(c for c in tree.iter(cw.Q + 'c') if c.get('r') == address)
        if missing_cache:
            cell = next(c for c in tree.iter(cw.Q + 'c') if c.find(cw.Q + 'f') is not None)
            check(cell.find(cw.Q + 'v') is not None, 'remove an existing real formula cache')
            cached = cell.find(cw.Q + 'v')
            if cached is not None:
                cell.remove(cached)
        else:
            for child in list(cell):
                cell.remove(child)
            cell.attrib.pop('t', None)
            if isinstance(value, str):
                cell.set('t', 'inlineStr')
                cw.ET.SubElement(cw.ET.SubElement(cell, cw.Q + 'is'), cw.Q + 't').text = value
            else:
                cw.ET.SubElement(cell, cw.Q + 'v').text = str(value)
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w') as dest:
            for info in book.infolist():
                dest.writestr(info, cw.ET.tostring(tree) if info.filename == part else book.read(info.filename))
    path.write_bytes(output.getvalue())


def setup(root, case):
    manifest = root / 'config_excel' / cw.MANIFEST
    target = root / 'data/game_data.json'
    mapping = json.loads(manifest.read_text(encoding='utf-8'))
    if case == 'manifest_missing': manifest.unlink()
    elif case == 'manifest_broken': manifest.write_text('{', encoding='utf-8')
    elif case.startswith('manifest_value_'):
        values = {'empty': {}, 'list': [], 'sheets_missing': {'version': 1},
                  'sheets_null': {'sheets': None}, 'sheets_list': {'sheets': []}}
        write_json(manifest, values[case.removeprefix('manifest_value_')])
    elif case.startswith('path_'):
        values = {'empty': '', 'null': None, 'number': 42, 'parent': '../mon.xlsx',
                  'backslash': '..\\mon.xlsx', 'absolute': str(root / 'mon.xlsx'),
                  'missing_file': 'absent.xlsx', 'wrong_sheet': 'level.xlsx'}
        mapping['sheets']['mon'] = values[case.removeprefix('path_')]
        write_json(manifest, mapping)
    elif case == 'required_mapping_missing':
        del mapping['sheets']['mon']; write_json(manifest, mapping)
    elif case.startswith('optional_'):
        _, name, variant = case.split('_', 2)
        mapping['sheets'].pop(name, None)
        if variant == 'null': mapping['sheets'][name] = None
        if variant == 'empty': mapping['sheets'][name] = ''
        write_json(manifest, mapping)
        if variant.startswith('absent'):
            (root / 'config_excel' / (name + '.xlsx')).unlink()
            if variant == 'absent_no_projection':
                data = json.loads(target.read_text(encoding='utf-8'))
                data.pop(full.SECTIONS[name], None); write_json(target, data)
    elif case == 'target_missing': target.unlink()
    elif case == 'target_broken': target.write_text('{', encoding='utf-8')
    elif case == 'cache_broken': target.with_name('.import_state.json').write_text('{', encoding='utf-8')
    elif case in ('cache_valid', 'cache_old', 'cache_wrong_target'):
        cw.incremental_import(root / 'config_excel', target)
        state = target.with_name('.import_state.json')
        data = json.loads(state.read_text(encoding='utf-8'))
        if case == 'cache_old': data['version'] = -1
        if case == 'cache_wrong_target': data['target_hash'] = 'bad'
        write_json(state, data)
    elif case.startswith('invalid_'):
        values = {'health': ('mon', 'health', -1), 'size': ('mon', 'size', 1.5),
                  'armour': ('mon', 'armourType', 9), 'weapon': ('mon', 'equipment', '{missing|1}'),
                  'drop': ('mon', 'res', '{999,1,1}'), 'id': ('mon', 'id', 2),
                  'reference': ('monGroup', 'mon', '{999,null,null,null,null,null,null,null,null,null}')}
        xml_cell(root, *values[case.removeprefix('invalid_')])
    elif case == 'missing_formula_cache': xml_cell(root, 'level', 'atkRatio', missing_cache=True)
    elif case == 'corrupt_xlsx': (root / 'config_excel/mon.xlsx').write_bytes(b'not a workbook')
    elif case == 'extra_json':
        data = json.loads(target.read_text(encoding='utf-8'))
        data['phase8_sentinel'] = {'unchanged': [1, 'x', None]}
        data['defaults']['autoCollectDelay'] = 12
        write_json(target, data)


CASES = ['valid', 'manifest_missing', 'manifest_broken',
         *['manifest_value_' + x for x in ('empty', 'list', 'sheets_missing', 'sheets_null', 'sheets_list')],
         'required_mapping_missing',
         *['path_' + x for x in ('empty', 'null', 'number', 'parent', 'backslash', 'absolute', 'missing_file', 'wrong_sheet')],
         *['optional_' + n + '_' + v for n in ('ship',) for v in
           ('fallback', 'null', 'empty', 'absent_existing_projection', 'absent_no_projection')],
         'target_missing', 'target_broken', 'cache_broken', 'cache_valid', 'cache_old', 'cache_wrong_target',
         *['invalid_' + x for x in ('health', 'size', 'armour', 'weapon', 'drop', 'id', 'reference')],
         'missing_formula_cache', 'corrupt_xlsx', 'extra_json']


def run_case(case, entry):
    root = fresh(case + '-' + entry)
    setup(root, case)
    before = snapshot(root)
    before_times = {str(p.relative_to(root)): p.stat().st_mtime_ns for p in root.rglob('*') if p.is_file()}
    stage = 'import' if entry == 'incremental' else 'construct'
    try:
        if entry == 'incremental':
            result = cw.incremental_import(root/'config_excel',root/'data/game_data.json')
        else:
            store=le.Store(root);stage='load';request=store.load()
            if entry=='store_load':result=request
            else:
                stage='save' if entry=='store_save' else 'validate'
                result=store.execute(request,entry=='store_save')
        outcome={'accepted':True,'stage':stage,'returned':result}
    except Exception as error:
        outcome={'accepted':False,'stage':stage,'error_type':type(error).__name__,'error':str(error)}
    parsed=[]
    if outcome['accepted'] and entry=='incremental':parsed=outcome['returned'].get('parsed',[])
    after = snapshot(root)
    if case == 'valid':
        check(outcome['accepted'], 'valid input must succeed: ' + entry + ': ' + str(outcome))
    outcome.update(case=case, entry=entry, files=delta(before, after), unchanged=sorted(k for k in before if after.get(k)==before[k]),
                   parsed=parsed, file_hashes=after)
    if not outcome['accepted'] or entry in ('store_load', 'store_validate'):
        check(after == before, case + '/' + entry + ': read/rejection must not write')
    if case == 'cache_valid' and entry == 'incremental':
        check(after == before and not parsed, 'unchanged import is a no-op')
        check(before_times == {str(p.relative_to(root)):p.stat().st_mtime_ns for p in root.rglob('*') if p.is_file()}, 'no-op preserves mtimes')
    # Retain complete return data, errors and effects; normalize only the sandbox prefix.
    return json.loads(json.dumps(outcome, ensure_ascii=False).replace(str(root).replace('\\','\\\\'), '<CASE>'))


def projection_and_excel():
    root = fresh('projection')
    target = root/'data/game_data.json'
    data = json.loads(target.read_text(encoding='utf-8'))
    data['phase8_sentinel']={'unknown':[1,2]}
    import config_snapshot as protocol
    current=protocol.Store(root).capture();files=dict(current.files);files['data/game_data.json']=full.encode(data)
    protocol.Store(root).publish(current,files,lambda _:None)
    cw.incremental_import(root/'config_excel',target)
    before = json.loads(target.read_text(encoding='utf-8'))
    xml_cell(root,'mon','health',321)
    result = cw.incremental_import(root/'config_excel',target)
    after = json.loads(target.read_text(encoding='utf-8'))
    check(result['parsed']==['mon'],'single-table parses only mon')
    check({k:v for k,v in before.items() if k!='enemies'}=={k:v for k,v in after.items() if k!='enemies'},'other JSON fields unchanged')
    state = json.loads(target.with_name('.import_state.json').read_text(encoding='utf-8'))
    check(state['version']==cw.CACHE_VERSION and state['target_hash']==cw.sha(target.read_bytes()),'cache and target hash')
    store=le.Store(root)
    store.paths=le.PreparedStore(root).paths
    store.directory=root/'config_excel'
    check(state['hashes']=={n:cw.sha(p.read_bytes()) for n,p in store.paths.items()},'source hashes')
    request=store.load(); old={p:p.read_bytes() for p in [*store.paths.values(),target,target.with_name('.import_state.json')]}
    # Level ratios are protected, so exercise formula evaluation on editable health cells.
    table=request['tables']['mon']
    health_column=le.get_column_letter(table['headers'].index('health')+1)
    table['rows'][0]['health']='=ROUND(1.005,2)'
    table['rows'][1]['health']='='+health_column+'4/3'
    first_id, second_id = (str(row['id']) for row in table['rows'][:2])
    result=store.execute(request,True)
    projected=json.loads(target.read_text(encoding='utf-8'))
    check(projected['enemies'][first_id]['health']==1.01 and projected['enemies'][second_id]['health']==1.01/3,'ROUND and Decimal/float reference boundary')
    backup=Path(result['backup'])
    for p in (store.paths['mon'],target,target.with_name('.import_state.json')):
        check((backup/p.relative_to(root)).read_bytes()==old[p],'backup original bytes')
    for name,p in store.paths.items():
        if name!='mon': check(p.read_bytes()==old[p],'unmodified workbook bytes '+name)
    with zipfile.ZipFile(io.BytesIO(old[store.paths['mon']])) as a, zipfile.ZipFile(store.paths['mon']) as b:
        part=dict(cw.sheet_parts(a))['mon']
        check(a.namelist()==b.namelist(),'ZIP members/order')
        for name in a.namelist():
            if name!=part: check(a.read(name)==b.read(name),'unchanged ZIP part '+name)
        tree=cw.ET.fromstring(b.read(part)); cell=next(c for c in tree.iter(cw.Q+'c') if c.get('r')==health_column+'4')
        check(cell.find(cw.Q+'f').text=='ROUND(1.005,2)' and float(cell.find(cw.Q+'v').text)==1.01,'formula and XML cache')
    check(cw.incremental_import(store.directory,target)['parsed']==[],'Store fingerprint accepted by incremental')
    before=snapshot(root)
    try: full.full_import(ROOT.parent/'太空战舰.xlsx',target)
    except ValueError as error:
        check(bool(str(error)),'historical aggregate rejection includes a reason')
    else: raise AssertionError('Historical aggregate unexpectedly accepted')
    check(snapshot(root)==before and not target.with_suffix('.json.tmp').exists(),'full-import failure preserves output')


def published_editor_protocol():
    """The real public importer/editor share complete snapshots and publication."""
    import config_publication as publication
    import config_snapshot as protocol
    root=AREA/'published-editor'
    root.mkdir()
    for name in ('data','config_excel'):shutil.copytree(ROOT/name,root/name)
    imported=cw.incremental_import(root/'config_excel',root/'data/game_data.json')
    initial=protocol.Store(root).capture()
    check(imported['generation']==initial.identity,'Public import publishes a complete validated generation')
    editor=le.Store(root);request=editor.load()
    check(set(request['tables'])==set(le.EDITABLE),'Public editor loads frozen source tables')
    editor.execute(request,False)
    check(protocol.Store(root).capture().identity==initial.identity,'Validation does not publish')
    request['tables']['mon']['rows'][0]['health']*=1.01
    saved=editor.execute(request,True);current=protocol.Store(root).capture()
    check(current.identity!=initial.identity and saved['snapshot']['generation']==current.identity,'Editor saves source and projection under one publication')
    check(initial.files['data/game_data.json']!=current.files['data/game_data.json'],'Existing consumer retains original captured bytes')
    try:editor.execute(request,True)
    except ValueError:pass
    else:raise AssertionError('Public editor accepted stale request')
    check(protocol.Store(root).capture().identity==current.identity,'Stale editor cannot write')
    source=root/'config_excel/mon.xlsx'
    with protocol.exclusive(source):
        try:editor.load()
        except protocol.Busy:pass
        else:raise AssertionError('Editor captured workbook while external writer held it')
    check(protocol.Store(root).capture().identity==current.identity,'Editable sources do not invalidate committed JSON readers')
    import openpyxl
    complete=AREA/'valid-current-master.xlsx'
    book=openpyxl.Workbook(write_only=True)
    captured=protocol.Store(root).capture_publication()[2]
    paths=le.PreparedStore(root).paths
    for name,path in paths.items():
        source_book=openpyxl.load_workbook(io.BytesIO(captured[path.relative_to(root).as_posix()]),read_only=True,data_only=True)
        sheet=book.create_sheet(name)
        for row in source_book[name].iter_rows(values_only=True):sheet.append(row)
        source_book.close()
    book.save(complete)
    full.full_import(complete,root/'data/game_data.json')
    current=protocol.Store(root).capture()
    check(json.loads(current.manifest)['external_inputs'][0]['sha256']==protocol.digest(complete.read_bytes()),'Full import publishes a validated frozen external selection')
    master=ROOT.parent/'太空战舰.xlsx'
    cw.sync_workbooks(master,root/'config_excel')
    split=protocol.Store(root).capture()
    check(split.files['data/game_data.json']==current.files['data/game_data.json'],'Split preserves runtime values until explicit import')
    check(split.files['config_excel/published/mon.xlsx']==current.files['config_excel/published/mon.xlsx'],'Split cannot silently replace a running template generation')
    before={name:raw for name,raw in split.files.items()}
    try:full.full_import(master,root/'data/game_data.json')
    except ValueError:pass
    else:raise AssertionError('Historical aggregate unexpectedly passed current validators')
    check(dict(protocol.Store(root).capture().files)==before,'Rejected full import preserves the published configuration')



def generation_protocol():
    import os
    import subprocess
    import config_snapshot as protocol
    root = AREA / 'generation-protocol'
    root.mkdir(exist_ok=True)
    store = protocol.Store(root)
    try: store.capture()
    except FileNotFoundError: pass
    else: raise AssertionError('Reader must not initialize loose files')
    def validate(files):
        for name, payload in files.items():
            if name.endswith('.json'): assert json.loads(payload)['value'] >= 0
    first = store.publish(None, {'data/a.json': b'{"value":1}', 'data/b.json': b'{"value":2}'}, validate)
    check(store.capture().identity == first.identity, 'Whole published generation captures')
    with protocol.exclusive(root / protocol.LOCK):
        result = subprocess.run([sys.executable, '-c',
            'import sys;sys.path.insert(0,sys.argv[1]);import config_snapshot as c;'
            '\ntry:c.Store(sys.argv[2]).capture()\nexcept c.Busy:sys.exit(23)', str(ROOT/'tools'), str(root)], timeout=15)
        check(result.returncode == 23, 'Second process cannot capture during writer lock')
    second = store.publish(first, {'data/a.json': b'{"value":3}', 'data/b.json': b'{"value":4}'}, validate)
    check(first.files['data/a.json'] == b'{"value":1}', 'Captured bytes remain frozen after later publication')
    try: store.publish(first, dict(first.files), validate)
    except protocol.Stale: pass
    else: raise AssertionError('Stale publisher accepted')
    check(store.capture().identity == second.identity, 'Stale refusal preserves generation')
    crash = subprocess.run([sys.executable, '-c',
        'import sys,os;sys.path.insert(0,sys.argv[1]);import config_snapshot as c;'
        's=c.Store(sys.argv[2]);old=s.capture();replace=os.replace;'
        '\ndef crash(a,b):\n replace(a,b)\n if str(b).endswith("a.json"):os._exit(79)'
        '\nc.os.replace=crash;s.publish(old,{"data/a.json":b"{\\\"value\\\":5}","data/b.json":b"{\\\"value\\\":6}"},lambda f:None)',
        str(ROOT/'tools'), str(root)], capture_output=True, timeout=15)
    check(crash.returncode == 79 and (root/protocol.JOURNAL).exists(), 'Actual process death retains installing journal')
    try:store.capture()
    except protocol.RecoveryRequired:pass
    else:raise AssertionError('Ordinary reader accepted interrupted publication')
    store.capture_publication()
    check(store.capture().identity == second.identity and not (root/protocol.JOURNAL).exists(), 'Explicit publisher recovers complete prior generation')
    original_replace = protocol.os.replace
    def fail_both(source, destination):
        if str(destination).endswith('b.json'): raise OSError('injected install/rollback refusal')
        return original_replace(source, destination)
    with patch.object(protocol.os, 'replace', fail_both):
        try: store.publish(second, {'data/a.json': b'{"value":7}', 'data/b.json': b'{"value":8}'}, validate)
        except protocol.RecoveryRequired: pass
        else: raise AssertionError('Rollback refusal was not retained')
    journal = json.loads((root/protocol.JOURNAL).read_bytes())
    check(journal['state']=='recovery_required' and (root/journal['recovery']).is_dir(), 'Failed rollback retains journal and original bytes')
    store.capture_publication()
    check(store.capture().identity==second.identity, 'Next publisher completes retained rollback')
    with protocol.exclusive(root/protocol.LOCK):
        try: protocol.Store(root/'data/..').capture()
        except protocol.Busy: pass
        else: raise AssertionError('Root alias escaped mutual exclusion')
    def outside_lock(files):
        with protocol.exclusive(root/protocol.LOCK): validate(files)
    second = store.publish(second, dict(second.files), outside_lock)
    check(store.capture().identity==second.identity, 'Candidate parsing does not hold the publication lock')
    native_lock_protocol(root, protocol)
    powershell_lock_protocol(root, protocol)
    (root/'data/a.json').write_bytes(b'{"value":9}')
    try: store.capture()
    except protocol.RecoveryRequired: pass
    else: raise AssertionError('Mixed/out-of-protocol generation accepted')
    (root/'data/a.json').write_bytes(second.files['data/a.json'])
    if os.name == 'nt':
        workbook = root/'external.xlsx';workbook.write_bytes(b'frozen workbook bytes')
        with protocol.exclusive(workbook):
            try: protocol.stable_external_bytes(workbook)
            except protocol.Busy: pass
            else: raise AssertionError('Workbook writer handle was not excluded')
        check(protocol.stable_external_bytes(workbook)==b'frozen workbook bytes', 'External workbook captures only after writer releases')
        directory=root/'config_excel';directory.mkdir(exist_ok=True)
        member='config_excel/source.xlsx';source=root/member;source.write_bytes(b'original')
        prior=store.capture()
        published=store.publish(prior,dict(prior.files),validate,inputs={member:b'original'},source_updates={member:b'updated'})
        check(source.read_bytes()==b'updated','Source update writes through exclusive stable handle')
        def fail_manifest(a,b):
            if Path(b)==root/protocol.MANIFEST:raise OSError('injected manifest refusal')
            return original_replace(a,b)
        with patch.object(protocol.os,'replace',fail_manifest):
            try:store.publish(published,dict(published.files),validate,inputs={member:b'updated'},source_updates={member:b'next'})
            except protocol.RecoveryRequired:pass
            else:raise AssertionError('Failed source publication accepted')
        check(source.read_bytes()==b'updated','Rollback restores edited workbook from retained original')
        store.capture_publication()
        check(store.capture().identity==published.identity,'Explicit publisher completes source transaction recovery')

        # Attempt real external writes exactly while rollback holds the verified
        # source handle. Both replacement and in-place edits must be excluded.
        from contextlib import contextmanager
        original_stream=protocol.stable_external_stream
        attempts=[]
        @contextmanager
        def observed_stream(path,*args,**kwargs):
            with original_stream(path,*args,**kwargs) as stream:
                if Path(path)==source and (root/protocol.JOURNAL).exists():
                    code="import sys,os; p=sys.argv[1]; denied=0\nfor action in (lambda:open(p,'wb'),lambda:os.unlink(p)):\n try:action()\n except PermissionError:denied+=1\nsys.exit(0 if denied==2 else 9)"
                    result=subprocess.run([sys.executable,'-c',code,str(source)],capture_output=True,timeout=15)
                    attempts.append(result.returncode)
                yield stream
        fired=False
        def fail_once(a,b):
            nonlocal fired
            if Path(b)==root/protocol.MANIFEST and not fired:
                fired=True;raise OSError('one failed installation')
            return original_replace(a,b)
        with patch.object(protocol,'stable_external_stream',observed_stream),patch.object(protocol.os,'replace',fail_once):
            try:store.publish(published,dict(published.files),validate,inputs={member:b'updated'},source_updates={member:b'rollback'})
            except OSError:pass
            else:raise AssertionError('Injected publication did not fail')
        check(attempts==[0] and source.read_bytes()==b'updated','Recovery holds source identity through verified restore; real external edits/deletes denied')
        fired=False
        with patch.object(protocol.os,'replace',fail_once):
            try:store.publish(published,dict(published.files),validate,inputs={member:b'updated'},source_updates={'config_excel/new.xlsx':b'created'})
            except OSError:pass
            else:raise AssertionError('New source failure did not occur')
        check(not (root/'config_excel/new.xlsx').exists() and not (root/protocol.JOURNAL).exists(),'New source rollback deletes via its still-exclusive verified handle')
        # Both actual CLIs must expose the protocol exception and real recovery target.
        cli_root=AREA/'recovery-cli';cli_root.mkdir()
        shutil.copytree(ROOT/'tools',cli_root/'tools',ignore=shutil.ignore_patterns('__pycache__'))
        (cli_root/'config_excel').mkdir();(cli_root/'data').mkdir()
        journal=cli_root/protocol.JOURNAL
        journal.write_bytes(protocol.encoded({'source_conflict':'config_excel/mon.xlsx'}))
        for tool,args in [('config_workbooks.py',['import']),('level_editor_store.py',['load'])]:
            result=subprocess.run([sys.executable,'-X','utf8',str(cli_root/'tools'/tool),*args],capture_output=True,text=True,encoding='utf-8',timeout=30)
            message=json.loads(result.stdout)
            check(result.returncode==1 and message['recovery_required'] and message['recovery_path']==str(journal) and 'source conflict' in message['message'],'Real CLI retains recovery location and reason: '+tool)



def powershell_lock_protocol(root, protocol):
    import os, subprocess, time, uuid
    if os.name != 'nt':return
    engine=shutil.which('powershell.exe')
    if engine is None:raise AssertionError('PowerShell lock consumer unavailable')
    script=AREA/'lock.ps1'
    script.write_text('param([string]$LockPath,[string]$Marker)\n'
        'try {$lease=[System.IO.File]::Open($LockPath,[System.IO.FileMode]::OpenOrCreate,[System.IO.FileAccess]::ReadWrite,[System.IO.FileShare]::None)} catch {exit 23}\n'
        'try {[System.IO.File]::WriteAllText($Marker,"held"); while (-not [System.IO.File]::Exists($Marker+".release")) {Start-Sleep -Milliseconds 50}} finally {$lease.Dispose()}\n',encoding='utf-8')
    marker=AREA/('powershell-held-'+uuid.uuid4().hex)
    command=[engine,'-NoProfile','-NonInteractive','-ExecutionPolicy','Bypass','-File',str(script),str(root/protocol.LOCK),str(marker)]
    with protocol.exclusive(root/protocol.LOCK):
        with (AREA/'powershell-blocked.log').open('wb') as log:
            child=subprocess.run(command,stdout=log,stderr=subprocess.STDOUT,timeout=30)
        check(child.returncode==23 and not marker.exists(),'PowerShell observes Python configuration lock')
    with (AREA/'powershell-holder.log').open('wb') as log:
        child=subprocess.Popen(command,stdout=log,stderr=subprocess.STDOUT)
        try:
            deadline=time.monotonic()+30
            while not marker.exists() and child.poll() is None and time.monotonic()<deadline:time.sleep(.05)
            check(marker.exists(),'PowerShell holds actual configuration lock')
            try:
                with protocol.exclusive(root/protocol.LOCK):raise AssertionError('Python bypassed PowerShell lock')
            except protocol.Busy:pass
            Path(str(marker)+'.release').touch()
            check(child.wait(timeout=30)==0,'PowerShell releases configuration lock')
        finally:
            if child.poll() is None:child.kill();child.wait()


def build_capture_protocol():
    import os,subprocess,time
    import config_snapshot as protocol
    root=AREA/'build-capture';root.mkdir()
    store=protocol.Store(root)
    first=store.publish(None,{'data/a.json':b'{"value":1}'},lambda _:None)
    captured=store.capture().runtime()
    second=store.publish(first,{'data/a.json':b'{"value":2}'},lambda _:None)
    python_output=AREA/'python-package';captured.materialize(python_output)
    check(protocol.Store(python_output).capture().files['data/a.json']==first.files['data/a.json'],'Python package materializes captured bytes after source generation changes')
    if os.name!='nt':return
    marker=AREA/'ps-captured';destination=AREA/'powershell-package';destination.mkdir()
    script=AREA/'capture-build.ps1'
    script.write_text('param($Helper,$Root,$Marker,$Destination)\n$ErrorActionPreference="Stop"\n. $Helper\n'
        '$captured=Get-ConfigurationSnapshot $Root\n[IO.File]::WriteAllText($Marker,"captured")\n'
        'while(-not [IO.File]::Exists($Marker+".release")){Start-Sleep -Milliseconds 50}\n'
        'foreach($name in $captured.Keys){$path=Join-Path $Destination $name;[IO.Directory]::CreateDirectory([IO.Path]::GetDirectoryName($path))|Out-Null;[IO.File]::WriteAllBytes($path,$captured[$name])}\n',encoding='utf-8')
    with (AREA/'powershell-capture.log').open('wb') as log:
        child=subprocess.Popen(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(script),str(ROOT/'tools/configuration_snapshot.ps1'),str(root),str(marker),str(destination)],stdout=log,stderr=subprocess.STDOUT)
        try:
            deadline=time.monotonic()+30
            while not marker.exists() and child.poll() is None and time.monotonic()<deadline:time.sleep(.05)
            check(marker.exists(),'Actual PowerShell build helper completed capture')
            store.publish(second,{'data/a.json':b'{"value":3}'},lambda _:None)
            Path(str(marker)+'.release').touch()
            check(child.wait(timeout=30)==0,'PowerShell materializes after source changes')
        finally:
            if child.poll() is None:child.kill();child.wait()
    check(protocol.Store(destination).capture().files['data/a.json']==second.files['data/a.json'],'Player staging uses its complete captured generation, not changed source paths')


def native_lock_protocol(root, protocol):
    import os
    import subprocess
    import time
    engine = os.environ.get('SPACE_BATTLESHIP_TEST_GODOT')
    bundle = ROOT/'bin/space_native.dll'
    if not engine or not bundle.is_file():
        print('CONFIG_NATIVE_LOCK: not run (runtime unavailable)', flush=True)
        return
    project = AREA/'native-lock'
    (project/'bin').mkdir(parents=True, exist_ok=True)
    shutil.copy2(bundle, project/'bin/space_native.dll')
    shutil.copy2(ROOT/'space_native.gdextension', project/'space_native.gdextension')
    (project/'project.godot').write_text('config_version=5\n[application]\nconfig/name="Config lock fixture"\n[rendering]\nrenderer/rendering_method="gl_compatibility"\n', encoding='utf-8')
    script = project/'lock.gd'
    script.write_text('extends SceneTree\nvar lease:RefCounted\nvar marker:String\n'
        'func _initialize():\n var args=OS.get_cmdline_user_args()\n lease=ClassDB.instantiate("SpaceSaveLock")\n'
        ' if not lease.acquire(args[0]):print("CONFIG_LOCK_BUSY");quit(23);return\n'
        ' print("CONFIG_LOCK_OWNED")\n'
        ' if args.size()==1:quit(0);return\n marker=args[1]\n FileAccess.open(marker,FileAccess.WRITE).store_string("held")\n'
        'func _process(_delta):\n if not marker.is_empty() and FileAccess.file_exists(marker+".release"):quit(0)\n return false\n', encoding='utf-8')
    env=os.environ.copy()
    for key, folder in [('APPDATA','roaming'),('LOCALAPPDATA','local')]:
        directory=AREA/'native-userdata'/folder;directory.mkdir(parents=True,exist_ok=True)
        env[key]=str(directory)
    # This fixture has no imported assets or global script classes. Register the
    # actual extension directly; runtime loading/instantiation below is required.
    (project/'.godot').mkdir(exist_ok=True)
    (project/'.godot/extension_list.cfg').write_text('res://space_native.gdextension\n',encoding='utf-8')
    cmd=[engine,'--headless','--rendering-method','gl_compatibility','--path',str(project),'--script',str(script),'--',str(root/protocol.LOCK)]
    with protocol.exclusive(root/protocol.LOCK):
        with (project/'blocked.log').open('wb') as log:
            blocked=subprocess.run(cmd,env=env,stdout=log,stderr=subprocess.STDOUT,timeout=45)
        check(blocked.returncode==23 and b'CONFIG_LOCK_BUSY' in (project/'blocked.log').read_bytes(), 'Godot observes Python root lock')
    with (project/'available.log').open('wb') as log:
        available=subprocess.run(cmd,env=env,stdout=log,stderr=subprocess.STDOUT,timeout=45)
    check(available.returncode==0 and b'CONFIG_LOCK_OWNED' in (project/'available.log').read_bytes(), 'Godot acquires after Python releases')
    import uuid
    marker=project/('held-'+uuid.uuid4().hex)
    holder_log=(project/'holder.log').open('wb')
    child=subprocess.Popen(cmd+[str(marker)],env=env,stdout=holder_log,stderr=subprocess.STDOUT)
    try:
        deadline=time.monotonic()+45
        while not marker.exists() and child.poll() is None and time.monotonic()<deadline:time.sleep(.05)
        if not marker.exists():
            child.kill(); child.wait()
        check(marker.exists(), 'Native holder confirms acquisition')
        try: protocol.Store(root).capture()
        except protocol.Busy: pass
        else: raise AssertionError('Python bypassed native root ownership')
        Path(str(marker)+'.release').touch()
        child.wait(timeout=45);holder_log.flush()
        check(child.returncode==0 and b'CONFIG_LOCK_OWNED' in (project/'holder.log').read_bytes(), 'Native release completes')
    finally:
        if child.poll() is None:child.kill();child.wait()
    holder_log.close()
    print('CONFIG_NATIVE_LOCK: passed both directions',flush=True)
    shutil.copy2(ROOT/'scripts/configuration_snapshot.gd',project/'configuration_snapshot.gd')
    store=protocol.Store(project)
    first=store.publish(None,{'data/a.json':b'{"generation":1}'},lambda files:None)
    probe=project/'capture.gd'
    probe.write_text('extends SceneTree\nvar snapshot:RefCounted\nvar marker:String\n'
        'func _initialize():\n snapshot=load("res://configuration_snapshot.gd").new()\n'
        ' if not snapshot.capture():push_error(snapshot.error);quit(4);return\n'
        ' marker=OS.get_cmdline_user_args()[0]\n FileAccess.open(marker,FileAccess.WRITE).store_string(snapshot.identity)\n'
        'func _process(_delta):\n if marker.is_empty() or not FileAccess.file_exists(marker+".release"):return false\n'
        ' var fresh=load("res://configuration_snapshot.gd").new()\n'
        ' if not fresh.capture() or fresh.identity==snapshot.identity or snapshot.json("res://data/a.json").generation!=1 or fresh.json("res://data/a.json").generation!=2:quit(5);return false\n'
        ' print("CONFIG_CAPTURE_STABLE");quit(0);return false\n',encoding='utf-8')
    marker=project/('captured-'+uuid.uuid4().hex)
    command=[engine,'--headless','--rendering-method','gl_compatibility','--path',str(project),'--script',str(probe),'--',str(marker)]
    with (project/'capture.log').open('wb') as log:
        child=subprocess.Popen(command,env=env,stdout=log,stderr=subprocess.STDOUT)
        try:
            deadline=time.monotonic()+45
            while not marker.exists() and child.poll() is None and time.monotonic()<deadline:time.sleep(.05)
            check(marker.exists() and marker.read_text()==first.identity,'Godot captured Python manifest identity')
            store.publish(first,{'data/a.json':b'{"generation":2}'},lambda files:None)
            Path(str(marker)+'.release').touch()
            check(child.wait(timeout=45)==0,'Godot keeps old captured bytes after publisher releases new generation')
        finally:
            if child.poll() is None:child.kill();child.wait()
    check(b'CONFIG_CAPTURE_STABLE' in (project/'capture.log').read_bytes(),'Actual reader parsed both immutable generations')



def main():
    generation_protocol()
    print(f"CONFIGURATION_PROTOCOL {checks} assertions passed",flush=True)
    build_capture_protocol()
    published_editor_protocol()
    print(f"CONFIGURATION_PUBLICATION cumulative {checks} assertions passed",flush=True)
    matrix=[]
    for case in CASES:
        for entry in ('incremental','store_load','store_validate','store_save'):
            matrix.append(run_case(case,entry))
        print('Recorded',case,flush=True)
    projection_and_excel()
    report={'matrix':matrix,'assertions':checks}
    (AREA/'INPUT_MATRIX.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'Configuration boundaries: {len(matrix)} input outcomes, {checks} assertions passed. Evidence: {AREA}',flush=True)


if __name__=='__main__': main()
