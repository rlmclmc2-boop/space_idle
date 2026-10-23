from config_snapshot import RecoveryRequired as ConfigurationRecoveryRequired
"""Level editor source transaction. Uses the existing XLSX projection contract."""
from ui_text import t as ui_text
import argparse
import ast
import copy
from decimal import Decimal, ROUND_HALF_UP
import io
import json
import math
from pathlib import Path
import re
import sys
import uuid
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter
from lxml import etree as ET
from config_workbooks import (Q, MANIFEST, CACHE_VERSION, RecoveryRequired, atomic_batch, sha,
                              sheet_parts, read_changed_file, read_hyperspace_bundle, entity_cache_signature)
from import_workbook import ROOT, SECTIONS, validate_projection, encode

EDITABLE = ('mon', 'monGroup', 'level')
LEVEL_RATIOS = ('atkRatio', 'lifeRatio', 'resRatio')


class PreparedStore:
    def __init__(self, root=ROOT):
        self.root = Path(root)
        self.directory = self.root / 'config_excel'
        self.target = self.root / 'data/game_data.json'
        self.manifest = self.directory / MANIFEST
        mapping = json.loads(self.manifest.read_text(encoding='utf-8'))['sheets']
        self.paths = {}
        for name in SECTIONS:
            if name in ('jewel', 'unlock', 'crew', 'crew_assignment', 'crew_config', 'planet', 'planet_build', 'planet_buff', 'galaxy', 'galaxy_build', 'galaxy_config', 'enhance_config', 'weapon_motion', 'enemy_weapon_base', 'battle_design') and name not in mapping:
                if not (self.directory / f'{name}.xlsx').is_file():
                    continue
                mapping[name] = f'{name}.xlsx'
            filename = mapping[name]
            if Path(filename).name != filename or '/' in filename or '\\' in filename:
                raise ValueError(ui_text('debug.level_editor_store.message_04', name=name))
            self.paths[name] = self.directory / filename

    def snapshot(self):
        paths = [*self.paths.values(), self.target, self.manifest]
        paths.extend(self.directory / name for name in ('hyperspace_config.xlsx', 'hyperspace_enemies.xlsx') if (self.directory / name).exists())
        from hyperspace_entities import load_schema, SCHEMA_PATH, EARLY_FROZEN_PATH
        paths.extend([SCHEMA_PATH, EARLY_FROZEN_PATH])
        paths.extend(self.target.parent / name for name in load_schema()['frozen_inputs'])
        return {str(p): sha(p.read_bytes()) for p in paths}

    def load(self):
        before = self.snapshot()
        tables = {}
        for name in EDITABLE:
            book = openpyxl.load_workbook(self.paths[name], data_only=False)
            try:
                sheet = book[name]
                headers = [c.value for c in sheet[1] if c.value is not None]
                rows = [{key: sheet.cell(r, c + 1).value for c, key in enumerate(headers)}
                        for r in range(4, sheet.max_row + 1) if sheet.cell(r, 1).value is not None]
                numeric = [key for c,key in enumerate(headers) if str(sheet.cell(3,c+1).value).lower() in ('int','float','number') or any(type(row.get(key)) in (int,float) for row in rows)]
                tables[name] = {'headers': headers, 'rows': rows, 'numeric_fields': numeric}
            finally:
                book.close()
        if before != self.snapshot():
            raise ValueError(ui_text('debug.level_editor_store.message_01'))
        data = json.loads(self.target.read_text(encoding='utf-8'))
        return {'tables': tables, 'snapshot': before, 'equipment': data['equipment'], 'resources': data['resources']}

    def workbook_bytes(self, name, table):
        """Patch sheet XML only; preserve styles, headers and other package parts.

        Formula addresses stay explicit in the editor. Evaluate the project's
        arithmetic/ROUND subset and refresh caches; unsupported syntax fails closed.
        """
        headers, rows = table['headers'], table['rows']
        cells = {f'{get_column_letter(c+1)}{r+4}': row.get(key)
                 for r, row in enumerate(rows) for c, key in enumerate(headers)}
        memo = {}

        def value(address, visiting=frozenset()):
            if address in memo:
                return memo[address]
            if address in visiting:
                raise ValueError(ui_text('debug.level_editor_store.message_05', name=name, address=address))
            raw = cells.get(address)
            if not isinstance(raw, str) or not raw.startswith('='):
                return raw
            expression = re.sub(r'\$?([A-Z]+)\$?(\d+)', lambda m: f'CELL("{m[1]}{m[2]}")', raw[1:])

            def walk(node):
                if isinstance(node, ast.Constant) and type(node.value) in (str, int, float):
                    return Decimal(str(node.value)) if type(node.value) in (int, float) else node.value
                if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                    return walk(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
                if isinstance(node, ast.BinOp):
                    a, b = walk(node.left), walk(node.right)
                    if isinstance(node.op, ast.Add): return a + b
                    if isinstance(node.op, ast.Sub): return a - b
                    if isinstance(node.op, ast.Mult): return a * b
                    if isinstance(node.op, ast.Div): return a / b
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and not node.keywords:
                    args = [walk(a) for a in node.args]
                    if node.func.id == 'CELL' and len(args) == 1:
                        result = value(args[0], visiting | {address})
                        if type(result) not in (int, float): raise ValueError(ui_text('debug.level_editor_store.message_12', args=str(args[0])))
                        return Decimal(str(result))
                    if node.func.id.upper() == 'ROUND' and len(args) == 2:
                        return args[0].quantize(Decimal(1).scaleb(-int(args[1])), rounding=ROUND_HALF_UP)
                raise ValueError(ui_text('debug.level_editor_store.message_06'))
            try:
                result = float(walk(ast.parse(expression, mode='eval').body))
                if not math.isfinite(result): raise ValueError(ui_text('debug.level_editor_store.message_11'))
            except Exception as error:
                raise ValueError(ui_text('debug.level_editor_store.message_101', name=name, address=address, raw=raw, error=error)) from error
            memo[address] = result
            return result

        with zipfile.ZipFile(self.paths[name]) as archive:
            part = dict(sheet_parts(archive))[name]
            tree = ET.fromstring(archive.read(part))
            body = tree.find(Q + 'sheetData')
            preserved, existing_ids = {}, set()
            if name == 'level':
                book = openpyxl.load_workbook(io.BytesIO(self.paths[name].read_bytes()), data_only=True)
                try:
                    sheet = book[name]
                    for old_row in body:
                        old_index = int(old_row.get('r'))
                        if old_index < 4: continue
                        record_id = sheet.cell(old_index, 1).value
                        existing_ids.add(record_id)
                        for old_cell in old_row:
                            column = re.sub(r'\d', '', old_cell.get('r'))
                            for key in LEVEL_RATIOS:
                                if key in headers and column == get_column_letter(headers.index(key)+1):
                                    preserved[(record_id, key)] = copy.deepcopy(old_cell)
                finally:
                    book.close()
            styles, exact_styles, row_attributes = {}, {}, {}
            for row in list(body):
                if int(row.get('r')) >= 4:
                    row_attributes[int(row.get('r'))] = dict(row.attrib)
                    for cell in row:
                        styles.setdefault(re.sub(r'\d', '', cell.get('r')), cell.get('s'))
                        exact_styles[cell.get('r')] = cell.get('s')
                    body.remove(row)
            for r, record in enumerate(rows, 4):
                row = ET.SubElement(body, Q + 'row', **row_attributes.get(r, {'r': str(r)}))
                for c, key in enumerate(headers, 1):
                    letter = get_column_letter(c)
                    address = f'{letter}{r}'
                    if name == 'level' and key in LEVEL_RATIOS:
                        original = preserved.get((record.get('id'), key))
                        if original is not None:
                            cell = copy.deepcopy(original)
                            cell.set('r', address)
                            row.append(cell)
                        elif record.get('id') not in existing_ids:
                            cell = ET.SubElement(row, Q + 'c', r=address)
                            ET.SubElement(cell, Q + 'v').text = '1'
                        continue
                    raw = record.get(key)
                    if raw is None: continue
                    cell = ET.SubElement(row, Q + 'c', r=address)
                    style = exact_styles.get(address, styles.get(letter))
                    if style: cell.set('s', style)
                    if isinstance(raw, str) and raw.startswith('='):
                        ET.SubElement(cell, Q + 'f').text = raw[1:]
                        ET.SubElement(cell, Q + 'v').text = str(value(address))
                    elif type(raw) in (int, float):
                        if not math.isfinite(raw): raise ValueError(ui_text('debug.level_editor_store.message_13', name=name, address=address))
                        ET.SubElement(cell, Q + 'v').text = str(int(raw)) if raw == int(raw) else str(raw)
                    else:
                        cell.set('t', 'inlineStr')
                        text = ET.SubElement(ET.SubElement(cell, Q + 'is'), Q + 't')
                        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                        text.text = str(raw)
            dimension = tree.find(Q + 'dimension')
            if dimension is None:
                dimension = ET.Element(Q + 'dimension')
                tree.insert(tree.index(body), dimension)
            dimension.set('ref', f'A1:{get_column_letter(len(headers))}{len(rows)+3}')
            output = io.BytesIO()
            with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as result:
                for item in archive.infolist():
                    result.writestr(item, ET.tostring(tree) if item.filename == part else archive.read(item.filename))
            return output.getvalue()

    def prepare(self, request):
        current = self.load()
        if current['snapshot'] != request['snapshot']:
            raise ValueError(ui_text('debug.level_editor_store.message_02'))
        tables = copy.deepcopy(request['tables'])
        original_levels = {r['id']: r for r in current['tables']['level']['rows']}
        for row in tables['level']['rows']:
            for key in LEVEL_RATIOS:
                row[key] = original_levels.get(row.get('id'), {}).get(key, 1)
        files = {}
        for name in EDITABLE:
            if tables[name]['headers'] != current['tables'][name]['headers']:
                raise ValueError(ui_text('debug.level_editor_store.message_07', name=name))
            ids = [r.get('id') for r in tables[name]['rows']]
            if any(type(i) not in (int, float) or not math.isfinite(i) or i < 1 or i != int(i) for i in ids) or len(set(ids)) != len(ids):
                raise ValueError(ui_text('debug.level_editor_store.message_08', name=name))
            if tables[name] != current['tables'][name]:
                files[self.paths[name]] = self.workbook_bytes(name, tables[name])
        data = json.loads(self.target.read_text(encoding='utf-8'))
        raw_files = {n: files.get(p, p.read_bytes()) for n, p in self.paths.items()}
        for name, raw in raw_files.items():
            ignored = [get_column_letter(tables['level']['headers'].index(k)+1) for k in LEVEL_RATIOS if k in tables['level']['headers']] if name == 'level' else []
            data[SECTIONS[name]] = read_changed_file(self.paths[name], name, raw, ignored_formula_columns=ignored)
        validate_projection(data, check_level_ratios=False)
        warnings = []
        data['source_files'] = {n: str(p.resolve()) for n, p in self.paths.items()}
        payload = encode(data)
        state = {'version': CACHE_VERSION, 'directory': str(self.directory.resolve()),
                 'target_hash': sha(payload), 'hashes': {n: sha(raw) for n, raw in raw_files.items()}}
        entity_paths = [self.directory / name for name in ('hyperspace_config.xlsx', 'hyperspace_enemies.xlsx')]
        if any(p.exists() for p in entity_paths):
            raw_entities = {p.name: p.read_bytes() for p in entity_paths}
            projected, _ = read_hyperspace_bundle(raw_entities, data, self.target.parent)
            entity_payloads = {name: (json.dumps(values, ensure_ascii=False, indent=2) + '\n').encode('utf-8') for name, values in projected.items()}
            files.update({self.target.with_name(name): raw for name, raw in entity_payloads.items()})
            state['entity_cache'] = entity_cache_signature(raw_entities, self.target.parent, entity_payloads)
            state['hyperspace_hashes'] = {name: sha(raw) for name, raw in raw_entities.items()}
        files[self.target] = payload
        files[self.target.with_name('.import_state.json')] = encode(state)
        return files, warnings

    def execute(self, request, save=False):
        files, warnings = self.prepare(request)
        if save:
            if request['snapshot'] != self.snapshot(): raise ValueError(ui_text('debug.level_editor_store.message_09'))
            backup = self.root / '.runtime/level-editor-backups' / uuid.uuid4().hex
            backups = {backup / p.relative_to(self.root): p.read_bytes() for p in files if p.exists()}
            backup_manifest = {'version': 1, 'files': [
                {'target': str((self.root / saved.relative_to(backup)).resolve()),
                 'original': str(saved.relative_to(backup)), 'sha256': sha(raw)}
                for saved, raw in backups.items()
            ]}
            atomic_batch({**backups, backup / 'manifest.json':
                          (json.dumps(backup_manifest, indent=2) + '\n').encode('utf-8')},
                         self.root / '.runtime/config-recovery')
            atomic_batch(files, self.root / '.runtime/config-recovery')
            return {**self.load(), 'warnings': warnings, 'backup': str(backup), 'message': ui_text('debug.level_editor_store.message_10')}
        return {'warnings': warnings, 'message': ui_text('debug.level_editor_store.message_03')}


class Store:
    """Public editor reads and writes one stable configuration/input capture."""
    def __init__(self,root=ROOT):self.root=Path(root).resolve()
    def load(self):
        from config_publication import execute
        return execute(self.root,'load')
    def execute(self,request,save=False):
        from config_publication import execute
        return execute(self.root,'save' if save else 'validate',request)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['load', 'validate', 'save'])
    parser.add_argument('--request', type=Path)
    args = parser.parse_args()
    try:
        store = Store()
        result = store.load() if args.action == 'load' else store.execute(json.loads(args.request.read_text(encoding='utf-8')), args.action == 'save')
        print(json.dumps({'ok': True, **result}, ensure_ascii=True, allow_nan=False))
    except Exception as error:
        result = {'ok': False, 'message': str(error), 'recovery_required': isinstance(error, (RecoveryRequired, ConfigurationRecoveryRequired))}
        if isinstance(error, (RecoveryRequired, ConfigurationRecoveryRequired)):
            result.update(recovery_path=(getattr(error,"recovery_path",None) or str(ROOT/".configuration-transaction.json")), rollback_failed=getattr(error,"rollback_failed",[]))
        print(json.dumps(result, ensure_ascii=True))
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
