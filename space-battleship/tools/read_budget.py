#!/usr/bin/env python3
"""Bounded, resumable project reads. Python 3; Excel subcommand needs openpyxl."""
import argparse
from collections import deque
import fnmatch
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CHARS = 1800
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")


def emit(rows, limit):
    """Never truncate a result: materialize overflow and return its locator."""
    lines = [str(row).rstrip("\n") for row in rows]
    body = "\n".join(lines) + ("\n" if lines else "")
    if len(body) <= limit:
        sys.stdout.write(body)
        return
    with tempfile.NamedTemporaryFile(
        mode="w", encoding="utf-8", newline="\n", suffix=".txt",
        prefix="space-idle-read-", delete=False,
    ) as saved:
        saved.write(body)
    print(f"OVER LIMIT: {len(lines)} lines, {len(body)} chars; saved: {saved.name}")
    print(f"Continue: python tools/read_budget.py slice --path \"{saved.name}\" --start 1 --count 20")


def git(*args):
    done = subprocess.run(["git", "-C", str(ROOT), *args], capture_output=True, text=True, encoding="utf-8", errors="replace")
    if done.returncode:
        raise RuntimeError(done.stderr.strip() or f"git exit {done.returncode}")
    return done.stdout


def tracked():
    return [p for p in git("ls-files", "-z").split("\0") if p]


def selected_paths(names):
    if not names or len(names) > 16:
        raise ValueError("give 1–16 explicit --path entries; locate filenames first")
    for name in names:
        path = Path(name).resolve()
        if not path.is_file():
            raise ValueError(f"missing file: {name}")
        yield path


def cmd_files(a):
    if not a.glob or a.glob in ("*", "**", "**/*"):
        raise ValueError("use a selective --glob, not the whole repository")
    rows = []
    for name in tracked():
        if fnmatch.fnmatch(name, a.glob) or fnmatch.fnmatch(Path(name).name, a.glob):
            path = ROOT / name
            rows.append(f"{path.stat().st_size:>10}  {name}")
    emit([f"MATCHES {len(rows)}", *rows], a.chars)


def cmd_grep(a):
    pattern = re.compile(a.query if a.regex else re.escape(a.query), re.IGNORECASE if a.ignore_case else 0)
    rows = []
    for path in selected_paths(a.path):
        with path.open(encoding="utf-8", errors="replace") as stream:
            for number, line in enumerate(stream, 1):
                if pattern.search(line):
                    rows.append(f"{path.relative_to(ROOT) if path.is_relative_to(ROOT) else path}:{number}: {line.rstrip()}")
    emit([f"MATCHES {len(rows)}", *rows], a.chars)


def cmd_slice(a):
    path = next(selected_paths([a.path]))
    if a.start < 1 or not 1 <= a.count <= 80:
        raise ValueError("start >= 1 and count 1–80 required")
    rows = []
    with path.open(encoding="utf-8", errors="replace") as stream:
        for number, line in enumerate(stream, 1):
            if number >= a.start:
                rows.append(f"{number}: {line.rstrip()}")
            if number >= a.start + a.count - 1:
                break
    emit([f"{path} lines {a.start}–{a.start + len(rows) - 1}", *rows], a.chars)


def cmd_json(a):
    path = next(selected_paths([a.path]))
    with path.open(encoding="utf-8") as stream:
        value = json.load(stream)
    if not (a.key or a.keys or a.where):
        raise ValueError("--key or --where required for value extraction; use --keys for index")
    # Slash separates levels when a literal key contains dots (such as UI text IDs).
    for part in (a.key.split("/") if "/" in a.key else a.key.split(".")) if a.key else []:
        value = value[int(part)] if isinstance(value, list) else value[part]
    if a.keys:
        if isinstance(value, dict):
            emit([f"OBJECT KEYS {len(value)}", *value.keys()], a.chars)
        elif isinstance(value, list):
            first = value[0] if value else None
            emit([f"ARRAY ITEMS {len(value)}", f"FIRST TYPE {type(first).__name__}",
                  *(first.keys() if isinstance(first, dict) else [])], a.chars)
        else:
            raise ValueError("--keys requires an object or array")
        return
    if a.where:
        field, expected = a.where.split("=", 1)
        if not isinstance(value, list):
            raise ValueError("--where requires a list at --key")
        value = [row for row in value if isinstance(row, dict) and str(row.get(field)) == expected]
    if a.fields:
        fields = a.fields.split(",")
        if isinstance(value, list):
            value = [{field: row.get(field) for field in fields} for row in value]
        elif isinstance(value, dict):
            value = {field: value.get(field) for field in fields}
    emit([json.dumps(value, ensure_ascii=False, indent=2)], a.chars)


def cmd_xlsx(a):
    from openpyxl import load_workbook
    path = next(selected_paths([a.path]))
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not a.sheet:
            emit([f"SHEETS {len(workbook.sheetnames)}", *workbook.sheetnames], a.chars)
            return
        sheet = workbook[a.sheet]
        if a.row < 1 or not 1 <= a.count <= 50:
            raise ValueError("row >= 1 and count 1–50 required")
        rows = [f"{path} [{a.sheet}] {'find ' + a.find if a.find else f'rows {a.row}–{a.row + a.count - 1}'}"]
        scan = sheet.iter_rows(min_row=a.row, max_row=None if a.find else a.row + a.count - 1, values_only=True)
        matches = 0
        for number, cells in enumerate(scan, a.row):
            if a.find and not any(a.find.casefold() in str(cells[col - 1]).casefold() for col in a.columns if col <= len(cells) and cells[col - 1] is not None):
                continue
            matches += 1
            entries = [f"{col}:{str(cells[col - 1])}" for col in a.columns if col <= len(cells) and cells[col - 1] is not None]
            rows.append(f"{number}  {' | '.join(entries)}")
        if a.find:
            rows.insert(1, f"MATCHES {matches}")
        emit(rows, a.chars)
    finally:
        workbook.close()


def cmd_diff(a):
    if not a.path:
        raise ValueError("give exact --path entries for diff")
    rows = git("diff", "--no-ext-diff", "--unified=3", a.base, "--", *a.path).splitlines()
    emit([f"DIFF {a.base}: {len(rows)} lines", *rows], a.chars)


def cmd_log(a):
    path = next(selected_paths([a.path]))
    pattern = re.compile(a.pattern, re.IGNORECASE)
    if not 0 <= a.context <= 20 or not 1 <= a.tail <= 100:
        raise ValueError("context 0–20 and tail 1–100 required")
    prior = deque(maxlen=a.context)
    tail = deque(maxlen=a.tail)
    selected = {}
    following = hits = total = 0
    with path.open(encoding="utf-8", errors="replace") as stream:
        for total, raw in enumerate(stream, 1):
            line = raw.rstrip("\n")
            tail.append((total, line))
            if pattern.search(line):
                hits += 1
                selected.update(prior)
                selected[total] = line
                following = a.context
            elif following:
                selected[total] = line
                following -= 1
            prior.append((total, line))
    rows = [f"LOG {path}: {total} lines; {hits} matching lines"]
    rows.extend(f"{n}: {line}" for n, line in (sorted(selected.items()) if hits else tail))
    emit(rows, a.chars)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="action", required=True)
    for name in ("files", "grep", "slice", "json", "xlsx", "diff", "log"):
        q = sub.add_parser(name)
        q.add_argument("--chars", type=int, default=DEFAULT_CHARS, help="max stdout chars; overflow saved with locator")
        if name != "files":
            q.add_argument("--path", action="append", required=True if name not in ("diff",) else False)
        if name == "files": q.add_argument("--glob", required=True)
        if name == "grep":
            q.add_argument("--query", required=True)
            q.add_argument("--regex", action="store_true")
            q.add_argument("--ignore-case", action="store_true")
        if name == "slice":
            q.add_argument("--start", type=int, required=True)
            q.add_argument("--count", type=int, default=20)
        if name == "json":
            q.add_argument("--key", default="")
            q.add_argument("--keys", action="store_true")
            q.add_argument("--where")
            q.add_argument("--fields")
        if name == "xlsx":
            q.add_argument("--sheet")
            q.add_argument("--find", help="case-insensitive value fragment in selected columns")
            q.add_argument("--row", type=int, default=1)
            q.add_argument("--count", type=int, default=10)
            q.add_argument("--columns", type=int, nargs="+", default=[1, 2, 3])
        if name == "diff": q.add_argument("--base", default="HEAD")
        if name == "log":
            q.add_argument("--pattern", default=r"error|fail|exception|traceback")
            q.add_argument("--context", type=int, default=2)
            q.add_argument("--tail", type=int, default=20)
    a = p.parse_args()
    if a.chars < 400 or a.chars > 10000:
        p.error("--chars must be 400–10000")
    if a.action in ("slice", "json", "xlsx", "log"):
        a.path = a.path[0] if len(a.path) == 1 else p.error("give one --path")
    try:
        globals()[f"cmd_{a.action}"](a)
    except (ValueError, KeyError, ImportError, RuntimeError, OSError, subprocess.SubprocessError) as exc:
        print(f"ERROR {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
