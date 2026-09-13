"""Read-only source audit or bounded Excel inspection; never replaces game data."""
import argparse
import ast
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_label(path):
    try:
        return Path(os.path.relpath(path.resolve(), ROOT)).as_posix()
    except ValueError:  # A source on another Windows drive has no relative path.
        return path.resolve().as_posix()


def differences(a, b, path=''):
    if isinstance(a, dict) and isinstance(b, dict):
        return [d for k in sorted(a.keys() | b.keys())
                for d in differences(a.get(k), b.get(k), path + '/' + k)]
    if isinstance(a, list) and isinstance(b, list) and len(a) == len(b):
        return [d for i, (x, y) in enumerate(zip(a, b))
                for d in differences(x, y, path + '/' + str(i))]
    return [] if a == b else [{'path': path, 'imported': a, 'runtime': b}]


def audit(source):
    paths = [source, ROOT/'data/game_data.json', ROOT/'project.godot', ROOT/'main.tscn']
    paths += sorted(ROOT.glob('*.cmd'))
    paths += [ROOT.parent/'启动.cmd']
    paths += sorted((ROOT/'scripts').glob('*.gd'))
    paths += sorted((ROOT.parent/'test').glob('test_*'))
    paths += [ROOT/'tools/import_workbook.py', Path(__file__).resolve()]
    paths = [p for p in paths if p.suffix != '.uid']
    before = {p: digest(p) for p in paths}
    formulas = openpyxl.load_workbook(source, data_only=False)
    cached = openpyxl.load_workbook(source, data_only=True)
    issues, unsupported = [], []
    memo = {}

    def cell_value(sheet, address, visiting=frozenset()):
        key = (sheet, address)
        if key in memo:
            return memo[key]
        if key in visiting:
            raise ValueError('circular reference')
        value = formulas[sheet][address].value
        if not isinstance(value, str) or not value.startswith('='):
            return value
        expression = re.sub(r'\b\$?([A-Z]+)\$?(\d+)\b',
                            lambda m: 'CELL("' + m[1] + m[2] + '")', value[1:])

        def walk(node):
            if isinstance(node, ast.Constant):
                return Decimal(str(node.value)) if isinstance(node.value, (int, float)) else node.value
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
                args = [walk(a) for a in node.args]
                if node.func.id == 'CELL' and len(args) == 1:
                    v = cell_value(sheet, args[0], visiting | {key})
                    return Decimal(str(v)) if isinstance(v, (int, float)) else v
                if node.func.id == 'ROUND' and len(args) == 2:
                    return Decimal(str(args[0])).quantize(Decimal(1).scaleb(-int(args[1])), rounding=ROUND_HALF_UP)
            if isinstance(node, ast.BinOp):
                a, b = walk(node.left), walk(node.right)
                if isinstance(node.op, ast.Mult): return a * b
                if isinstance(node.op, ast.Add): return a + b
                if isinstance(node.op, ast.Sub): return a - b
                if isinstance(node.op, ast.Div): return a / b
            raise ValueError('unsupported formula syntax')

        result = walk(ast.parse(expression, mode='eval').body)
        memo[key] = result
        return result

    sheets = []
    for sheet in formulas:
        count = 0
        for row in sheet:
            for cell in row:
                if cell.data_type == 'e':
                    issues.append({'cell': sheet.title+'!'+cell.coordinate, 'error': cell.value})
                if cell.data_type != 'f':
                    continue
                count += 1
                actual = cached[sheet.title][cell.coordinate].value
                try:
                    expected = cell_value(sheet.title, cell.coordinate)
                    matches = (Decimal(str(actual)) == expected if isinstance(expected, Decimal) and isinstance(actual, (int, float)) else actual == expected)
                    if not matches:
                        issues.append({'cell': sheet.title+'!'+cell.coordinate, 'formula': cell.value,
                                       'cached': actual, 'calculated': float(expected) if isinstance(expected, Decimal) else expected})
                except (ValueError, TypeError, ArithmeticError, SyntaxError) as exc:
                    unsupported.append({'cell': sheet.title+'!'+cell.coordinate, 'reason': str(exc)})
        sheets.append({'name': sheet.title, 'range': sheet.calculate_dimension(), 'formulas': count})
    formulas.close()
    cached.close()
    area = ROOT.parent / 'test/work'
    area.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='knowledge-audit-', dir=area) as folder:
        target = Path(folder)/'game_data.json'
        target.write_bytes((ROOT/'data/game_data.json').read_bytes())
        result = subprocess.run([sys.executable, str(ROOT/'tools/import_workbook.py'), str(source), str(target)], capture_output=True)
        comparison = differences(json.loads(target.read_text(encoding='utf-8')), json.loads((ROOT/'data/game_data.json').read_text(encoding='utf-8'))) if result.returncode == 0 else None
    report = {
        'checked_at_utc': datetime.now(timezone.utc).isoformat(),
        'scope': 'Saved workbook, cached formulas, temporary import versus runtime JSON; no source writes.',
        'sources': [{'path': source_label(p), 'sha256': h} for p, h in before.items()],
        'changed_during_audit': [str(p) for p, h in before.items() if not p.exists() or digest(p) != h],
        'sheets': sheets, 'formula_cache_issues': issues, 'unsupported_formulas': unsupported,
        'import_exit_code': result.returncode,
        'import_summary': result.stdout.decode('utf-8', errors='replace').strip(),
        'import_error': result.stderr.decode('utf-8', errors='replace').strip(),
        'runtime_differences': comparison,
    }
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=ROOT.parent/'太空战舰.xlsx')
    parser.add_argument('--sheet')
    parser.add_argument('--range', dest='area')
    parser.add_argument('--output', type=Path, help='Explicit report destination; omit to print')
    args = parser.parse_args()
    if args.sheet:
        if not args.area:
            parser.error('--sheet requires --range for minimum reading')
        book = openpyxl.load_workbook(args.source, data_only=False)
        values = openpyxl.load_workbook(args.source, data_only=True)
        report = [{'cell': c.coordinate, 'value': values[args.sheet][c.coordinate].value,
                   **({'formula': c.value} if c.data_type == 'f' else {})}
                  for row in book[args.sheet][args.area] for c in row if c.value is not None]
        book.close()
        values.close()
    else:
        report = audit(args.source)
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        if args.output.resolve() in [args.source.resolve(), (ROOT/'data/game_data.json').resolve()] or args.output.suffix != '.json':
            parser.error('output must be a separate .json report')
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload+'\n', encoding='utf-8')
        print('Report:', args.output)
    else:
        print(payload)
    if not args.sheet and (report['changed_during_audit'] or report['formula_cache_issues'] or report['unsupported_formulas'] or report['import_exit_code'] or report['runtime_differences']):
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
