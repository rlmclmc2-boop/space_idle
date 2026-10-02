"""Excel projection and validation shared by full and incremental imports."""
from ui_text import t as ui_text
import argparse
import ast
import copy
import json
import math
import os
import pathlib
import re
import sys
import openpyxl
from galaxy_config import validate as validate_galaxy

ROOT = pathlib.Path(__file__).resolve().parents[1]
SECTIONS = {"level":"levels", "equipment":"equipment", "mon":"enemies", "monGroup":"groups", "res":"resources", "config":"config", "ship":"ship", "hightech":"hightech", "unlock":"unlock", "crew":"crew", "crew_assignment":"crew_assignment", "crew_config":"crew_config", "planet":"planet", "planet_build":"planet_build", "planet_buff":"planet_buff", "enhance_config":"enhance_config"}
DEFAULTS = {"autoCollectDelay":5.0,"loopDelay":6.0,"deathRetreatDistance":300.0,"deathRetreatDuration":1.2,"projectilePixelsPerUnit":28.0,"startingIron":0.0,"startingTitanium":0.0}
SECTIONS.update({name:name for name in ('galaxy','galaxy_build','galaxy_config','weapon_motion','enemy_weapon_base')})
FALLBACKS = {"enemyWeaponMissingLevel":"Use player weapon row 1 when the enemy weapon row is missing.","enemyCannonMissingDamageAndCooldown":"Use player cannon row 1 for missing fields."}

def clean(value):
    return str(value).translate(str.maketrans({"｛":"{","｝":"}","，":",","；":";"})).strip("{} ")

def read_rows(sheet):
    values=sheet.iter_rows(values_only=True)
    header=next(values)
    # Rows 2–3 are human column descriptions and types, never configuration data.
    next(values,None)
    next(values,None)
    return [{key:value for key,value in zip(header,row) if key is not None} for row in values if row and row[0] is not None]

def convert_sheet(name, rows):
    if name in ('weapon_motion', 'enemy_weapon_base'):
        result = {}
        for row in rows:
            key = row.get('id')
            if not isinstance(key,str) or not key.strip() or key in result:
                raise ValueError(f'{name}: missing or duplicate ID {key}')
            result[key] = row
        return result
    if name in ('galaxy','galaxy_build','galaxy_config'):
        result = {}
        for source in rows:
            row = {key:('' if value is None else value) for key,value in source.items()}
            key = row.get('key')
            if not isinstance(key,str) or not key.strip() or key in result:
                raise ValueError(f'{name}: missing or duplicate key {key}')
            result[key] = row
        return result
    if name in ('crew', 'crew_assignment', 'crew_config', 'planet', 'planet_build', 'planet_buff', 'enhance_config'):
        result = {}
        for source in rows:
            row = {k: ('' if v is None else v) for k, v in source.items()}
            if name == 'crew':
                for retired in ('baseLevel', 'maxLevel', 'expGroup', 'baseExp', 'expGrowth'):
                    row.pop(retired, None)
            key = row.get('id')
            if name in ('planet', 'planet_buff'):
                if type(key) not in (int,float) or not math.isfinite(key) or key < 1 or key != int(key):
                    raise ValueError(f'{name}: expected positive numeric ID')
                row['id'] = int(key)
                key = str(int(key))
            if not isinstance(key, str) or not key.strip():
                raise ValueError(f'{name}: missing string ID')
            if key in result:
                raise ValueError(f'{name}: duplicate {key}')
            result[key] = row
        return result
    if name in ('mon', 'monGroup', 'res', 'level'):
        ids = [row.get('id') for row in rows]
        if any(type(i) not in (int, float) or not math.isfinite(i) or i < 1 or i != int(i) for i in ids) or len(set(ids)) != len(ids):
            raise ValueError(ui_text('debug.import_workbook.message_02', name=name))
    if name == 'config':
        rows=[row for row in rows if row.get('name') not in ('jewelCreat','jewelCombine','equipmentSocket','baseCriDmg')]
        keys = [row.get('name') for row in rows]
        if any(not isinstance(key, str) or not key.strip() for key in keys) or len(set(keys)) != len(keys):
            raise ValueError(ui_text('debug.import_workbook.message_03'))
    if name in ("hightech", "unlock"):
        result={}
        for row in rows:
            key=row.get("name")
            if not isinstance(key,str) or not key.strip() or key in result:
                raise ValueError(ui_text('debug.import_workbook.message_08', name=name))
            result[key]=row
        return result
    if name=="equipment":
        result={}
        for row in rows:
            if row.get('level') != 1:
                continue  # Legacy higher-level rows are no longer configuration.
            result.setdefault(row["name"],[]).append(row)
        for items in result.values():
            if any(type(r.get('level')) not in (int, float) or not math.isfinite(r['level']) or r['level'] < 1 or r['level'] != int(r['level']) for r in items):
                raise ValueError(ui_text('debug.import_workbook.message_09'))
            if len({r['level'] for r in items}) != len(items):
                raise ValueError(ui_text('debug.import_workbook.message_10'))
            items.sort(key=lambda r:r["level"])
        return result
    if name=="mon":
        result={}
        for row in rows:
            mounts = []
            for part in clean(row["equipment"]).split(","):
                fields = part.split("|")
                if len(fields) != 2 or not fields[0].strip() or not fields[1].strip().isdigit() or int(fields[1]) < 1:
                    raise ValueError(ui_text('debug.import_workbook.message_16', id=row["id"]))
                mounts.extend({"name": fields[0].strip()} for _ in range(int(fields[1])))
            row["equipment"] = mounts
            drop=clean(row["res"]).split(",")
            if len(drop)==4:
                drop=[drop[0],drop[1],drop[2]+"."+drop[3]]
            row["drops"]=[{"resourceId":int(drop[0]),"amount":float(drop[1]),"chance":float(drop[2])}]
            result[str(row["id"])]=row
        return result
    if name=="monGroup":
        return {str(r["id"]):{"description":r["des"],"slots":[None if v.strip()=="null" else int(v) for v in clean(r["mon"]).split(",")]} for r in rows}
    if name=="level":
        for row in rows:
            ratio = row.get("jewelRatio", 1)
            if type(ratio) not in (int, float) or not math.isfinite(ratio) or ratio < 0:
                raise ValueError(ui_text('debug.import_workbook.message_11'))
            row["jewelRatio"] = ratio
            row["groups"]=[{"id":int(p.split("|")[0]),"position":float(p.split("|")[1])} for p in clean(row["monGroup"]).split(",")]
        return rows
    if name=="res":
        return {str(r["id"]):r["name"] for r in rows}
    if name=="config":
        return {r["name"]:r["para_1"] for r in rows}
    if name=="ship":
        result={}
        for row in rows:
            key=row.get("name")
            if not isinstance(key,str) or not key.strip() or key in result:
                raise ValueError(ui_text('debug.import_workbook.message_12'))
            result[key]={"name":key,"des":row.get("des"),"weaponSlots":row.get("para_1"),"defenseSlots":row.get("para_2"),"movement":row.get("para_3"),"unlock":row.get("para_4"),"size":row.get("para_5"),"sameEquipmentLimit":row.get("para_6")}
        return result
    raise ValueError(ui_text('debug.import_workbook.message_01', name=name))

def projection_base(previous, source):
    data=copy.deepcopy(previous or {})
    data.pop('jewel', None)
    data.get('source_files', {}).pop('jewel', None)
    for retired in ('jewelCreat','jewelCombine','equipmentSocket','baseCriDmg'):data.get('config',{}).pop(retired,None)
    data.pop('crew_level', None)
    data.get('source_files', {}).pop('crew_level', None)
    for row in data.get('crew', {}).values():
        for retired in ('baseLevel', 'maxLevel', 'expGroup', 'baseExp', 'expGrowth'):
            row.pop(retired, None)
    data.pop('charge', None)
    data.get('source_files', {}).pop('charge', None)
    data["source"]=source
    data["defaults"]={**DEFAULTS,**data.get("defaults",{})}
    data.setdefault("fallbacks",dict(FALLBACKS))
    return data

def positive(value, label, allow_zero=False):
    if not isinstance(value, (int,float)) or not math.isfinite(value) or value < 0 or (value == 0 and not allow_zero):
        raise ValueError(ui_text('debug.import_workbook.message_101', label=label, positive="nonnegative" if allow_zero else "positive", value=repr(value)))

def validate_description(row, field='description', section='hightech'):
    description=row.get(field)
    label=f'{section} {row["name"]} {field}'
    if not isinstance(description,str) or not description.strip():
        raise ValueError(ui_text('debug.import_workbook.message_102', label=label))
    for token in re.findall(r'para\d+', description):
        positive(row.get(token if section=='hightech' else token.replace('para','para_')),f'{label} {token}',True)
    blocks=re.findall(r'\{([^{}]+)\}',description)
    if '{' in re.sub(r'\{[^{}]+\}','',description) or '}' in re.sub(r'\{[^{}]+\}','',description):
        raise ValueError(ui_text('debug.import_workbook.message_103', label=label))
    for block in blocks:
        parts=block.replace('，',',').split(',')
        if any(option.strip() not in ('向上取整','不含自身','百分比显示','保留两位小数','即100.3%展示为100%','百分比','四舍五入保留整数百分比部分') for option in parts[1:]):
            raise ValueError(ui_text('debug.import_workbook.message_106', label=label))
        formula=re.sub(r'para\d+','1.0',parts[0]).replace('过去一分钟的铁生成量','1.0').replace('过去一分钟的宝石碎片生成量','1.0').replace('等级','1.0').replace('lv','1.0').replace('（','(').replace('）',')').replace('^','**')
        if not re.fullmatch(r'[0-9. +*/()\-]+',formula):
            raise ValueError(ui_text('debug.import_workbook.message_107', label=label))
        try:
            tree=ast.parse(formula.strip(),mode='eval')
        except SyntaxError as error:
            raise ValueError(ui_text('debug.import_workbook.message_124', label=label)) from error
        for node in ast.walk(tree):
            if not isinstance(node,(ast.Expression,ast.BinOp,ast.UnaryOp,ast.Add,ast.Sub,ast.Mult,ast.Div,ast.Pow,ast.UAdd,ast.USub,ast.Constant)) or (isinstance(node,ast.Constant) and (type(node.value) not in (int,float) or not math.isfinite(node.value))):
                raise ValueError(ui_text('debug.import_workbook.message_125', label=label))


def validate_weapon_motion(data):
    motion = data.get('weapon_motion', {})
    expected = {'player_projectile_pixels_per_unit','enemy_projectile_pixels_per_unit','player_cannon_speed_multiplier','chain_carrier_speed','missile_ejection_gap','missile_launch_speed','missile_turn_rate','missile_orphan_lifetime','missile_reacquire_interval','missile_departure_angle','missile_ignition_at','missile_seek_start','missile_cruise_at','missile_lifetime','missile_brake_range','missile_brake_angle','missile_min_guided_speed','missile_brake_factor','missile_hit_radius','missile_launch_edge_margin','missile_launch_forward_y'}
    if not isinstance(motion,dict):raise ValueError('weapon_motion: expected table')
    if 'weapon_motion' in data and set(motion) != expected:
        raise ValueError(f'weapon_motion: missing {sorted(expected-set(motion))}; unknown {sorted(set(motion)-expected)}')
    for key,row in motion.items():
        value = row.get('value')
        if type(value) not in (int,float):raise ValueError(f'weapon_motion {key}: expected numeric value')
        if key == 'missile_launch_forward_y':
            if type(value) not in (int,float) or not math.isfinite(value) or not -1 <= value <= 0:
                raise ValueError(f'weapon_motion {key}: expected finite direction in [-1,0]')
        else:
            positive(value, f'weapon_motion {key}', key in ('missile_ejection_gap','missile_departure_angle','missile_ignition_at','missile_seek_start'))
        if not str(row.get('unit','')).strip() or not str(row.get('description','')).strip():
            raise ValueError(f'weapon_motion {key}: unit and description required')
    if motion and float(motion['missile_cruise_at']['value']) <= float(motion['missile_ignition_at']['value']):
        raise ValueError('weapon_motion: cruise_at must exceed ignition_at')
    enemy_base=data.get('enemy_weapon_base',{})
    if 'enemy_weapon_base' in data and (not isinstance(enemy_base,dict) or set(enemy_base)!={'laser','missile','cannon','longLaser'}):raise ValueError('enemy_weapon_base: require all four base weapons')
    if motion:
        salvo=data['equipment']['missile'][0].get('para1')
        if type(salvo) not in (int,float) or not math.isfinite(salvo) or salvo<1 or salvo!=int(salvo):raise ValueError('equipment missile.para1: expected positive integer salvo quantity')
    for key,row in enemy_base.items():
        for field in ('dmg','cd','dmgtype'):
            positive(row.get(field), f'enemy_weapon_base {key}.{field}')
        # Require consumed fields, while retaining nullable unused parameters
        # and the beam's intentionally optional charge duration.
        required = ('para1','para2') if key in ('missile','longLaser') else ('para1',)
        for field in required:
            positive(row.get(field), f'enemy_weapon_base {key}.{field}', key == 'longLaser' and field == 'para1')
        if key == 'longLaser' and 'para3' not in row:
            raise ValueError('enemy_weapon_base longLaser.para3: charge column required (nullable)')
        for field in ('para1','para2','para3'):
            if row.get(field) is not None:positive(row[field], f'enemy_weapon_base {key}.{field}', True)
    if motion:
        data.get('defaults',{}).pop('projectilePixelsPerUnit',None)
    if enemy_base:
        data.setdefault('fallbacks',{}).update({'enemyWeaponMissingLevel':'Use enemy_weapon_base when the enemy weapon row is missing.','enemyCannonMissingDamageAndCooldown':'Use enemy_weapon_base for missing hostile weapon fields.'})

def validate_projection(data, *, check_level_ratios=True):
    validate_weapon_motion(data)
    validate_galaxy(data)
    validate_crew(data)
    validate_unlocks(data)
    validate_planets(data)
    equipment=data["equipment"]
    levels=data["levels"]
    groups=data["groups"]
    enemies=data["enemies"]
    config=data["config"]
    ships=data.get("ship",{})
    if ships:
        if len(ships) != 5:
            raise ValueError(ui_text('debug.import_workbook.message_04', ships=len(ships)))
        for key,row in ships.items():
            if not isinstance(row.get('des'),str) or not row['des'].strip():
                raise ValueError(ui_text('debug.import_workbook.message_126', key=key))
            for field in ('weaponSlots','defenseSlots','movement','size'):
                positive(row.get(field),f'ship {key} {field}', field == 'unlock')
            if row['weaponSlots'] != int(row['weaponSlots']) or row['weaponSlots'] < 1:
                raise ValueError(ui_text('debug.import_workbook.message_127', key=key))
            positive(row.get('sameEquipmentLimit'),f'ship {key} sameEquipmentLimit')
            if row['sameEquipmentLimit'] != int(row['sameEquipmentLimit']):
                raise ValueError(ui_text('debug.import_workbook.message_128', key=key))
            if row['defenseSlots'] != int(row['defenseSlots']) or row['defenseSlots'] < 1:
                raise ValueError(ui_text('debug.import_workbook.message_129', key=key))
    for key in ('armour','shield','laser','missile','cannon'):
        items=equipment.get(key,[])
        bases = [r for r in items if r.get('level') == 1]
        if len(bases) != 1:
            raise ValueError(ui_text('debug.import_workbook.message_108', key=key))
        for r in bases:
            positive(r['para1'] if key in ('armour','shield') else r['dmg'],f'{key} Lv.{r["level"]} value')
            if key not in ('armour','shield'):
                positive(r['cd'], f'{key} cooldown')
                positive(r['para2'] if key=='missile' else r['para1'],f'{key} projectile speed')
            growth = 'para2' if key == 'armour' else 'para4' if key == 'shield' else 'dmgMulti'
            positive(r.get(growth), f'{key} {growth}', True)
    for key, items in equipment.items():
        for r in (r for r in items if r.get('level') == 1):
            for field, value in r.items():
                if re.fullmatch(r'cost_\d+', str(field)) and value is not None:
                    positive(value, f'{key} {field}', True)
                    suffix = field.removeprefix('cost_')
                    positive(r.get('cost_multi_' + suffix, r.get('costMulti_' + suffix)), f'{key} cost multiplier {suffix}', True)
    if [r['id'] for r in levels] != list(range(1,len(levels)+1)) or not levels:
        raise ValueError(ui_text('debug.import_workbook.message_104'))
    for level in levels:
        for key in (('length','atkRatio','lifeRatio','resRatio') if check_level_ratios else ('length',)):
            positive(level[key],f'level {level["id"]} {key}')
        positive(level.get('planetExpRatio'), f'level {level["id"]} planetExpRatio', True)
        positions=[g['position'] for g in level['groups']]
        if not positions or positions!=sorted(set(positions)) or not all(0<=p<=1 for p in positions):
            raise ValueError(ui_text('debug.import_workbook.message_109', id=level["id"]))
        for g in level['groups']:
            if str(g['id']) not in groups: raise ValueError(ui_text('debug.import_workbook.message_131', id=g["id"]))
    for gid,g in groups.items():
        if len(g['slots'])!=10: raise ValueError(ui_text('debug.import_workbook.message_110', gid=gid))
        for enemy_id in g['slots']:
            if enemy_id is not None and str(enemy_id) not in enemies: raise ValueError(ui_text('debug.import_workbook.message_132', enemy_id=enemy_id))
    for enemy in enemies.values():
        eid = enemy['id']
        size = enemy['size']
        if type(size) not in (int, float) or not math.isfinite(size) or size != int(size) or size < 1:
            raise ValueError(ui_text('debug.import_workbook.message_05', eid=eid))
        if enemy['armourType'] not in (0, 1, 2): raise ValueError(ui_text('debug.import_workbook.message_06', eid=eid))
        for weapon in enemy['equipment']:
            key = weapon['name']
            base = key.replace('_mon', '').replace('-mon', '')
            candidates = equipment.get(key, []) + equipment.get(base, [])
            if base not in ('laser', 'cannon', 'missile') or not any(r['level'] == 1 for r in candidates):
                raise ValueError(ui_text('debug.import_workbook.message_13', eid=eid, key=key))
        positive(enemy['health'],f'enemy {enemy["id"]} health')
        positive(enemy['dmgMultiple'],f'enemy {enemy["id"]} damage multiplier',True)
        for drop in enemy['drops']:
            if str(drop['resourceId']) not in data['resources']: raise ValueError(ui_text('debug.import_workbook.message_14', eid=eid))
            positive(drop['amount'],'drop amount',True)
            if not 0<=drop['chance']<=1: raise ValueError(ui_text('debug.import_workbook.message_133'))
    for key in ('dmgReduce','autoCollectReduce'):
        if not isinstance(config[key],(int,float)) or not 0<=config[key]<1: raise ValueError(ui_text('debug.import_workbook.message_111', key=key))
    positive(config['movement'],'movement')
    positive(config['backRange'],'backRange',True)
    positive(config.get('offlineMax'),'offlineMax (hours)',True)
    positive(config.get('chronoParticlesPerSecond'),'chronoParticlesPerSecond')
    positive(config.get('chronoDefaultSpeed'),'chronoDefaultSpeed')
    for ship_key in data['ship']:
        key = 'playerVisualScale' + ship_key
        positive(config.get(key), key)
    for size in range(1, 7):
        key = f'enemyVisualScaleSize{size}'
        positive(config.get(key), key)
    speeds = config.get('chronoSpeeds')
    if not isinstance(speeds,str) or not speeds.strip():
        raise ValueError(ui_text('debug.import_workbook.chrono_speeds_format'))
    parsed_speeds = {}
    for entry in speeds.split(','):
        fields = entry.strip().split('|')
        if len(fields) != 2:
            raise ValueError(ui_text('debug.import_workbook.chrono_speeds_format'))
        try:
            multiplier, cost = map(float, fields)
        except ValueError as error:
            raise ValueError(ui_text('debug.import_workbook.chrono_speeds_number')) from error
        if not math.isfinite(multiplier) or multiplier <= 0 or not math.isfinite(cost) or cost < 0 or multiplier in parsed_speeds:
            raise ValueError(ui_text('debug.import_workbook.chrono_speeds_values'))
        parsed_speeds[multiplier] = cost
    if config['chronoDefaultSpeed'] not in parsed_speeds or parsed_speeds[config['chronoDefaultSpeed']] != 0:
        raise ValueError(ui_text('debug.import_workbook.chrono_default'))
    auto_gen = config.get('autoGenRes')
    if auto_gen is not None:
        if not isinstance(auto_gen, str):
            raise ValueError(ui_text('debug.import_workbook.message_112'))
        parts = [part.strip() for part in auto_gen.replace('，', ',').split(',')]
        if len(parts) != 4 or not parts[0] or not parts[1] or not parts[2] or not parts[3]:
            raise ValueError(ui_text('debug.import_workbook.message_113'))
        if not re.fullmatch(r'\d+', parts[1]) or str(int(parts[1])) not in data['resources']:
            raise ValueError(ui_text('debug.import_workbook.message_114', parts=repr(parts[1])))
        try:
            interval, amount, speed = (float(parts[0]), float(parts[2]), float(parts[3]))
        except ValueError as error:
            raise ValueError(ui_text('debug.import_workbook.message_134')) from error
        positive(interval, 'autoGenRes interval')
        positive(amount, 'autoGenRes amount', True)
        positive(speed, 'autoGenRes speed')
    threshold=config.get('hightechCostGrowthLevel')
    positive(threshold,'hightechCostGrowthLevel',True)
    if threshold != int(threshold):
        raise ValueError('hightechCostGrowthLevel must be an integer')
    limit=config.get('hightechLimit')
    positive(limit,'hightechLimit')
    positive(config.get('techPointGet'),'techPointGet')
    parts=str(config.get('scientistCost','')).replace('，',',').split(',')
    if len(parts)<2: raise ValueError(ui_text('debug.import_workbook.message_105'))
    try:
        positive(float(parts[0]),'scientistCost multiplier')
        seen=set()
        for part in parts[1:]:
            rid,amount=part.split('|')
            if rid not in data['resources'] or rid in seen: raise ValueError(ui_text('debug.import_workbook.message_135'))
            seen.add(rid)
            positive(float(amount),'scientistCost amount',True)
    except (ValueError,TypeError) as error:
        raise ValueError(ui_text('debug.import_workbook.message_115')) from error
    for key,row in data['hightech'].items():
        validate_description(row)
        if not isinstance(row.get('des'),str) or not row['des'].strip():
            raise ValueError(ui_text('debug.import_workbook.message_116', key=key))
        positive(row.get('tpCostBase'),f'hightech {key} tpCostBase')
        if math.floor(row['tpCostBase']+0.5)<1:
            raise ValueError(ui_text('debug.import_workbook.message_117', key=key))
        positive(row.get('tpCostMutiple'),f'hightech {key} tpCostMutiple',True)
        positive(row.get('tpCostMutiple2'),f'hightech {key} tpCostMutiple2',True)
        positive(row.get('para1'),f'hightech {key} para1',True)
        if row.get('para2') is not None:
            positive(row['para2'],f'hightech {key} para2',True)
        if key in ('超时空炼铁炉', '宝石熔炼炉'):
            positive(row.get('para1'),f'hightech {key} interval')
            positive(row.get('para2'),f'hightech {key} para2',True)

    for field in ('reactorEnergyBase','reactorEnergyGrowth','reactorUpgradeBase','reactorUpgradeGrowth','reactorBoostExponent','reactorInitialLevel','reactorPercentScale','reactorAllocationStep'):
        positive(data['config'].get(field),field)
    if data['config']['reactorEnergyGrowth'] <= 1 or data['config']['reactorUpgradeGrowth'] <= 1 or data['config']['reactorInitialLevel'] != int(data['config']['reactorInitialLevel']):
        raise ValueError('Invalid reactor growth or initial level')
    if data['config']['reactorAllocationStep'] != int(data['config']['reactorAllocationStep']):
        raise ValueError('reactorAllocationStep must be an integer')
    uranium = data['config'].get('reactorUraniumId')
    if type(uranium) not in (int,float) or uranium != int(uranium) or str(int(uranium)) not in data['resources']:
        raise ValueError('Invalid reactorUraniumId')
    modules = data['config'].get('reactorModules')
    if not isinstance(modules,str) or set(modules.split(',')) != {'weapons','defence','smelting','condensation'}:
        raise ValueError('Invalid reactorModules')

    data["defaults"].pop("maxEquipmentLevel", None)

def validate_unlocks(data):
    rows = data.get('unlock')
    if not isinstance(rows, dict) or not rows:
        raise ValueError('Missing unlock table; import unlock.xlsx first')
    targets = set()
    allowed = {kind: set(data.get(kind, {})) for kind in ('equipment', 'ship', 'hightech')}
    allowed['equipment'] = {key for key in allowed['equipment'] if not key.endswith(('_mon', '-mon'))}
    allowed['feature'] = {'jewels','reactor','crew_level'} | ({'galaxy'} if data.get('galaxy') else set())
    allowed['reactor_module'] = {'condensation'}
    allowed['crew'] = set(data.get('crew', {}))
    allowed['planet'] = set(data.get('planet', {}))
    for name, row in rows.items():
        kind, target = row.get('type'), row.get('target')
        identity = (kind, target)
        if row.get('name') != name or kind not in allowed or target not in allowed[kind] or identity in targets:
            raise ValueError(f'Invalid or duplicate unlock target: {name}')
        targets.add(identity)
        gate = row.get('level')
        if type(gate) not in (int, float) or not math.isfinite(gate) or gate != int(gate) or not 0 <= gate <= len(data['levels']):
            raise ValueError(f'Invalid unlock level: {name}')
        if row.get('mode') not in ('cleared', 'reached'):
            raise ValueError(f'Invalid unlock mode: {name}')
        if any(not isinstance(row.get(field), str) or not row[field].strip() for field in ('title', 'desc')):
            raise ValueError(f'Missing unlock title/desc: {name}')
    # Crew may explicitly be initially available (empty unlockId).
    missing = {(kind, key) for kind, keys in allowed.items() if kind != 'crew' for key in keys} - targets
    if missing:
        raise ValueError(f'Missing unlock targets: {sorted(missing)}')
    # Legacy columns may arrive from an old master. Never restore a second gate source.
    for kind in ('ship', 'hightech'):
        for row in data.get(kind, {}).values(): row.pop('unlock', None)
    for items in data.get('equipment', {}).values():
        for row in items: row.pop('unlock', None)
    data['config'].pop('jewelDropLevel', None)


def validate_crew(data):
    """Shared full/incremental checks; no crew values are defaulted by code."""
    def integer(value, label, zero=False):
        positive(value, label, zero)
        if isinstance(value, bool) or value != int(value):
            raise ValueError(f'{label}: expected integer')
    if data.get('crew'):
        settings = data.get('crew_config', {})
        positive(settings.get('base_exp', {}).get('value'), 'crew_config base_exp')
        multiplier = settings.get('exp_multiplier', {}).get('value')
        for key in ('equip_bonus', 'tech_ai_per_level', 'tech_speed', 'gem_bonus', 'charge_bonus'):
            row = settings.get(key, {})
            positive(row.get('value'), 'crew_config ' + key, True)
            if not isinstance(row.get('des'), str) or not row['des'].strip():
                raise ValueError('crew_config ' + key + ': missing des')
        integer(settings['tech_ai_per_level']['value'], 'tech_ai_per_level', True)
        for key in ('name_level', 'exp_bar', 'badge_tip', 'planet_exp', 'dedicated_ai'):
            if not isinstance(settings.get(key, {}).get('des'), str) or not settings[key]['des'].strip():
                raise ValueError('crew_config ' + key + ': missing des')
        for key, target in {'equip_bonus':'equipment', 'tech_ai_per_level':'hightech', 'tech_speed':'hightech', 'gem_bonus':'jewel', 'charge_bonus':'reactor'}.items():
            if settings[key].get('targetType') != target:
                raise ValueError('crew_config ' + key + ': invalid targetType')
        positive(multiplier, 'crew_config exp_multiplier')
        if multiplier < 1:
            raise ValueError('crew_config exp_multiplier must be at least 1')
    jobs = data.get('crew_assignment', {})
    contracts = json.loads((ROOT/'data/ui_text_contract.json').read_text(encoding='utf-8'))['entries'] if jobs else {}
    for id, row in jobs.items():
        for field in ('targetType', 'effectType', 'description'):
            if not isinstance(row.get(field), str) or not row[field].strip():
                raise ValueError(f'crew_assignment {id}: missing {field}')
        for field in ('baseValue', 'powerScale', 'levelScale', 'interval'):
            positive(row.get(field), f'crew_assignment {id} {field}', True)
        if row.get('levelScale', 0) != 0:
            raise ValueError(f'{id}: crew level effect scaling is retired; levelScale must be 0')
        integer(row.get('maxCrew'), f'crew_assignment {id} maxCrew')
        if row['effectType'] in ('AUTO_UPGRADE','AUTO_SCIENTIST','AUTO_COMBINE','AUTO_REACTOR'):
            positive(row['baseValue'], f'{id} baseValue')
            positive(row['interval'], f'{id} interval')
            if row['effectType'] in ('AUTO_UPGRADE','AUTO_SCIENTIST'):
                modes = [mode.strip() for mode in str(row.get('upgradeModes','')).split(',')]
                if not modes or len(set(modes))!=len(modes) or any(mode!='max' and (not mode.isdigit() or not 0<int(mode)<=2147483647) for mode in modes):
                    raise ValueError(f'{id}: upgradeModes requires positive integers or max')
        for field in ('titleTextId','descTextId'):
            if not isinstance(row.get(field,''),str):
                raise ValueError(f'{id}: {field} must be text')
            key = row.get(field, '')
            if key and (key not in contracts or not set(contracts[key]['params']).issubset(set() if field=='titleTextId' else {'value','interval','description','mode'})):
                raise ValueError(f'{id}: unknown or incompatible {field}: {key}')
    for id, row in data.get('crew', {}).items():
        for field in ('name', 'description', 'icon'):
            if not isinstance(row.get(field), str) or not row[field].strip():
                raise ValueError(f'crew {id}: missing {field}')
        for field in ('equipmentSlotCount',):
            integer(row.get(field), f'crew {id} {field}', field == 'equipmentSlotCount')
        positive(row.get('basePower'), f'crew {id} basePower')
        for job in jobs.values():
            try:
                value = float(job['baseValue'])*float(row['basePower'])**float(job['powerScale'])
                if not math.isfinite(value) or (job['effectType'] in ('AUTO_UPGRADE','AUTO_SCIENTIST','AUTO_COMBINE','AUTO_REACTOR') and (value<=0 or not math.isfinite(float(job['interval'])/value))):
                    raise ValueError(f'crew {id}: nonfinite effective power or interval')
            except OverflowError as error:
                raise ValueError(f'crew {id}: effect power overflow') from error
        if row.get('defaultAssignment') and row['defaultAssignment'] not in jobs:
            raise ValueError(f'crew {id}: unknown defaultAssignment')
        if bool(row.get('defaultAssignment')) != bool(row.get('defaultTargetId')):
            raise ValueError(f'crew {id}: defaultAssignment requires defaultTargetId')
        if row.get('unlockId') and row['unlockId'] not in data.get('unlock', {}):
            raise ValueError(f'crew {id}: unknown unlockId')


def validate_planets(data):
    unlock_edges = {}
    for id, row in data.get('planet_buff', {}).items():
        target = row.get('value', 0) if row.get('buff_type') == 'planet_unlock' else 0
        if type(target) not in (int, float) or not math.isfinite(target) or target < 0 or target != int(target):
            raise ValueError(f'planet_buff {id}: invalid planet_unlock value')
        if row.get('buff_type') == 'planet_unlock':
            if not target or row.get('stack') != 'max' or row.get('target') != 'planet':
                raise ValueError(f'planet_buff {id}: planet_unlock requires planet/max and a positive ID')
            target = str(int(target))
            source = str(int(row['planet_id']))
            if target not in data.get('planet', {}) or target == source:
                raise ValueError(f'planet_buff {id}: unknown or self unlock target')
            if row.get('source') != 'conquer' or row.get('condition') != 'conquered':
                raise ValueError(f'planet_buff {id}: planet unlock requires conquest')
            if target in unlock_edges and unlock_edges[target] != source:
                raise ValueError(f'planet_buff {id}: conflicting unlock sources')
            unlock_edges[target] = source

        required = 'id planet_id source source_id condition buff_type target value stack order des'.split()
        if any(key not in row for key in required):
            raise ValueError(f'planet_buff {id}: missing field')
        if type(row['id']) is not int or row['id'] < 1 or str(row['id']) != id:
            raise ValueError(f'planet_buff {id}: invalid numeric ID')
        if type(row['planet_id']) is not int or str(row['planet_id']) not in data.get('planet', {}):
            raise ValueError(f'planet_buff {id}: unknown planet')
        if row['source'] not in ('conquer','planet','building','event') or row['condition'] not in ('always','conquered','building_complete'):
            raise ValueError(f'planet_buff {id}: unsupported source or condition')
        if row['condition']=='building_complete' and str(row['source_id']) not in data.get('planet_build', {}):
            raise ValueError(f'planet_buff {id}: unknown building')
        if (row['buff_type'],row['target']) not in (('level_bonus','equipment'),('level_bonus','hightech'),('free_charge','all'),('level_bonus','enhancement'),('crew_exp_share','all'),('planet_unlock','planet')):
            raise ValueError(f'planet_buff {id}: unsupported buff target')
        if row['stack'] not in ('add','mul','max'):
            raise ValueError(f'planet_buff {id}: unsupported stack')
        if row['buff_type']=='crew_exp_share' and (row['stack']!='max' or row['value'] not in (0,1)):
            raise ValueError(f'planet_buff {id}: crew_exp_share requires max and 0/1')
        for key in ('value','order'):positive(row[key],f'planet_buff {id} {key}',True)
        if row['order'] != int(row['order']) or (row['buff_type'] != 'free_charge' and row['stack'] != 'mul' and row['value'] != int(row['value'])):
            raise ValueError(f'planet_buff {id}: expected integer level/order')
        if not isinstance(row['des'], str) or not row['des'].strip():
            raise ValueError(f'planet_buff {id}: missing description')
    if any(row.get('buff_type') == 'planet_unlock' for row in data.get('planet_buff', {}).values()):
        first = min(data.get('planet', {}), key=int)
        if first in unlock_edges:raise ValueError('planet_buff: first planet must not require conquest')
        for planet_id in data['planet']:
            seen = set()
            current = planet_id
            while current != first:
                if current in seen or current not in unlock_edges:
                    raise ValueError(f'planet_buff: missing or cyclic unlock path for {planet_id}')
                seen.add(current)
                current = unlock_edges[current]
    for id, row in data.get('planet', {}).items():
        start = row.get('reforgeStartLevel', 1)
        if type(start) not in (int,float) or not math.isfinite(start) or start != int(start) or not 1 <= start <= len(data.get('levels',[])):
            raise ValueError(f'planet {id}: invalid reforgeStartLevel')
        if str(row.get('id')) != id or not isinstance(row.get('name'), str) or not row['name'].strip():
            raise ValueError(f'planet {id}: invalid identity or name')
        positive(row.get('baseTime'), f'planet {id} baseTime')
        minimum = row.get('minTime', 5)
        if type(minimum) not in (int, float):
            raise ValueError(f'planet {id}: invalid minTime')
        positive(minimum, f'planet {id} minTime')
        positive(row.get('baseExp'), f'planet {id} baseExp')
        if type(row.get('id')) is not int or row['id']<1 or not isinstance(row.get('tags',''),str):
            raise ValueError(f'planet {id}: invalid numeric ID or tags')
        gate = data.get('unlock', {}).get(row.get('unlockId'), {})
        if gate.get('type') != 'planet' or gate.get('target') != id:
            raise ValueError(f'planet {id}: invalid unlockId')

    settings = data.get('enhance_config', {})
    if settings:
        required = 'cost_base cost_growth threshold_1 threshold_2 threshold_3 base_critical_rate base_critical_multiplier critical_growth proficiency_growth adaptation_growth counter_log_base bonus_round_scale repeat_probability repeat_growth repeat_delay memory_interval memory_heal_fraction memory_buffer_fraction deferred_duration deferred_interval deferred_curve_coefficient deferred_percent_scale deferred_clear_probability branch_threshold_1 branch_threshold_2 branch_threshold_3 branch_position_offset_1 branch_position_offset_2 branch_position_offset_3'.split()
        for key in required:
            row = settings.get(key, {})
            positive(row.get('value'), 'enhance_config ' + key, key in ('branch_position_offset_1','branch_position_offset_2','branch_position_offset_3','base_critical_rate','critical_growth','proficiency_growth','adaptation_growth','repeat_probability','repeat_growth','memory_heal_fraction','memory_buffer_fraction','deferred_clear_probability'))
            if not isinstance(row.get('des'), str) or not row['des'].strip() or not isinstance(row.get('unit'),str) or not row['unit'].strip():
                raise ValueError('enhance_config ' + key + ': missing unit or description')
        for key in ('base_critical_rate','repeat_probability','deferred_clear_probability'):
            if settings[key]['value'] > 1:raise ValueError('enhance_config ' + key + ': probability exceeds one')
        thresholds = [settings['threshold_' + str(i)]['value'] for i in (1,2,3)]
        if any(isinstance(n,bool) or n != int(n) for n in thresholds) or not thresholds[0] < thresholds[1] < thresholds[2]:
            raise ValueError('enhance_config: expected ascending integer thresholds')
        if settings['counter_log_base']['value'] <= 1:raise ValueError('enhance_config: logarithm base must exceed one')
        if settings['cost_base']['value'] != int(settings['cost_base']['value']):raise ValueError('enhance_config: cost_base must be positive integer fragments')
        growth=settings['cost_growth']['value']
        if isinstance(growth,bool) or growth != int(growth) or growth < 1:raise ValueError('enhance_config: cost_growth must be integer >=1')
        branch_thresholds=[settings['branch_threshold_'+str(i)]['value'] for i in (1,2,3)]
        if any(n != int(n) for n in branch_thresholds) or not branch_thresholds[0] < branch_thresholds[1] < branch_thresholds[2]:raise ValueError('enhance_config: expected ascending integer branch thresholds')
        offsets=[settings['branch_position_offset_'+str(i)]['value'] for i in (1,2,3)]
        if any(n != int(n) for n in offsets) or not offsets[0] < offsets[1] < offsets[2]:raise ValueError('enhance_config: expected ascending integer branch position offsets')
        precision=settings['deferred_percent_scale']['value']
        if precision != int(precision):raise ValueError('enhance_config: deferred_percent_scale must be positive integer')

    if settings:
        for key in ['proficiency_a_damage_bonus', 'proficiency_b1_interval', 'proficiency_b1_growth', 'proficiency_b2_interval_multiplier', 'proficiency_b3_resistance', 'repeat_a_probability', 'repeat_b1_targets', 'repeat_b2_repeats', 'repeat_b3_probability', 'critical_a_probability', 'critical_b1_attacks', 'critical_b1_damage_bonus', 'critical_b2_probability', 'critical_b2_damage_bonus', 'critical_b2_stacks', 'critical_b2_duration', 'critical_b3_guaranteed_rate', 'adaptation_a_capacity_bonus', 'adaptation_b1_reduction', 'adaptation_b1_stacks', 'adaptation_b2_interval', 'adaptation_b2_duration', 'adaptation_b2_capacity_multiplier', 'adaptation_b3_resistance_bonus', 'memory_a_bonus', 'memory_b1_probability', 'memory_b1_reduction', 'memory_b1_duration', 'memory_b2_resistance', 'memory_b2_duration', 'memory_b2_lockout', 'memory_b3_shield_charge_bonus', 'memory_b3_armour_capacity_bonus', 'deferred_a_clear_probability', 'deferred_b1_reduction', 'deferred_b1_duration', 'deferred_b2_probability_conversion', 'deferred_b3_forced_probability', 'deferred_b3_damage_scale']:
            row=settings.get(key,{})
            positive(row.get('value'),'enhance_config '+key,key not in ['proficiency_b1_interval', 'proficiency_b2_interval_multiplier', 'critical_b2_duration', 'adaptation_b2_interval', 'adaptation_b2_duration', 'memory_b1_duration', 'memory_b2_duration', 'memory_b2_lockout', 'deferred_b1_duration'])
            if not isinstance(row.get('des'),str) or not row['des'].strip() or not isinstance(row.get('unit'),str) or not row['unit'].strip():raise ValueError('enhance_config '+key+': missing unit or description')
        for key in ['critical_b3_guaranteed_rate','memory_b1_reduction','adaptation_b1_reduction','deferred_b1_reduction','deferred_b2_probability_conversion','proficiency_b3_resistance', 'repeat_a_probability', 'repeat_b3_probability', 'critical_a_probability', 'critical_b2_probability', 'memory_b1_probability', 'memory_b2_resistance', 'deferred_a_clear_probability', 'deferred_b2_probability_conversion', 'deferred_b3_forced_probability']:
            if settings[key]['value']>1:raise ValueError('enhance_config '+key+': fraction exceeds one')
        for key in ['repeat_b1_targets', 'repeat_b2_repeats', 'critical_b1_attacks', 'critical_b2_stacks', 'adaptation_b1_stacks']:
            if settings[key]['value']!=int(settings[key]['value']):raise ValueError('enhance_config '+key+': count must be integer')
        if settings['repeat_b1_targets']['value']>10 or settings['repeat_b2_repeats']['value']>10:raise ValueError('enhance_config: derived attack count exceeds bounded budget 10')

    for id, row in data.get('planet_build', {}).items():
        required = 'id name des min_planet planet_rule planet_value order unlock_explore build_explore extra_crew type config1 config2'.split()
        if any(key not in row for key in required):
            raise ValueError(f'planet_build {id}: missing field')
        for key in ('min_planet','order','unlock_explore','build_explore','extra_crew'):
            positive(row[key], f'planet_build {id} {key}',key not in ('min_planet','build_explore'))
            if row[key]!=int(row[key]):raise ValueError(f'planet_build {id}: {key} must be integer')
        if row['planet_rule'] not in ('all','only','exclude','tag'):
            raise ValueError(f'planet_build {id}: unsupported rule')
        if row['planet_rule'] in ('only','exclude'):
            values=str(row['planet_value']).split(',')
            if any(not x.isdigit() or int(x)<1 for x in values):raise ValueError(f'planet_build {id}: invalid planet IDs')
        if row['planet_rule']=='tag' and (not row['planet_value'] or ',' in row['planet_value']):
            raise ValueError(f'planet_build {id}: expected one tag')
        if row['type'] not in ('auto_explore','refinery','equipment','shipyard'):
            raise ValueError(f'planet_build {id}: unsupported effect handler')
        for key in ('config1','config2'):positive(row[key], f'planet_build {id} {key}',True)


def encode(data):
    return json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False).encode("utf-8")

def full_import(source, target):
    previous=json.loads(target.read_text(encoding="utf-8")) if target.exists() else {}
    data=projection_base(previous,source.name)

    book=openpyxl.load_workbook(source,data_only=True,read_only=True)
    try:
        for name,section in SECTIONS.items():
            if name not in book.sheetnames:
                if name in ('ship','unlock','crew','crew_assignment','crew_config','planet','planet_build','planet_buff','enhance_config','weapon_motion','enemy_weapon_base'):
                    continue  # Older master workbooks predate optional projections.
                raise ValueError(ui_text('debug.import_workbook.message_15', name=name))
            data[section]=convert_sheet(name,read_rows(book[name]))
    finally:
        book.close()
    validate_projection(data)
    target.parent.mkdir(parents=True,exist_ok=True)
    temporary=target.with_suffix(".json.tmp")
    try:
        temporary.write_bytes(encode(data))
        os.replace(temporary,target)
    finally:
        if temporary.exists(): temporary.unlink()
    return data

def main():
    if hasattr(sys.stdout,"reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    parser=argparse.ArgumentParser()
    parser.add_argument("source",nargs="?",type=pathlib.Path,default=ROOT.parent/"太空战舰.xlsx")
    parser.add_argument("target",nargs="?",type=pathlib.Path,default=ROOT/"data"/"game_data.json")
    args=parser.parse_args()
    data=full_import(args.source,args.target)
    print(f'Imported {len(data["levels"])} levels, {len(data["enemies"])} enemies, {len(data["groups"])} groups, {sum(map(len,data["equipment"].values()))} equipment rows.')

if __name__=="__main__":
    main()
