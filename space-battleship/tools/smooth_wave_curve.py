"""Preview/apply smooth combat bands from Excel-authored stage tail factors.

Edit atkMultiplier/lifeMultiplier on each selected stage's last monGroup row,
then run with --apply and import via config_workbooks.py. Earlier stage tails
anchor the next head. This writes only monGroup.xlsx; level ratios/drop data
remain unchanged. Intermediate absolute combat ratios grow geometrically. Optional authorAtkScale
shapes attack bands only; tail factors remain the final runtime values.
"""
import argparse
import math
import pathlib
import os
import openpyxl
from import_workbook import ROOT, convert_sheet, read_rows


def factor(value, label):
    if value in (None, ''):
        return 1.0
    if type(value) not in (int, float) or not math.isfinite(value) or value <= 0:
        raise ValueError(f'{label}: expected a positive finite number or blank')
    return float(value)


def smooth(directory, stages, apply=False):
    level_book = openpyxl.load_workbook(directory / 'level.xlsx', read_only=True, data_only=True)
    try:
        levels = {int(r['id']): r for r in convert_sheet('level', read_rows(level_book['level']))}
    finally:
        level_book.close()
    path = directory / 'monGroup.xlsx'
    book = openpyxl.load_workbook(path)
    try:
        sheet = book['monGroup']
        columns = {cell.value: cell.column for cell in sheet[1]}
        rows = {int(sheet.cell(i, columns['id']).value): i for i in range(4, sheet.max_row + 1)
                if sheet.cell(i, columns['id']).value is not None}
        changes = []
        for stage in sorted(set(stages)):
            level = levels[stage]
            points = level['groups']
            previous = levels.get(stage - 1)
            summary = {'stage': stage}
            for kind, key, entry in [('atkRatio', 'atkMultiplier', 'entryAtkRatio'),
                                     ('lifeRatio', 'lifeMultiplier', 'entryLifeRatio')]:
                tail = factor(sheet.cell(rows[points[-1]['id']], columns[key]).value, key)
                anchor = float(previous[kind]) if previous else 1.0
                previous_factor = factor(sheet.cell(rows[previous['groups'][-1]['id']], columns[key]).value, key) if previous else 1.0
                head = float(level.get(entry, anchor))
                shape_column = columns.get('authorAtkScale') if key == 'atkMultiplier' else None
                tail_shape = factor(sheet.cell(rows[points[-1]['id']], shape_column).value, 'authorAtkScale') if shape_column else 1.0
                start, end = anchor * previous_factor, float(level[kind]) * tail / tail_shape
                for i, point in enumerate(points):
                    t = i / (len(points) - 1) if len(points) > 1 else 1.0
                    base = head * (1 - t) + float(level[kind]) * t
                    absolute = start ** (1 - t) * end ** t
                    shape = factor(sheet.cell(rows[point['id']], shape_column).value, 'authorAtkScale') if shape_column else 1.0
                    value = tail if i == len(points) - 1 else absolute / base * shape
                    factor(value, f'{stage}/{i + 1} {key}')
                    changes.append((rows[point['id']], columns[key], value))
                summary[key] = {'tail': tail, 'absolute_stage_growth': float(level[kind]) * tail / start,
                                'unshaped_adjacent_wave_growth': (end / start) ** (1 / (len(points) - 1)) if len(points) > 1 else 1.0}
            print(summary)
        if apply:
            for row, column, value in changes:
                sheet.cell(row, column, value)
            temporary = path.with_suffix('.xlsx.tmp')
            try:
                book.save(temporary)
                os.replace(temporary, path)
            finally:
                temporary.unlink(missing_ok=True)
        return changes
    finally:
        book.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--directory', type=pathlib.Path, default=ROOT / 'config_excel')
    parser.add_argument('--stages', type=int, nargs='+', default=list(range(11, 21)))
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    smooth(args.directory, args.stages, args.apply)
