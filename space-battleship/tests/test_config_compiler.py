import copy
import importlib.util
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('config_compile',ROOT/'config_tools/compile.py')
compiler=importlib.util.module_from_spec(spec);spec.loader.exec_module(compiler)


class CompilerCases(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=ROOT/'tests/work')
        self.root=Path(self.temp.name)
        shutil.copy2(ROOT/'tests/excel/core_rules.xlsx',self.root/'core_rules.xlsx')
        shutil.copy2(ROOT/'tests/excel/core_rules.json',self.root/'core_rules.json')
        self.manifest=self.root/'core_rules.json'

    def tearDown(self):self.temp.cleanup()

    def edit(self,action):
        path=self.root/'core_rules.xlsx'
        book=openpyxl.load_workbook(path);action(book);book.save(path);book.close()

    def expect_error(self,code):
        before=compiler.compile_bundle(ROOT/'tests/excel/core_rules.json')
        output=self.root/'out';compiler.publish(before,output)
        pointer=(output/'active-manifest.json').read_bytes()
        with self.assertRaises(compiler.CompilationError) as caught:compiler.compile_bundle(self.manifest)
        self.assertIn(code,{e['code'] for e in caught.exception.errors})
        self.assertEqual(pointer,(output/'active-manifest.json').read_bytes())
        return caught.exception.errors

    def test_deterministic_and_provenance(self):
        a=compiler.compile_bundle(self.manifest);b=compiler.compile_bundle(self.manifest)
        self.assertEqual(a,b)
        self.assertEqual(a['provenance']['res']['[71]']['name']['row'],4)

    def test_duplicate(self):
        self.edit(lambda b:b['res'].append([71,'duplicate']))
        self.expect_error('DuplicateId')

    def test_reference(self):
        def change(b):
            s=b['equipment'];headers=[c.value for c in s[1]];s.cell(4,headers.index('res_1')+1,999)
        self.edit(change)
        self.assertTrue(any(e['column'] for e in self.expect_error('MissingReference')))

    def test_numeric_text(self):
        self.edit(lambda b:setattr(b['res']['A4'],'value','71'))
        self.expect_error('LiteralType')

    def test_missing_key(self):
        self.edit(lambda b:setattr(b['res']['A4'],'value',None))
        self.expect_error('EmptyPrimaryKey')

    def test_fixture_does_not_relax_production_requiredness(self):
        schema=json.loads(compiler.SCHEMA.read_text(encoding='utf-8'))
        self.assertNotIn('null',schema['tables']['res']['fields']['name']['types'])
        self.assertNotIn('null',schema['tables']['mon']['fields']['health']['types'])
        self.edit(lambda b:setattr(b['res']['B4'],'value',None))
        self.expect_error('LiteralType')

    def test_formula_without_cache(self):
        self.edit(lambda b:setattr(b['res']['A4'],'value','=70+1'))
        self.expect_error('FormulaCacheMissing')

    def test_unknown_sheet(self):
        self.edit(lambda b:b.create_sheet('secret'))
        self.expect_error('UnexpectedSheet')

    def test_typed_table_bad_header(self):
        self.edit(lambda b:b.create_sheet('bad_typed').append(['wrong']))
        manifest=json.loads(self.manifest.read_text())
        manifest['workbooks'][0]['sheets'].append(dict(sheet='bad_typed',table='hyperspace_config_00'))
        self.manifest.write_text(json.dumps(manifest))
        self.expect_error('HeaderMismatch')

    def test_unknown_rule(self):
        self.edit(lambda b:setattr(b['amount_case']['B4'],'value','free_money'))
        self.expect_error('UnknownRuleKey')

    def test_nested_reference_and_battle_order(self):
        self.edit(lambda b:setattr(b['level']['C4'],'value','{99999|0.6,9201|0.4}'))
        errors=self.expect_error('MissingReference')
        self.assertIn('BattlePointOrder',{e['code'] for e in errors})

    def test_topology_cycle(self):
        schema=json.loads(compiler.SCHEMA.read_text(encoding='utf-8'))['tables']['ai_tech_tree']
        def add(book):
            s=book.create_sheet('ai_tech_tree');columns=schema['columns']
            s.append(columns);s.append(['independent fixture']*len(columns));s.append(['literal']*len(columns))
            for key,dependency in [('test_a','test_b:1'),('test_b','test_a:1')]:
                row=dict(id=key,name='fixture',description='fixture',tier='basic',base_seconds=3,growth=2,effect='ai_count',factor=1,requirements=dependency)
                s.append([row.get(c) for c in columns])
        self.edit(add)
        manifest=json.loads(self.manifest.read_text());manifest['workbooks'][0]['sheets'].append(dict(sheet='ai_tech_tree',table='ai_tech_tree'))
        self.manifest.write_text(json.dumps(manifest))
        self.expect_error('DependencyCycle')

    def test_row_reordering_semantics(self):
        first=compiler.compile_bundle(self.manifest)
        def swap(b):
            s=b['amount_case'];a=[c.value for c in s[4]];z=[c.value for c in s[5]]
            for i,(x,y) in enumerate(zip(a,z),1):s.cell(4,i,y);s.cell(5,i,x)
        self.edit(swap)
        second=compiler.compile_bundle(self.manifest)
        self.assertEqual(first['semanticHash'],second['semanticHash'])
        self.assertNotEqual(first['sourceManifestHash'],second['sourceManifestHash'])

    def test_manifest_domain_isolation(self):
        value=json.loads(self.manifest.read_text());value['kind']='production'
        self.manifest.write_text(json.dumps(value))
        self.expect_error('ManifestIsolation')

    def test_active_pointer_failure(self):
        bundle=compiler.compile_bundle(self.manifest);output=self.root/'out'
        compiler.publish(bundle,output);old=(output/'active-manifest.json').read_bytes()
        real_replace=compiler.os.replace
        def fail(src,dst):
            if Path(dst).name=='active-manifest.json':raise OSError('injected rename')
            return real_replace(src,dst)
        with patch.object(compiler.os,'replace',side_effect=fail):
            with self.assertRaises(OSError):compiler.publish(bundle,output)
        self.assertEqual(old,(output/'active-manifest.json').read_bytes())


if __name__=='__main__':
    (ROOT/'tests/work').mkdir(parents=True,exist_ok=True)
    unittest.main()
