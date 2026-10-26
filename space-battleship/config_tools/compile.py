"""Validate manifest-listed Excel workbooks before publishing runtime JSON.

The emitted JSON retains the existing GameSession table shape. Validation and
publication happen before any running configuration is replaced.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
import tempfile
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_MANIFEST = Path(__file__).with_name("production_manifest.json")
DEFAULT_OUTPUT = ROOT / "generated" / "config.json"

# This file is also read by GameSession before native systems are created.
HYPERSPACE_CONTRACT = json.loads(Path(__file__).with_name("hyperspace_parameter_contract.json").read_text(encoding="utf-8"))
HYPERSPACE_INT_KEYS = frozenset(HYPERSPACE_CONTRACT["intKeys"])
HYPERSPACE_POSITIVE_KEYS = frozenset(HYPERSPACE_CONTRACT["positiveKeys"])
HYPERSPACE_NONNEGATIVE_KEYS = frozenset(HYPERSPACE_CONTRACT["nonnegativeKeys"])
HYPERSPACE_PROBABILITY_KEYS = frozenset(HYPERSPACE_CONTRACT["probabilityKeys"])


def quantized_range_has_value(low: object, high: object, precision: object) -> bool:
    """Match Hyperspace._quantized's lattice and its 1e-7 boundary tolerance."""
    if not all(valid_type(value, "number") for value in (low, high, precision)) or precision <= 0 or low > high:
        return False
    lower, upper = low / precision - 1e-7, high / precision + 1e-7
    if not math.isfinite(lower) or not math.isfinite(upper):
        return False
    return math.ceil(lower) <= math.floor(upper)


def stable_id(value: object) -> str | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return str(value)
    if isinstance(value, str) and value and value == value.strip():
        return value
    return None


def literal_kind(value: object) -> str:
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (int, float)):
        return "number" if math.isfinite(value) else "invalid"
    if isinstance(value, str):
        return "text" if value and value == value.strip() else "invalid"
    return "invalid"


def valid_type(value: object, kind: str) -> bool:
    if kind == "literal":
        return literal_kind(value) != "invalid"
    if kind == "bool":
        return isinstance(value, bool)
    if kind == "int":
        return isinstance(value, int) and not isinstance(value, bool)
    if kind == "number":
        return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)
    if kind == "text":
        return isinstance(value, str) and bool(value.strip()) and value == value.strip()
    if kind == "amount_text":
        return isinstance(value, str) and bool(re.fullmatch(r"(?:0|[1-9]\d*)(?:\.\d+)?", value))
    raise ValueError(f"unknown schema type: {kind}")


class Compiler:
    def __init__(self, manifest_path: Path):
        self.manifest_path = manifest_path.resolve()
        self.manifest = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        self.source = (self.manifest_path.parent / self.manifest["sourceRoot"]).resolve()
        self.errors: list[dict] = []
        self.rows: dict[tuple[str, str], list[dict]] = {}
        self.runtime: dict[str, dict[str, list[dict]]] = {}
        self.provenance: dict[tuple[str, str], dict[tuple, int]] = {}

    def error(self, book: str, sheet: str, row: int, column: str, key: str, code: str, detail: str) -> None:
        self.errors.append({"workbook": book, "sheet": sheet, "row": row,
                            "column": column, "key": key, "code": code, "detail": detail})

    def row_error(self, book: str, sheet: str, record: dict, field: str, code: str, detail: str) -> None:
        self.error(book + ".xlsx", sheet, record["_row"], record["_columns"].get(field, ""),
                   str(record["_key"]), code, detail)

    def read(self) -> None:
        if self.manifest.get("schemaVersion") != 1 or self.manifest.get("kind") not in ("production", "test") or not isinstance(self.manifest.get("fullGame"), bool):
            raise ValueError("InvalidManifest")
        books = self.manifest.get("workbooks")
        if not isinstance(books, list) or not books:
            raise ValueError("InvalidManifest")
        expected = {b["file"] for b in books}
        if len(expected) != len(books) or any(Path(name).name != name or not name.endswith(".xlsx") for name in expected):
            raise ValueError("InvalidManifest")
        actual = {p.name for p in self.source.glob("*.xlsx")}
        if actual != expected:
            self.error("manifest", "", 0, "", "", "InputSetMismatch",
                       f"missing={sorted(expected-actual)} extra={sorted(actual-expected)}")
            return
        production_source = (ROOT / "config_excel").resolve()
        if self.manifest["kind"] == "test" and self.source == production_source:
            raise ValueError("TestManifestUsesProductionInput")
        if self.manifest["kind"] == "production" and self.source != production_source:
            raise ValueError("ProductionManifestUsesTestInput")
        for book in books:
            self._read_book(book)

    def _read_book(self, book: dict) -> None:
        name = book["file"]
        stem = Path(name).stem
        path = self.source / name
        try:
            formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
            cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            self.error(name, "", 0, "", "", "InvalidWorkbook", str(exc))
            return
        schemas = {s["name"]: s for s in book["sheets"]}
        if set(formulas.sheetnames) != set(schemas):
            self.error(name, "", 0, "", "", "SheetSetMismatch",
                       f"expected={sorted(schemas)} actual={sorted(formulas.sheetnames)}")
        self.runtime[stem] = {}
        for sheet_name, schema in schemas.items():
            if sheet_name not in formulas:
                continue
            self._read_sheet(name, stem, formulas[sheet_name], cached[sheet_name], schema)
        formulas.close()
        cached.close()

    def _read_sheet(self, book: str, stem: str, formula_sheet, cached_sheet, schema: dict) -> None:
        sheet = schema["name"]
        header = [c.value for c in formula_sheet[1]]
        if header != schema["columns"]:
            self.error(book, sheet, 1, "", "", "HeaderMismatch", "sheet columns differ from manifest")
            return
        columns = {name: get_column_letter(i + 1) for i, name in enumerate(header) if name is not None}
        required = set(schema["required"])
        types = schema["types"]
        primary = schema["primary"]
        seen: set[tuple] = set()
        rows: list[dict] = []
        output: list[dict] = []
        provenance: dict[tuple, int] = {}
        for frow, crow in zip(formula_sheet.iter_rows(min_row=4), cached_sheet.iter_rows(min_row=4)):
            if not any(c.value is not None for c in frow):
                continue
            row_number = frow[0].row
            values = {name: crow[i].value for i, name in enumerate(header) if name is not None}
            key = tuple("<none>" if schema.get("allowNullPrimary") and values.get(k) is None
                        else stable_id(values.get(k)) for k in primary)
            label = "/".join(str(part) for part in key)
            if any(part is None for part in key):
                self.error(book, sheet, row_number, columns[primary[0]], label,
                           "InvalidPrimaryKey", "primary key must be stable and nonempty")
            elif key in seen:
                self.error(book, sheet, row_number, columns[primary[0]], label,
                           "DuplicatePrimaryKey", "primary key already occurs in this sheet")
            else:
                seen.add(key)
                provenance[key] = row_number
            for i, name in enumerate(header):
                if name is None:
                    continue
                formula = frow[i]
                value = values[name]
                if formula.data_type == "f":
                    if stem == "hyperspace_config":
                        self.error(book, sheet, row_number, columns[name], label,
                                   "FormulaForbidden", "typed literal content must be literal")
                    if value is None or crow[i].data_type == "e":
                        self.error(book, sheet, row_number, columns[name], label,
                                   "FormulaCacheMissing", "formula has no valid saved result")
                if formula.data_type == "e" or crow[i].data_type == "e":
                    self.error(book, sheet, row_number, columns[name], label,
                               "ExcelError", "Excel error values cannot be published")
                if value is None or value == "":
                    if name in required:
                        self.error(book, sheet, row_number, columns[name], label,
                                   "RequiredValueMissing", "required field is empty")
                    continue
                kind = types[name]
                if kind == "literal" and name == "value" and schema.get("allowedKeys"):
                    kind = schema["allowedKeys"].get(str(values.get(primary[0])), "literal")
                if not valid_type(value, kind):
                    self.error(book, sheet, row_number, columns[name], label,
                               "InvalidFieldType", f"expected {kind}, got {type(value).__name__}")
            if schema.get("allowedKeys") and len(primary) == 1:
                rule = str(values.get(primary[0]))
                if rule not in schema["allowedKeys"]:
                    self.error(book, sheet, row_number, columns[primary[0]], label,
                               "UnknownRuleKey", "key absent from registered rule schema")
            if schema.get("fixedValues") and str(values.get(primary[0])) in schema["fixedValues"]:
                if values.get("value") != schema["fixedValues"][str(values.get(primary[0]))]:
                    self.error(book, sheet, row_number, columns.get("value", ""), label,
                               "FixedRuleChanged", "structural rule requires a schema revision")
            if "data_type" in values and "value" in values and values["value"] is not None:
                declared = {"float": "number", "number": "number", "string": "text",
                            "int": "int", "bool": "bool"}.get(str(values["data_type"]))
                if declared is None or not valid_type(values["value"], declared):
                    self.error(book, sheet, row_number, columns["value"], label,
                               "TypedLiteralMismatch", f"declared={values['data_type']}")
            for field in ("chance", "probability"):
                value = values.get(field)
                if value is not None and (not valid_type(value, "number") or not 0 <= value <= 1):
                    self.error(book, sheet, row_number, columns[field], label,
                               "ProbabilityRange", "probability must be between zero and one")
            record = dict(values)
            record.update(_row=row_number, _columns=columns, _key=label)
            rows.append(record)
            # Runtime consumers distinguish an explicit null field from an absent key.
            output.append(values)
        self.rows[(stem, sheet)] = rows
        self.runtime[stem][sheet] = output
        self.provenance[(stem, sheet)] = provenance

    def table(self, book: str, sheet: str) -> list[dict]:
        return self.rows.get((book, sheet), [])

    def ids(self, book: str, sheet: str, field: str) -> set[str]:
        return {str(row[field]) for row in self.table(book, sheet) if row.get(field) is not None}

    def reference(self, book: str, sheet: str, row: dict, field: str, target: set[str]) -> None:
        value = row.get(field)
        if value is not None and str(value) not in target:
            self.row_error(book, sheet, row, field, "MissingReference", f"{value!r} is not a known target")

    def validate_relations(self) -> None:
        self._validate_mainline()
        self._validate_growth()
        self._validate_hyperspace()

    def _validate_mainline(self) -> None:
        equipment = self.ids("equipment", "equipment", "name")
        ships = self.ids("ship", "ship", "name")
        crew = self.ids("crew", "crew", "id")
        planets = self.ids("planet", "planet", "id")
        resources = self.ids("res", "res", "id")
        groups = self.ids("monGroup", "monGroup", "id")
        enemies = self.ids("mon", "mon", "id")
        unlocks = self.ids("unlock", "unlock", "name")
        legacy = self.ids("hightech", "hightech", "name")
        for row in self.table("unlock", "unlock"):
            target = {"equipment": equipment, "ship": ships, "crew": crew,
                      "planet": planets, "hightech": legacy}.get(row.get("type"))
            if target is not None:
                self.reference("unlock", "unlock", row, "target", target)
        for row in self.table("crew", "crew"):
            self.reference("crew", "crew", row, "unlockId", unlocks)
        for row in self.table("planet", "planet"):
            self.reference("planet", "planet", row, "unlockId", unlocks)
        for row in self.table("equipment", "equipment"):
            for field in ("res_1", "res_2"):
                self.reference("equipment", "equipment", row, field, resources)
            if row.get("cd") is not None and row["cd"] < 0:
                self.row_error("equipment", "equipment", row, "cd", "CooldownRange", "cooldown cannot be negative")
        for row in self.table("config", "config"):
            if row.get("name") == "startEquip":
                for name in str(row.get("para_1", "")).split(","):
                    if name.strip() not in equipment:
                        self.row_error("config", "config", row, "para_1", "MissingReference", f"equipment {name}")
        for row in self.table("mon", "mon"):
            for name in re.findall(r"([A-Za-z][A-Za-z0-9_-]*)\|", str(row.get("equipment", ""))):
                if name not in equipment:
                    self.row_error("mon", "mon", row, "equipment", "MissingReference", f"equipment {name}")
            drops = row.get("res")
            if drops is not None:
                expression = str(drops)
                blocks = expression[1:-1].split(";") if expression.startswith("{") and expression.endswith("}") else []
                if not blocks:
                    self.row_error("mon", "mon", row, "res", "InvalidDropExpression", "expected {resource,amount,chance} blocks")
                for block in blocks:
                    fields = block.split(",")
                    if len(fields) != 3:
                        self.row_error("mon", "mon", row, "res", "InvalidDropExpression", "expected resource,amount,chance")
                        continue
                    resource, amount, chance = fields
                    if resource not in resources:
                        self.row_error("mon", "mon", row, "res", "MissingReference", f"resource {resource}")
                    try:
                        numeric_amount = Decimal(amount)
                        numeric_chance = Decimal(chance)
                        if not numeric_amount.is_finite() or numeric_amount < 0:
                            raise ValueError("amount")
                        if not numeric_chance.is_finite() or not 0 <= numeric_chance <= 1:
                            raise ValueError("chance")
                    except (ValueError, InvalidOperation) as exc:
                        self.row_error("mon", "mon", row, "res", "EmbeddedDropRange", f"invalid {exc}")
        for row in self.table("monGroup", "monGroup"):
            members = str(row.get("mon", "")).strip("{} ").split(",")
            if len(members) not in (10, 15):
                self.row_error("monGroup", "monGroup", row, "mon", "InvalidFormation", "expected ten or fifteen slots")
            for member in members:
                if member.strip() != "null" and member.strip() not in enemies:
                    self.row_error("monGroup", "monGroup", row, "mon", "MissingReference", f"enemy {member}")
        for row in self.table("level", "level"):
            if not valid_type(row.get("length"), "number") or row["length"] <= 0:
                self.row_error("level", "level", row, "length", "StageLengthRange", "length must be positive")
            for field in ("atkRatio", "lifeRatio", "resRatio", "jewelRatio", "entryAtkRatio", "entryLifeRatio"):
                value = row.get(field)
                if value is not None and (not valid_type(value, "number") or value <= 0):
                    self.row_error("level", "level", row, field, "StageRatioRange", "stage ratio must be positive")
            if row.get("planetExpRatio") is not None and (not valid_type(row["planetExpRatio"], "number") or row["planetExpRatio"] < 0):
                self.row_error("level", "level", row, "planetExpRatio", "StageRatioRange", "planet experience ratio cannot be negative")
            try:
                reward_groups = json.loads(str(row.get("rewardReferenceGroups") or "[]"))
                if not isinstance(reward_groups, list): raise ValueError("not an array")
            except (ValueError, TypeError):
                self.row_error("level", "level", row, "rewardReferenceGroups", "InvalidRewardGroups", "expected JSON array")
                reward_groups = []
            for group in reward_groups:
                if str(group) not in groups:
                    self.row_error("level", "level", row, "rewardReferenceGroups", "MissingReference", f"group {group}")
            text = str(row.get("monGroup", ""))
            matches = re.findall(r"(\d+)\|([0-9.]+)", text)
            if not matches or text.count("|") != len(matches):
                self.row_error("level", "level", row, "monGroup", "InvalidStagePoints", "invalid point encoding")
            previous = 0.0
            for group, position in matches:
                if group not in groups:
                    self.row_error("level", "level", row, "monGroup", "MissingReference", f"group {group}")
                pos = float(position)
                if not previous < pos <= 1:
                    self.row_error("level", "level", row, "monGroup", "StagePointOrder", f"position {position}")
                previous = pos
        if self.manifest["fullGame"] and self.table("level", "level"):
            stages = sorted(int(row["id"]) for row in self.table("level", "level") if isinstance(row.get("id"), int))
            if stages != list(range(1, 221)):
                self.error("level.xlsx", "level", 0, "A", "", "StageCoverage", "expected exactly stages 1..220")

    def _validate_growth(self) -> None:
        planets = self.ids("planet", "planet", "id")
        buildings = self.ids("planet_build", "planet_build", "id")
        galaxy_keys = self.ids("galaxy", "galaxy", "key")
        research = self.ids("ai_tech_tree", "ai_tech_tree", "id")
        for row in self.table("planet_buff", "planet_buff"):
            self.reference("planet_buff", "planet_buff", row, "planet_id", planets)
        edges: dict[str, str] = {}
        for row in self.table("planet_build", "planet_build"):
            self.reference("planet_build", "planet_build", row, "min_planet", planets)
            self.reference("planet_build", "planet_build", row, "previous_id", buildings)
            if row.get("previous_id"):
                edges[str(row["id"])] = str(row["previous_id"])
        self._check_cycles(edges, "planet_build", "planet_build", "previous_id")
        for row in self.table("galaxy_build", "galaxy_build"):
            self.reference("galaxy_build", "galaxy_build", row, "galaxy_key", galaxy_keys)
        galaxy_edges: dict[str, str] = {}
        for row in self.table("galaxy", "galaxy"):
            self.reference("galaxy", "galaxy", row, "next_galaxy", galaxy_keys)
            if row.get("next_galaxy"):
                galaxy_edges[str(row["key"])] = str(row["next_galaxy"])
        self._check_cycles(galaxy_edges, "galaxy", "galaxy", "next_galaxy", id_field="key")
        tech_edges: dict[str, set[str]] = defaultdict(set)
        for row in self.table("ai_tech_tree", "ai_tech_tree"):
            for term in str(row.get("requirements") or "").split(";"):
                if not term:
                    continue
                tech = term.split(":", 1)[0]
                if tech not in research:
                    self.row_error("ai_tech_tree", "ai_tech_tree", row, "requirements", "MissingReference", f"tech {tech}")
                tech_edges[str(row["id"])].add(tech)
        self._check_cycles(tech_edges, "ai_tech_tree", "ai_tech_tree", "requirements")

    def _check_cycles(self, edges: dict, book: str, sheet: str, field: str, id_field: str = "id") -> None:
        active: set[str] = set()
        done: set[str] = set()

        def visit(node: str) -> None:
            if node in active:
                row = next((r for r in self.table(book, sheet) if str(r.get(id_field)) == node), None)
                if row:
                    self.row_error(book, sheet, row, field, "DependencyCycle", f"cycle at {node}")
                return
            if node in done:
                return
            active.add(node)
            children = edges.get(node, ())
            if isinstance(children, str): children = (children,)
            for child in children: visit(child)
            active.remove(node)
            done.add(node)

        for node in edges: visit(node)

    def _validate_hyperspace(self) -> None:
        parameter_rows = self.table("hyperspace_config", "基础参数")
        parameters = {str(row.get("key")): row.get("value") for row in parameter_rows}
        hyperspace_book = next(book for book in self.manifest["workbooks"] if book["file"] == "hyperspace_config.xlsx")
        parameter_schema = next(sheet for sheet in hyperspace_book["sheets"] if sheet["name"] == "基础参数")
        expected_keys = set(parameter_schema["allowedKeys"])
        missing_keys = expected_keys - set(parameters)
        if missing_keys:
            self.error("hyperspace_config.xlsx", "基础参数", 0, "A", ",".join(sorted(missing_keys)),
                       "MissingHyperspaceParameter", "registered parameter row is missing")
        contracted_keys = HYPERSPACE_INT_KEYS | HYPERSPACE_POSITIVE_KEYS | HYPERSPACE_NONNEGATIVE_KEYS | HYPERSPACE_PROBABILITY_KEYS | {"amplification_rate"}
        if contracted_keys != expected_keys:
            self.error("hyperspace_config.xlsx", "基础参数", 0, "", "", "ParameterContractMismatch",
                       f"missing={sorted(expected_keys-contracted_keys)} extra={sorted(contracted_keys-expected_keys)}")
        for row in parameter_rows:
            key, value = str(row.get("key")), row.get("value")
            if key not in expected_keys:
                continue  # UnknownRuleKey is reported by the generic sheet reader.
            expected_type = "int" if key in HYPERSPACE_INT_KEYS else "float"
            if row.get("data_type") != expected_type or not valid_type(value, "int" if expected_type == "int" else "number"):
                self.row_error("hyperspace_config", "基础参数", row, "value", "HyperspaceParameterType", f"{key} must use {expected_type}")
                continue
            if abs(value) >= HYPERSPACE_CONTRACT["magnitudeExclusive"]:
                self.row_error("hyperspace_config", "基础参数", row, "value", "HyperspaceParameterMagnitude", f"{key} exceeds the runtime number boundary")
            if key in HYPERSPACE_POSITIVE_KEYS and value <= 0:
                self.row_error("hyperspace_config", "基础参数", row, "value", "HyperspacePositiveParameter", f"{key} must be positive")
            elif key in HYPERSPACE_NONNEGATIVE_KEYS and value < 0:
                self.row_error("hyperspace_config", "基础参数", row, "value", "HyperspaceNonnegativeParameter", f"{key} cannot be negative")
            elif key in HYPERSPACE_PROBABILITY_KEYS and not 0 <= value <= 1:
                self.row_error("hyperspace_config", "基础参数", row, "value", "HyperspaceProbabilityRange", f"{key} must be within zero and one")
            elif key == "amplification_rate" and value <= -1:
                self.row_error("hyperspace_config", "基础参数", row, "value", "HyperspaceAmplificationRange", "1 + amplification_rate must be positive")
            if key in HYPERSPACE_CONTRACT["hardCaps"] and (value > HYPERSPACE_CONTRACT["hardCaps"][key] or key == "maximum_ultimate" and value != 1):
                self.row_error("hyperspace_config", "基础参数", row, "value", "HyperspaceHardCap", f"{key} exceeds its gameplay cap")
            if key in HYPERSPACE_CONTRACT["fixedValues"] and value != HYPERSPACE_CONTRACT["fixedValues"][key]:
                self.row_error("hyperspace_config", "基础参数", row, "value", "FixedRuleChanged", f"{key} is fixed by the gameplay contract")
        for child, parent in HYPERSPACE_CONTRACT["capacityRelations"]:
            if valid_type(parameters.get(child), "number") and valid_type(parameters.get(parent), "number") and parameters[child] > parameters[parent]:
                row = next(row for row in parameter_rows if row.get("key") == child)
                self.row_error("hyperspace_config", "基础参数", row, "value", "HyperspaceCapacityRelation", f"{child} exceeds {parent}")
        stage_ids = self.ids("level", "level", "id")
        for key in HYPERSPACE_CONTRACT["stageReferences"]:
            if valid_type(parameters.get(key), "int") and str(parameters[key]) not in stage_ids:
                row = next(row for row in parameter_rows if row.get("key") == key)
                self.row_error("hyperspace_config", "基础参数", row, "value", "MissingReference", f"{key} is not an authored stage")
        routes = self.ids("hyperspace_config", "路线", "route_id")
        affixes = self.ids("hyperspace_config", "词缀定义", "affix_id")
        legends = self.ids("hyperspace_config", "传说定义", "effect_id")
        enemies = self.ids("hyperspace_enemies", "敌人", "enemy_id")
        enemy_sizes = {str(row["enemy_id"]): row["size"] for row in self.table("hyperspace_enemies", "敌人")}
        groups = self.ids("hyperspace_enemies", "编队", "group_id")
        rewards = self.ids("hyperspace_enemies", "奖励参考", "reference_id")
        resource_ids = self.ids("res", "res", "id")
        stage_ids = self.ids("level", "level", "id")
        main_groups = self.ids("monGroup", "monGroup", "id")
        weapons = self.ids("equipment", "equipment", "name")
        for row in self.table("hyperspace_config", "路线"):
            self.reference("hyperspace_config", "路线", row, "weapon", weapons)
        for row in self.table("hyperspace_config", "词缀区间"):
            self.reference("hyperspace_config", "词缀区间", row, "affix_id", affixes)
            if valid_type(row.get("minimum"), "number") and valid_type(row.get("maximum"), "number") and row["minimum"] > row["maximum"]:
                self.row_error("hyperspace_config", "词缀区间", row, "minimum", "RangeReversed", "minimum exceeds maximum")
            if valid_type(row.get("minimum"), "number") and valid_type(row.get("maximum"), "number") and valid_type(parameters.get("value_precision"), "number") and parameters["value_precision"] > 0 and row["minimum"] <= row["maximum"] and not quantized_range_has_value(row["minimum"], row["maximum"], parameters["value_precision"]):
                self.row_error("hyperspace_config", "词缀区间", row, "minimum", "QuantizedRangeEmpty", "no value lies on the configured precision lattice")
        for sheet in ("传说随机参数", "传说常量"):
            for row in self.table("hyperspace_config", sheet):
                self.reference("hyperspace_config", sheet, row, "effect_id", legends)
                if sheet == "传说随机参数" and valid_type(row.get("minimum"), "number") and valid_type(row.get("maximum"), "number") and row["minimum"] > row["maximum"]:
                    self.row_error("hyperspace_config", sheet, row, "minimum", "RangeReversed", "minimum exceeds maximum")
                if sheet == "传说随机参数" and valid_type(row.get("minimum"), "number") and valid_type(row.get("maximum"), "number") and valid_type(parameters.get("value_precision"), "number") and parameters["value_precision"] > 0 and row["minimum"] <= row["maximum"] and not quantized_range_has_value(row["minimum"], row["maximum"], parameters["value_precision"]):
                    self.row_error("hyperspace_config", sheet, row, "minimum", "QuantizedRangeEmpty", "no value lies on the configured precision lattice")
        for row in self.table("hyperspace_enemies", "武器安装"):
            self.reference("hyperspace_enemies", "武器安装", row, "enemy_id", enemies)
            self.reference("hyperspace_enemies", "武器安装", row, "weapon_id", weapons)
        positions: dict[str, list[dict]] = defaultdict(list)
        for row in self.table("hyperspace_enemies", "编队槽位"):
            self.reference("hyperspace_enemies", "编队槽位", row, "group_id", groups)
            self.reference("hyperspace_enemies", "编队槽位", row, "enemy_id", enemies)
            if not isinstance(row.get("slot_index"), int) or not 0 <= row["slot_index"] < 15:
                self.row_error("hyperspace_enemies", "编队槽位", row, "slot_index", "InvalidSlotIndex", "slot must be 0..14")
            radius = 18 * enemy_sizes.get(str(row.get("enemy_id")), 1)
            if not all(valid_type(row.get(axis), "number") for axis in ("x", "y")) or not radius <= row["x"] <= 572 - radius or not radius <= row["y"] <= 960 - radius:
                self.row_error("hyperspace_enemies", "编队槽位", row, "x", "FormationBounds", "authored size envelope must fit 572x960 battle field")
            for other in positions[str(row.get("group_id"))]:
                other_radius = 18 * enemy_sizes.get(str(other.get("enemy_id")), 1)
                if math.hypot(row["x"] - other["x"], row["y"] - other["y"]) < radius + other_radius:
                    self.row_error("hyperspace_enemies", "编队槽位", row, "x", "FormationOverlap", "authored size envelopes overlap")
            positions[str(row.get("group_id"))].append(row)
        slots_by_group: dict[str, set[str]] = defaultdict(set)
        for row in self.table("hyperspace_enemies", "编队槽位"):
            slots_by_group[str(row.get("group_id"))].add(str(row.get("slot_index")))
        for row in self.table("hyperspace_enemies", "编队"):
            self.reference("hyperspace_enemies", "编队", row, "reward_reference_id", rewards)
            if str(row.get("group_id")) not in slots_by_group:
                self.row_error("hyperspace_enemies", "编队", row, "group_id", "FormationMissing", "group has no slots")
        for row in self.table("hyperspace_enemies", "奖励参考"):
            self.reference("hyperspace_enemies", "奖励参考", row, "late_reference_level_id", stage_ids)
            self.reference("hyperspace_enemies", "奖励参考", row, "late_reference_group_id", main_groups)
        coverage: Counter = Counter()
        ordered_groups: dict[tuple[str, str], list[int]] = defaultdict(list)
        assigned_groups: set[str] = set()
        for row in self.table("hyperspace_enemies", "路线编排"):
            self.reference("hyperspace_enemies", "路线编排", row, "route_id", routes)
            self.reference("hyperspace_enemies", "路线编排", row, "group_id", groups)
            coverage[(str(row.get("route_id")), str(row.get("tier")))] += 1
            ordered_groups[(str(row.get("route_id")), str(row.get("tier")))].append(row.get("order"))
            group = str(row.get("group_id"))
            if group in assigned_groups:
                self.row_error("hyperspace_enemies", "路线编排", row, "group_id", "RepeatedRouteGroup", "group appears more than once")
            assigned_groups.add(group)
            matching = next((g for g in self.table("hyperspace_enemies", "编队") if str(g.get("group_id")) == group), None)
            if matching and matching.get("tier") != row.get("tier"):
                self.row_error("hyperspace_enemies", "路线编排", row, "tier", "RouteTierMismatch", "group tier differs from route tier")
        for (route, tier), orders in ordered_groups.items():
            if sorted(orders) != list(range(len(orders))):
                self.error("hyperspace_enemies.xlsx", "路线编排", 0, "C", route, "RouteOrder", f"{tier}: order must start at zero and be contiguous")
        if self.manifest["fullGame"] and routes:
            for route in routes:
                for tier, count in (("normal", 4), ("elite", 4), ("boss", 1), ("ultimate", 1)):
                    if coverage[(route, tier)] != count:
                        self.error("hyperspace_enemies.xlsx", "路线编排", 0, "", route,
                                   "RouteCoverage", f"{tier}: expected {count}, got {coverage[(route, tier)]}")
        for row in self.table("hyperspace_enemies", "成员分配"):
            self.reference("hyperspace_enemies", "成员分配", row, "group_id", groups)
            if str(row.get("slot_index")) not in slots_by_group[str(row.get("group_id"))]:
                self.row_error("hyperspace_enemies", "成员分配", row, "slot_index", "MissingReference", "unknown formation slot")
        budgets: dict[str, set[int]] = defaultdict(set)
        for row in self.table("hyperspace_enemies", "早段回退预算"):
            self.reference("hyperspace_enemies", "早段回退预算", row, "reference_id", rewards)
            self.reference("hyperspace_enemies", "早段回退预算", row, "resource_id", resource_ids)
            ordinal = row.get("reference_member_ordinal")
            if isinstance(ordinal, int):
                budgets[str(row.get("reference_id"))].add(ordinal)
            if valid_type(row.get("amount"), "number") and row["amount"] < 0:
                self.row_error("hyperspace_enemies", "早段回退预算", row, "amount", "NegativeBudget", "zero is valid, negative is not")
        for row in self.table("hyperspace_enemies", "奖励参考"):
            reference = str(row.get("reference_id"))
            if reference in budgets and budgets[reference] != set(range(row["reference_member_count"])):
                self.row_error("hyperspace_enemies", "奖励参考", row, "reference_member_count", "BudgetCoverage", "published budget must cover every member ordinal")

    def source_hash(self) -> str:
        digest = hashlib.sha256()
        for entry in sorted(self.manifest["workbooks"], key=lambda item: item["file"]):
            path = self.source / entry["file"]
            digest.update(entry["file"].encode())
            digest.update(hashlib.sha256(path.read_bytes()).digest())
        return digest.hexdigest()

    def semantic_hash(self) -> str:
        normalized = {}
        for book in self.manifest["workbooks"]:
            if book["category"] != "active":
                continue
            stem = Path(book["file"]).stem
            normalized[stem] = {}
            for sheet in book["sheets"]:
                rows = self.runtime[stem][sheet["name"]]
                if "order" in sheet["columns"]:
                    rows = sorted(rows, key=lambda row: tuple(str(row.get(k, "")) for k in sheet["primary"]))
                normalized[stem][sheet["name"]] = rows
        encoded = json.dumps(normalized, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
        return hashlib.sha256(encoded.encode()).hexdigest()


def compile_config(manifest: Path, output: Path, report: Path | None = None,
                   source: Path | None = None) -> dict:
    try:
        compiler = Compiler(manifest)
        if source is not None:
            compiler.source = source.resolve()
        compiler.read()
        if not compiler.errors:
            compiler.validate_relations()
    except (ValueError, KeyError, OSError, TypeError) as exc:
        detail = str(exc)
        code = detail if detail in ("InvalidManifest", "TestManifestUsesProductionInput",
                                    "ProductionManifestUsesTestInput") else "CompilerFailure"
        result = {"ok": False, "errors": [{"workbook": "manifest", "sheet": "", "row": 0,
                  "column": "", "key": "", "code": code, "detail": detail}],
                  "sourceHash": "", "semanticHash": ""}
        if report:
            report.parent.mkdir(parents=True, exist_ok=True)
            report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return result
    result = {"ok": not compiler.errors, "errors": compiler.errors,
              "sourceHash": compiler.source_hash() if not compiler.errors else "",
              "semanticHash": compiler.semantic_hash() if not compiler.errors else ""}
    if compiler.errors:
        if report:
            report.parent.mkdir(parents=True, exist_ok=True)
            report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return result
    payload = json.dumps(compiler.runtime, ensure_ascii=False, indent=2, default=str) + "\n"
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="\n", dir=output.parent,
                                         prefix=output.name + ".", suffix=".tmp", delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        if temporary.read_text(encoding="utf-8") != payload:
            raise OSError("configuration readback mismatch")
        os.replace(temporary, output)
    finally:
        if temporary and temporary.exists():
            temporary.unlink()
    result["outputHash"] = hashlib.sha256(payload.encode()).hexdigest()
    if report:
        report.parent.mkdir(parents=True, exist_ok=True)
        report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--source", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    try:
        result = compile_config(args.manifest, args.output, args.report, args.source)
    except (ValueError, KeyError, OSError) as exc:
        print(json.dumps({"ok": False, "code": "CompilerFailure", "detail": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
