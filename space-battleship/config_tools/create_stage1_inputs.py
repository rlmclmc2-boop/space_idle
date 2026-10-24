"""One-time baseline schema capture and independent fixture authoring.

Not a runtime dependency: compile.py reads the reviewed, versioned schema.
This script never modifies production Excel. Deliberately requires --write.
"""
import argparse
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def main():
    p = argparse.ArgumentParser(); p.add_argument('--write', action='store_true', required=True); p.parse_args()
    tables = {}
    manifest = dict(schemaVersion=1, kind='production', inputRoot='../config_excel', workbooks=[], excludedSources={
        'battle_design.xlsx':'candidate metadata', 'level_monGroup_generated.xlsx':'generated copy',
        'published/':'release copies', '../../太空战舰.xlsx':'historical split input'})
    for path in sorted((ROOT/'config_excel').glob('*.xlsx')):
        if path.name in manifest['excludedSources']: continue
        book = openpyxl.load_workbook(path, data_only=False)
        entry = dict(path='../config_excel/'+path.name, classification='retained' if path.stem=='hightech' else 'active', sheets=[])
        for i, sheet in enumerate(book):
            table = path.stem if len(book.sheetnames)==1 else path.stem+'_'+str(i).zfill(2)
            rows = list(sheet.values)
            columns = list(rows[0])
            while columns and columns[-1] is None: columns.pop()
            fields = {}
            for col, name in enumerate(columns):
                if name is None: continue
                values = [r[col] if col < len(r) else None for r in rows[3:] if r[0] is not None]
                kinds = {'null' if v is None else 'bool' if isinstance(v,bool) else 'number' if isinstance(v,(float,int)) else 'string' for v in values}
                fields[name] = dict(types=sorted(kinds or {'null'}))
                if kinds <= {'number','null'} and any(v is not None for v in values):
                    if all(v is None or type(v) is int for v in values): fields[name]['integer']=True
            primary = [next(c for c in columns if c)]
            composites = {
                'hyperspace_config_06':['affix_id','tier'], 'hyperspace_config_09':['effect_id','parameter'],
                'hyperspace_config_10':['effect_id','constant_path'], 'hyperspace_config_11':['operation','resource'],
                'hyperspace_enemies_01':['enemy_id','order'], 'hyperspace_enemies_03':['group_id','slot_index'],
                'hyperspace_enemies_04':['route_id','tier','order'], 'hyperspace_enemies_06':['group_id','slot_index','order'],
                'hyperspace_enemies_07':['reference_id','reference_member_ordinal','drop_order']}
            primary = composites.get(table,primary)
            for field in primary:
                fields[field]['types'] = [t for t in fields[field]['types'] if t != 'null']
            spec=dict(columns=columns,fields=fields,primary=primary)
            if table=='hyperspace_enemies_06':
                fields['order']['types'].append('null')
                spec['nullablePrimary']=['order']
            if path.stem.startswith('hyperspace'): spec['literalOnly']=True
            if {'data_type','value'} <= fields.keys(): spec['typedValues']=True
            if table=='equipment':
                spec['references']=[dict(field=f,table='res',target='id') for f in ('res_1','res_2')]
                for f in ('cost_1','cost_2','costMulti_1','costMulti_2'): fields[f]['min']=0
            if table=='ai_tech_tree': spec['topologies']=[dict(field='requirements',id='id')]
            structured={
                'monGroup':dict(field='mon',table='mon',target='id',kind='slots'),
                'level':dict(field='monGroup',table='monGroup',target='id',kind='positions'),
                'mon':dict(field='equipment',table='equipment',target='name',kind='weapon_level')}
            if table in structured:spec['structuredReferences']=[structured[table]]
            tables[table]=spec
            entry['sheets'].append(dict(sheet=sheet.title, table=table))
        manifest['workbooks'].append(entry)
        book.close()
    # Test-only tables belong to the same versioned compiler schema.
    def spec(columns, numeric=(), primary=None):
        return dict(columns=columns,primary=primary or [columns[0]],fields={c:dict(types=['number'] if c in numeric else ['string']) for c in columns},literalOnly=True)
    tables['case']=spec(['case_id','config_fixture','seed','duration','comparison_mode'],['seed','duration'])
    tables['amount_case']=spec(['case_id','operation','left','right','parameter','count','expected','error','tolerance'],['parameter','count','tolerance'])
    for f in ('right','expected','error'):
        tables['amount_case']['fields'][f]['types'].append('null')
    tables['amount_case']['fields']['operation']['enum']=['add','sub_nonnegative','mul','div','pow','floor','ceil','sum_geometric','round_significant','compare','to_ratio','from_decimal']
    tables['amount_case']['references']=[dict(field='case_id',table='case',target='case_id')]
    (ROOT/'config_tools/schema.json').write_text(json.dumps(dict(version=1,tables=tables),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (ROOT/'config_tools/production.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    fixture=ROOT/'tests/excel'; fixture.mkdir(parents=True,exist_ok=True)
    book=openpyxl.Workbook();book.remove(book.active)
    def sheet(name, rows):
        s=book.create_sheet(name); cols=tables[name]['columns'];s.append(cols);s.append(['stage1 independent fixture']*len(cols));s.append(['literal']*len(cols))
        for row in rows: s.append(row)
    cases=[
        ['N01-add','add','1e500','1e499',0,0,'1.1e500',None],
        ['N01-mul','mul','2e500','3e400',0,0,'6e900',None],
        ['N01-div','div','6e900','3e400',0,0,'2e500',None],
        ['N01-sub','sub_nonnegative','1.000000000001e500','1e500',0,0,'1e488',None],
        ['N01-clamp','sub_nonnegative','1','2',0,0,'0',None],
        ['N01-pow','pow','2e250',None,2,0,'4e500',None],
        ['N01-ceil','ceil','6.01',None,0,0,'7',None],
        ['N01-floor','floor','6.99',None,0,0,'6',None],
        ['N01-int-floor','floor','113',None,0,0,'113',None],
        ['N01-int-ceil','ceil','201',None,0,0,'201',None],
        ['N01-series','sum_geometric','2',None,2,4,'30',None],
        ['N01-linear','sum_geometric','2e500',None,1,10000,'2e504',None],
        ['N01-div-zero','div','1','0',0,0,None,'AmountDivisionByZero'],
        ['N01-overflow','mul','1e9000000000000000','10',0,0,None,'AmountRangeExceeded']]
    for row in cases: row.append(1e-3 if row[0]=='N01-sub' else 1e-12)
    sheet('case',[[row[0],'core_rules',17,1/60,'amount_relative'] for row in cases]);sheet('amount_case',cases)
    # Genuine minimal game definitions, independent from production values.
    for name, data in {
        'res':[['id','name'],[71,'test_iron']],
        'enemy_weapon_base':[['id','dmg','cd','dmgtype','para1','para2','para3'],['laser',3,2,1,100,1,1]],
        'equipment':[['name','des','level','dmgtype','dmg','cd','res_1','cost_1','costMulti_1','para1','para2'],['armour','fixture',1,2,0,0,71,7,0.1,30,0.1],['laser_mon','fixture enemy',1,1,3,2,71,7,0.1,100,None]],
        'mon':[['id','des','equipment','dmgMultiple','health','armourType','res','size','shield','shieldType','shieldRecovery','shieldDelay'],[9101,'fixture','{laser_mon|1}',1,9,0,'{71,2,1}',2,0,0,0,0]],
        'monGroup':[['id','des','mon'],[9201,'fixture','{null,null,null,9101,null,null,null,null,null,null}']],
        'level':[['id','length','monGroup','atkRatio','lifeRatio','resRatio','jewelRatio','planetExpRatio'],[9301,100,'{9201|0.5}',1,1,1,1,0]]
    }.items():
        s=book.create_sheet(name);cols=tables[name]['columns'];s.append(cols);s.append(['independent']*len(cols));s.append(['literal']*len(cols))
        for values in data[1:]:
            mapped=dict(zip(data[0],values));s.append([mapped.get(c) for c in cols])
    book.save(fixture/'core_rules.xlsx')
    # Fixture authoring never changes production requiredness.
    (ROOT/'config_tools/schema.json').write_text(json.dumps(dict(version=1,tables=tables),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    test_manifest=dict(schemaVersion=1,kind='test',inputRoot='.',workbooks=[dict(path='core_rules.xlsx',classification='active',sheets=[dict(sheet=s,table=s) for s in book.sheetnames])])
    (fixture/'core_rules.json').write_text(json.dumps(test_manifest,indent=2)+'\n',encoding='utf-8')


if __name__=='__main__': main()
