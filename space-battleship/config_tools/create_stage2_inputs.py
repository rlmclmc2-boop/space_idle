"""Author stage-2 runtime Excel inputs; never called by the compiler/runtime."""
from pathlib import Path
import argparse,json
import openpyxl

ROOT=Path(__file__).resolve().parents[1]

def main():
    p=argparse.ArgumentParser();p.add_argument('--write',action='store_true',required=True);p.parse_args()
    schema=json.loads((ROOT/'config_tools/schema.json').read_text(encoding='utf-8'))
    production=json.loads((ROOT/'config_tools/production.json').read_text(encoding='utf-8'))
    rules={'clear_wait_seconds':3,'maximum_travel_seconds':3,'retreat_seconds':1.2,'departure_seconds':1.6,
           'back_distance':200,'auto_collect_delay':5,'fragment_collect_delay':2,'auto_collect_reduce':0.5,
           'resistance':0.5,'jewel_probability':0.2,'maximum_speed':10}
    # Logical formation constants exported from enemy_formation.positions and
    # existing native player point. No scene/editor dependency in compilation.
    geometry={'center_x':286,'player_y':520,'single_row_y':140,'back_row_y':110,'middle_row_y':270,'front_row_y':350,
              'columns':5,'maximum_spacing':160,'formation_width':456,'target_slot_center':4.5}
    def table(columns,numeric,primary):
        return dict(columns=columns,primary=[primary],literalOnly=True,fields={c:dict(types=['number'] if c in numeric else ['string']) for c in columns})
    for name,values in [('battle_rules',rules),('battle_geometry',geometry)]:
        schema['tables'][name]=table(['key','value'],['value'],'key')
        schema['tables'][name]['requiredKeys']=list(values)
        schema['tables'][name]['fields']['key']['enum']=list(values)
    columns=['id','initial_stage','ship_id','iron_id','uranium_id','fragment_id','chrono_id','initial_amount','initial_chrono','initial_slot_level','death_option']
    schema['tables']['initial_profile']=table(columns,['initial_stage','iron_id','uranium_id','initial_slot_level','death_option'],'id')
    schema['tables']['initial_profile']['fields']['initial_slot_level'].update(integer=True,min=1,max=1)
    schema['tables']['initial_profile']['references']=[dict(field='initial_stage',table='level',target='id'),dict(field='ship_id',table='ship',target='name'),dict(field='iron_id',table='res',target='id'),dict(field='uranium_id',table='res',target='id')]
    session_columns=['case_id','duration','expected_auto','expected_manual','expected_stage','resource']
    schema['tables']['session_case']=table(session_columns,['duration','expected_stage','resource'],'case_id')
    def write_sheet(book,name,cols,rows):
        if name in book:del book[name]
        s=book.create_sheet(name);s.append(cols);s.append(['runtime contract']*len(cols));s.append(['literal']*len(cols))
        for row in rows:s.append(row)
    for name,rows in [('battle_rules',list(rules.items())),('battle_geometry',list(geometry.items())),('initial_profile',[['default',1,'Frigate',1,2,'jewel','chrono','0','0',1,0]])]:
        path=ROOT/'config_excel'/(name+'.xlsx');book=openpyxl.Workbook();book.remove(book.active)
        write_sheet(book,name,schema['tables'][name]['columns'],rows);book.save(path)
        production['workbooks']=[e for e in production['workbooks'] if e['path']!='../config_excel/'+path.name]
        production['workbooks'].append(dict(path='../config_excel/'+path.name,classification='active',sheets=[dict(sheet=name,table=name)]))
    path=ROOT/'tests/excel/core_rules.xlsx';book=openpyxl.load_workbook(path)
    test_rules=dict(rules,clear_wait_seconds=0.1,maximum_travel_seconds=0.2,retreat_seconds=0.1,departure_seconds=0,auto_collect_delay=0.2,fragment_collect_delay=0.1)
    test_geometry=dict(geometry,center_x=0,player_y=100,single_row_y=0,back_row_y=0,middle_row_y=20,front_row_y=40)
    for name,rows in [('battle_rules',list(test_rules.items())),('battle_geometry',list(test_geometry.items())),('initial_profile',[['default',1,'fixture_ship',71,72,'jewel','chrono','0','0',1,0]]),('session_case',[['E01-session',20,'2','3',1,71]])]:
        write_sheet(book,name,schema['tables'][name]['columns'],rows)
    for name,rows in [('ship',[['fixture_ship','independent fixture',1,1,50,None,1,999]]),('config',[['startEquip','independent fixture','armour,laser']])]:
        write_sheet(book,name,schema['tables'][name]['columns'],rows)
    # Explicit independent definitions; no production content fills missing rows.
    if not any(book['res'].cell(r,1).value==72 for r in range(4,book['res'].max_row+1)):book['res'].append([72,'test_uranium'])
    level=book['level'];headers=[c.value for c in level[1]]
    level.cell(4,headers.index('id')+1,1)
    level.cell(4,headers.index('monGroup')+1,'{9201|0.5,9201|0.9}')
    book.save(path)
    manifest=dict(schemaVersion=1,kind='test',inputRoot='.',workbooks=[dict(path=path.name,classification='active',sheets=[dict(sheet=s,table=s) for s in book.sheetnames])])
    (ROOT/'tests/excel/core_rules.json').write_text(json.dumps(manifest,indent=2)+'\n')
    (ROOT/'config_tools/schema.json').write_text(json.dumps(schema,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (ROOT/'config_tools/production.json').write_text(json.dumps(production,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

if __name__=='__main__':main()
