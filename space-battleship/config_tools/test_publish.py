"""Black-box publication probes against copies of the real consumer workbooks."""

from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

from compile import DEFAULT_MANIFEST, ROOT, compile_config


class PublicationTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory(prefix="space-config-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.source = self.root / "input"
        shutil.copytree(ROOT / "config_excel", self.source, ignore=shutil.ignore_patterns("published"))
        self.output = self.root / "active.json"
        self.report = self.root / "validation.json"
        self.baseline = compile_config(DEFAULT_MANIFEST, self.output, self.report, self.source)
        self.assertTrue(self.baseline["ok"], self.baseline["errors"][:3])
        self.original = self.output.read_bytes()

    def change(self, file: str, sheet: str, coordinate: str, value: object) -> None:
        path = self.source / file
        workbook = openpyxl.load_workbook(path)
        workbook[sheet][coordinate] = value
        workbook.save(path)
        workbook.close()

    def reject(self, expected: str) -> None:
        result = compile_config(DEFAULT_MANIFEST, self.output, self.report, self.source)
        self.assertFalse(result["ok"])
        self.assertIn(expected, {error["code"] for error in result["errors"]})
        self.assertEqual(self.output.read_bytes(), self.original, "failed candidate replaced active config")
        self.assertEqual(json.loads(self.report.read_text(encoding="utf-8")), result)
        self.assertTrue(all({"workbook", "sheet", "row", "column", "key", "code", "detail"} <= error.keys()
                            for error in result["errors"]))

    def test_duplicate_key_rejected_without_replacing_active(self) -> None:
        workbook = openpyxl.load_workbook(self.source / "equipment.xlsx")
        sheet = workbook["equipment"]
        primary_column = next(i for i, cell in enumerate(sheet[1], 1) if cell.value == "name")
        sheet.cell(5, primary_column).value = sheet.cell(4, primary_column).value
        workbook.save(self.source / "equipment.xlsx")
        workbook.close()
        self.reject("DuplicatePrimaryKey")

    def test_missing_foreign_key_rejected(self) -> None:
        workbook = openpyxl.load_workbook(self.source / "hyperspace_enemies.xlsx")
        sheet = workbook["编队槽位"]
        column = next(i for i, cell in enumerate(sheet[1], 1) if cell.value == "enemy_id")
        sheet.cell(4, column).value = 99999999
        workbook.save(self.source / "hyperspace_enemies.xlsx")
        workbook.close()
        self.reject("MissingReference")

    def test_formula_without_cache_rejected(self) -> None:
        self.change("level.xlsx", "level", "A4", "=2+2")
        self.reject("FormulaCacheMissing")

    def test_out_of_bounds_formation_rejected(self) -> None:
        workbook = openpyxl.load_workbook(self.source / "hyperspace_enemies.xlsx")
        sheet = workbook["编队槽位"]
        column = next(i for i, cell in enumerate(sheet[1], 1) if cell.value == "x")
        sheet.cell(4, column).value = 900
        workbook.save(self.source / "hyperspace_enemies.xlsx")
        workbook.close()
        self.reject("FormationBounds")

    def test_route_gap_rejected(self) -> None:
        workbook = openpyxl.load_workbook(self.source / "hyperspace_enemies.xlsx")
        sheet = workbook["路线编排"]
        column = next(i for i, cell in enumerate(sheet[1], 1) if cell.value == "order")
        sheet.cell(4, column).value = 44
        workbook.save(self.source / "hyperspace_enemies.xlsx")
        workbook.close()
        self.reject("RouteOrder")

    def test_unknown_typed_rule_rejected(self) -> None:
        workbook = openpyxl.load_workbook(self.source / "hyperspace_config.xlsx")
        sheet = workbook["基础参数"]
        column = next(i for i, cell in enumerate(sheet[1], 1) if cell.value == "key")
        sheet.cell(4, column).value = "unregistered_rule"
        workbook.save(self.source / "hyperspace_config.xlsx")
        workbook.close()
        self.reject("UnknownRuleKey")


if __name__ == "__main__":
    unittest.main()
