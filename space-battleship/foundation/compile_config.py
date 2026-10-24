"""Snapshot explicit Excel inputs and atomically publish an immutable bundle.

No legacy importer, game scene, JSON data projection or editor is imported.
The committed schema is a declaration; live values never select their types.
"""
from __future__ import annotations
import argparse
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import tempfile
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SCHEMA = Path(__file__).with_name('config') / 'schema.json'


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False).encode('utf-8')


def digest(value):
    return hashlib.sha256(canonical(value)).hexdigest()


class InvalidConfig(Exception):
    def __init__(self, errors):
        self.errors = errors
        super().__init__(f'{len(errors)} configuration errors')


def literal_kind(value):
    if value is None: return 'null'
    if type(value) is bool: return 'bool'
    if type(value) in (float, int): return 'number'
    return 'string' if isinstance(value, str) else 'invalid'


def compile_bundle(manifest_path):
    manifest_path = Path(manifest_path).resolve()
    manifest_bytes = manifest_path.read_bytes()
    manifest = json.loads(manifest_bytes)
    schema_bytes = SCHEMA.read_bytes()
    schema = json.loads(schema_bytes)
    errors, tables, provenance, classes, sources = [], {}, {}, {}, []
    locations = {}

    def error(book, sheet, row, column, key, code, detail):
        errors.append(dict(workbook=book, sheet=sheet, row=row, column=column,
                           key=str(key), code=code, detail=str(detail)))

    if manifest.get('version') != 1 or manifest.get('domain') not in ('production', 'fixture'):
        raise InvalidConfig([dict(workbook=str(manifest_path), sheet='', row=0, column='', key='', code='ManifestDomain', detail='version/domain')])
    domain = manifest['domain']
    if domain == 'production' and manifest['workbooks'] != schema['productionSources']:
        raise InvalidConfig([dict(workbook=str(manifest_path), sheet='', row=0, column='', key='', code='ManifestCoverage', detail='production source declarations differ from schema')])
    input_root = (manifest_path.parent / manifest['root']).resolve()
    production_root = (ROOT / 'config_excel').resolve()
    if (domain == 'production' and input_root != production_root) or (domain == 'fixture' and input_root == production_root):
        raise InvalidConfig([dict(workbook=str(manifest_path), sheet='', row=0, column='', key='', code='ManifestIsolation', detail='input domain')])
    # Capture every input before parsing any workbook.
    captured = []
    seen_files = set()
    for entry in manifest['workbooks']:
        relative = Path(entry['file'])
        path = (input_root / relative).resolve()
        if relative.is_absolute() or '..' in relative.parts or not path.is_relative_to(input_root) or path in seen_files:
            error(str(relative), '', 0, '', '', 'ManifestPath', 'duplicate or escaping input')
            continue
        if domain == 'fixture' and path.is_relative_to(production_root):
            error(str(relative), '', 0, '', '', 'ManifestIsolation', 'production input in fixture')
            continue
        seen_files.add(path)
        try: raw = path.read_bytes()
        except OSError as exc:
            error(str(relative), '', 0, '', '', 'InputRead', exc)
            continue
        captured.append((entry, raw))
    if errors: raise InvalidConfig(errors)

    for entry, raw in captured:
        book_name, classification = entry['file'], entry['classification']
        if classification not in ('active', 'retired', 'candidate', 'generated', 'published'):
            error(book_name, '', 0, '', '', 'Classification', classification)
            continue
        formulas = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        cached = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        declared = [s['sheet'] for s in entry['sheets']]
        sources.append(dict(file=book_name, classification=classification, sha256=hashlib.sha256(raw).hexdigest(), sheets=declared))
        if set(declared) != set(formulas.sheetnames) or len(declared) != len(set(declared)):
            error(book_name, '', 0, '', '', 'SheetSet', 'all sheets must be explicit and unique')
            continue
        for declaration in entry['sheets']:
            sheet_name, name = declaration['sheet'], declaration['table']
            if sheet_name not in formulas: continue
            if classification in ('candidate', 'generated', 'published'): continue
            if name not in schema['tables'] or name in tables:
                error(book_name, sheet_name, 0, '', name, 'TableIdentity', 'unknown or duplicate table')
                continue
            spec = schema['tables'][name]
            fs, vs = formulas[sheet_name], cached[sheet_name]
            formula_rows, value_rows = fs.iter_rows(), vs.iter_rows()
            header = [c.value for c in next(formula_rows)]
            next(value_rows)
            if header != spec['columns']:
                error(book_name, sheet_name, 1, '', name, 'Columns', header)
                continue
            for _ in range(2): next(formula_rows, None); next(value_rows, None)
            records, origins, identities = [], {}, set()
            for row_number, (cells, values) in enumerate(zip(formula_rows, value_rows), 4):
                if all(c.value is None for c in cells): continue
                row = {}
                row_bad = False
                for column, field in enumerate(header):
                    cell, cache = cells[column], values[column]
                    if field is None:
                        if cell.value is not None: error(book_name, sheet_name, row_number, cell.column_letter, '', 'UnnamedColumn', cell.value)
                        continue
                    value = cache.value
                    contract = spec['fields'][field]
                    code = None
                    if cell.data_type == 'f' and (spec.get('literalOnly') or cache.value is None):
                        code = 'FormulaForbidden' if spec.get('literalOnly') else 'FormulaCacheMissing'
                    elif cell.data_type == 'e' or cache.data_type == 'e': code = 'CellError'
                    elif value is None and contract.get('required'): code = 'Required'
                    elif literal_kind(value) not in contract['types']: code = 'LiteralType'
                    elif value is not None and 'enum' in contract and value not in contract['enum']: code = 'UnknownRule'
                    elif literal_kind(value) == 'number':
                        if not math.isfinite(value): code = 'NonFinite'
                        elif contract.get('integer') and (value != int(value) or abs(value) > 9007199254740991): code = 'ExactInteger'
                        elif 'min' in contract and value < contract['min']: code = 'Range'
                        elif 'max' in contract and value > contract['max']: code = 'Range'
                    if contract.get('literal') == 'amount_text' and value is not None:
                        if not isinstance(value, str) or not re.fullmatch(r'\+?[0-9]+(?:\.[0-9]*)?(?:[eE][+-]?[0-9]+)?', value): code = 'AmountLiteral'
                    if contract.get('literal') == 'int64_text' and value is not None:
                        if not isinstance(value, str) or not re.fullmatch(r'0|[1-9][0-9]*', value) or int(value) > 9223372036854775807: code = 'IntegerEncoding'
                    if code:
                        row_bad = True
                        error(book_name, sheet_name, row_number, get_column_letter(column + 1), field, code, value)
                    row[field] = value
                if row_bad: continue
                identity = canonical([row.get(k) for k in spec['primary']]).decode()
                if identity in identities: error(book_name, sheet_name, row_number, 'A', identity, 'DuplicateId', name)
                identities.add(identity)
                origins[identity] = {field: dict(workbook=book_name, sheet=sheet_name, row=row_number, column=get_column_letter(j + 1)) for j, field in enumerate(header) if field is not None}
                locations[(name, identity)] = (book_name, sheet_name, row_number)
                records.append(row)
            records.sort(key=lambda r: canonical([r.get(k) for k in spec['primary']]))
            if domain == 'production' and spec.get('requiredKeys'):
                present = {r[spec['primary'][0]] for r in records}
                for missing in set(spec['requiredKeys']) - present:
                    error(book_name, sheet_name, 0, 'A', missing, 'MissingRule', 'required registered key')
            tables[name], provenance[name], classes[name] = records, origins, classification
        formulas.close(); cached.close()

    def row_error(table, row, field, code, detail):
        spec = schema['tables'][table]
        identity = canonical([row.get(k) for k in spec['primary']]).decode()
        origin = provenance[table][identity][field]
        error(origin['workbook'], origin['sheet'], origin['row'], origin['column'], identity, code, detail)

    for name, records in tables.items():
        spec = schema['tables'][name]
        for row in records:
            for relation in spec.get('references', []):
                value = row.get(relation['field'])
                if value is None: continue
                allowed = {r[relation['target']] for r in tables.get(relation['table'], [])}
                if value not in allowed: row_error(name, row, relation['field'], 'MissingReference', relation['table'] + ':' + str(value))
            type_field = 'data_type' if 'data_type' in row else 'type' if name == 'galaxy_config' else None
            if type_field and 'value' in row:
                expected = {'int':'number','float':'number','bool':'bool','string':'string','null':'null'}.get(row[type_field])
                value = row['value']
                if expected is None or literal_kind(value) != expected or (row[type_field] == 'int' and value != int(value)):
                    row_error(name, row, 'value', 'TypedLiteral', row[type_field])
        if 'orderScope' in spec:
            scopes = {}
            for row in records:
                scope = tuple(row[k] for k in spec['orderScope'])
                scopes.setdefault(scope, []).append(row)
            for rows in scopes.values():
                zero_marker = spec.get('nullOrderMeaning') == 'zero_budget_member' and len(rows) == 1 and rows[0]['order'] is None and rows[0].get('reference_member_ordinal') is None
                if not zero_marker and (any(r['order'] is None for r in rows) or sorted(r['order'] for r in rows) != list(range(len(rows)))):
                    row_error(name, rows[0], 'order', 'OrderSequence', 'orders must be unique, zero-based and continuous within scope')

    def brace_list(value):
        text = str(value).translate(str.maketrans('｛｝，','{},')).strip()
        if not text.startswith('{') or not text.endswith('}'): raise ValueError('braced list required')
        return text[1:-1].split(',')

    for table, field, target, target_key, syntax in [
        ('monGroup','mon','mon','id','slots'), ('mon','equipment','equipment','name','weapons'),
        ('level','monGroup','monGroup','id','points')]:
        allowed = {str(r[target_key]) for r in tables.get(target, [])}
        for row in tables.get(table, []):
            previous = -1
            try: tokens = brace_list(row[field])
            except ValueError:
                row_error(table, row, field, 'StructuredReference', row[field])
                continue
            for token in tokens:
                token = token.strip()
                if syntax == 'slots' and token == 'null': continue
                parts = token.split('|')
                if parts[0] not in allowed: row_error(table, row, field, 'MissingReference', token)
                if syntax in ('weapons','points'):
                    try:
                        value = float(parts[1])
                        if len(parts) != 2 or not math.isfinite(value): raise ValueError()
                        if syntax == 'points' and not 0 <= previous < value <= 1: # first point may be zero
                            if not (previous == -1 and 0 <= value <= 1): raise ValueError()
                        if syntax == 'weapons' and (value < 1 or int(value) != value): raise ValueError()
                        previous = value
                    except (ValueError, IndexError): row_error(table, row, field, 'StructuredReference', token)
    for row in tables.get('mon', []):
        if row.get('res') is None: continue
        text = row['res'].translate(str.maketrans('｛｝，','{},'))
        groups = re.findall(r'\{([^{}]+)\}', text) if '{' in text else [text]
        remainder = re.sub(r'\{[^{}]+\}', '', text).replace(',', '').replace(';', '').strip() if '{' in text else ''
        if not groups or remainder:
            row_error('mon', row, 'res', 'DropReference', text)
            continue
        for block in groups:
            for segment in block.split(';'):
                try:
                    resource, amount, chance = segment.split(',')
                    if int(resource) not in {r['id'] for r in tables.get('res', [])} or not re.fullmatch(r'\+?[0-9]+(?:\.[0-9]*)?(?:[eE][+-]?[0-9]+)?', amount.strip()) or not 0 <= float(chance) <= 1: raise ValueError()
                except (ValueError, OverflowError): row_error('mon', row, 'res', 'DropReference', segment)
    for row in tables.get('unlock', []):
        kind, target = row['type'], row['target']
        if kind in schema['unlockTargets']:
            table, field = schema['unlockTargets'][kind]
            if str(target) not in {str(r[field]) for r in tables.get(table, [])}:
                row_error('unlock', row, 'target', 'MissingReference', str(target))
        elif target not in schema['unlockRuleIds'].get(kind, []):
            row_error('unlock', row, 'target', 'UnknownRule', str(target))
        if row['level'] != 0 and row['level'] not in {r['id'] for r in tables.get('level', [])}:
            row_error('unlock', row, 'level', 'MissingReference', row['level'])
    # Research graph is both validated and exported once.
    graph = {}
    for row in tables.get('ai_tech_tree', []):
        links = []
        for requirement in re.split(r'[;,]', str(row.get('requirements') or '')):
            if not requirement: continue
            parts = requirement.split(':')
            identity = parts[0]; links.append(identity)
            if len(parts) != 2 or not re.fullmatch(r'[0-9]+(?:\*k\+[0-9]+)?', parts[1]):
                row_error('ai_tech_tree', row, 'requirements', 'RequirementSyntax', requirement)
            if identity not in {r['id'] for r in tables.get('ai_tech_tree', [])}: row_error('ai_tech_tree', row, 'requirements', 'MissingReference', identity)
        graph[row['id']] = links
    ordered, visiting, complete = [], set(), set()
    def visit(node):
        if node in visiting: raise ValueError(node)
        if node in complete: return
        visiting.add(node)
        for dependency in graph.get(node, []): visit(dependency)
        visiting.remove(node); complete.add(node); ordered.append(node)
    try:
        for node in graph: visit(node)
    except ValueError as exc:
        row = next(r for r in tables['ai_tech_tree'] if r['id'] == str(exc))
        row_error('ai_tech_tree', row, 'requirements', 'DependencyCycle', exc)
    for name in ('hyper_ranges', 'hyper_legend_ranges'):
        for row in tables.get(name, []):
            if row['minimum'] > row['maximum']: row_error(name, row, 'minimum', 'IntervalOrder', row['maximum'])
    # Do not compact holes or silently accept overlapping explicit slots.
    occupied = {}
    for row in tables.get('hyper_slots', []):
        identity = (row['group_id'], row['x'], row['y'])
        if identity in occupied: row_error('hyper_slots', row, 'x', 'CoincidentSlots', identity)
        occupied[identity] = True
    if domain == 'production' and set(tables) != set(schema['productionTables']):
        error(str(manifest_path), '', 0, '', '', 'ManifestCoverage', 'active/retired tables must be complete')
    if errors: raise InvalidConfig(errors)
    indices = {name: {canonical([r[k] for k in schema['tables'][name]['primary']]).decode(): i for i, r in enumerate(rows)} for name, rows in tables.items()}
    semantic = dict(schemaVersion=1, domain=domain, tables=tables, tableClasses=classes)
    return dict(**semantic, semanticHash=digest(semantic), sourceManifestHash=digest(sources), schemaHash=hashlib.sha256(schema_bytes).hexdigest(),
                indices=indices, ruleParameters={}, researchOrder=ordered, provenance=provenance, sources=sources)


def publish(bundle, output):
    output = Path(output); versions = output / 'versions'; versions.mkdir(parents=True, exist_ok=True)
    raw = canonical(bundle); checksum = hashlib.sha256(raw).hexdigest()
    destination = versions / (checksum + '.json')
    def atomic(path, content):
        fd, temporary = tempfile.mkstemp(prefix='.compile-', dir=path.parent)
        try:
            with os.fdopen(fd, 'wb') as stream: stream.write(content); stream.flush(); os.fsync(stream.fileno())
            if Path(temporary).read_bytes() != content: raise OSError('readback mismatch')
            os.replace(temporary, path)
        finally:
            if Path(temporary).exists(): Path(temporary).unlink()
    if not destination.exists(): atomic(destination, raw)
    if destination.read_bytes() != raw: raise OSError('immutable version mismatch')
    pointer = dict(schemaVersion=1, bundle='versions/' + destination.name, checksum=checksum, semanticHash=bundle['semanticHash'])
    atomic(output / 'active.json', canonical(pointer))
    return pointer


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--manifest', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    try:
        bundle = compile_bundle(args.manifest); pointer = publish(bundle, args.output)
        print(json.dumps(dict(ok=True, tables=len(bundle['tables']), rows=sum(map(len,bundle['tables'].values())), **pointer)))
    except InvalidConfig as exc:
        args.output.mkdir(parents=True, exist_ok=True)
        (args.output / 'errors.json').write_bytes(canonical(exc.errors))
        print(json.dumps(dict(ok=False, errors=exc.errors[:20]), ensure_ascii=True)); raise SystemExit(1)
    except (OSError, KeyError, TypeError, ValueError) as exc:
        error = dict(workbook=str(args.manifest), sheet='', row=0, column='', key='', code='CompilationIO' if isinstance(exc, OSError) else 'ManifestShape', detail=str(exc))
        print(json.dumps(dict(ok=False, errors=[error]), ensure_ascii=True)); raise SystemExit(1)


if __name__ == '__main__': main()
