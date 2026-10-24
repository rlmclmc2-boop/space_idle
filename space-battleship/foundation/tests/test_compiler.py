import copy
import importlib.util
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location('foundation_compiler', ROOT / 'foundation/compile_config.py')
compiler = importlib.util.module_from_spec(spec)
spec.loader.exec_module(compiler)


class ConfigContracts(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary.name)
        shutil.copy2(ROOT / 'foundation/tests/excel/foundation.xlsx', self.directory / 'foundation.xlsx')
        data = json.loads((ROOT / 'foundation/tests/fixture.json').read_text(encoding='utf-8'))
        data['root'] = '.'
        self.manifest = self.directory / 'fixture.json'
        self.manifest.write_text(json.dumps(data), encoding='utf-8')

    def tearDown(self):
        self.temporary.cleanup()

    def edit(self, function):
        path = self.directory / 'foundation.xlsx'
        book = openpyxl.load_workbook(path)
        function(book)
        book.save(path)
        book.close()

    def reject(self, code):
        output = self.directory / 'active'
        compiler.publish(compiler.compile_bundle(ROOT / 'foundation/tests/fixture.json'), output)
        before = (output / 'active.json').read_bytes()
        with self.assertRaises(compiler.InvalidConfig) as caught:
            compiler.compile_bundle(self.manifest)
        errors = caught.exception.errors
        self.assertIn(code, {e['code'] for e in errors})
        self.assertEqual(before, (output / 'active.json').read_bytes())
        self.assertTrue(all(set(e) == {'workbook','sheet','row','column','key','code','detail'} for e in errors))
        return errors

    def test_deterministic_snapshot(self):
        one = compiler.compile_bundle(self.manifest)
        self.assertEqual(one, compiler.compile_bundle(self.manifest))
        self.assertEqual(one['provenance']['res']['[71]']['id']['row'], 4)

    def test_row_reordering_is_not_a_semantic_change(self):
        before = compiler.compile_bundle(self.manifest)
        def swap(book):
            sheet = book['res']; one = [c.value for c in sheet[4]]; two = [c.value for c in sheet[5]]
            for j, (a, b) in enumerate(zip(one, two), 1): sheet.cell(4,j,b); sheet.cell(5,j,a)
        self.edit(swap)
        after = compiler.compile_bundle(self.manifest)
        self.assertEqual(before['semanticHash'], after['semanticHash'])
        self.assertNotEqual(before['sourceManifestHash'], after['sourceManifestHash'])

    def test_duplicate_id(self):
        self.edit(lambda b: b['res'].append([71, 'duplicate']))
        self.reject('DuplicateId')

    def test_text_number_is_rejected(self):
        self.edit(lambda b: setattr(b['res']['A4'], 'value', '71'))
        self.reject('LiteralType')

    def test_required_id(self):
        self.edit(lambda b: setattr(b['res']['A4'], 'value', None))
        self.reject('Required')

    def test_non_integral_id(self):
        self.edit(lambda b: setattr(b['res']['A4'], 'value', 71.5))
        self.reject('ExactInteger')

    def test_formula_without_cached_value(self):
        self.edit(lambda b: setattr(b['res']['B4'], 'value', '="iron"'))
        self.reject('FormulaCacheMissing')

    def test_literal_table_rejects_formula(self):
        self.edit(lambda b: setattr(b['amount_case']['C4'], 'value', '=10^500'))
        self.reject('FormulaForbidden')

    def test_cell_error(self):
        self.edit(lambda b: setattr(b['res']['B4'], 'value', '#VALUE!'))
        self.reject('CellError')

    def test_amount_encoding(self):
        self.edit(lambda b: setattr(b['amount_case']['C4'], 'value', '-1e500'))
        self.reject('AmountLiteral')

    def test_integer_encoding(self):
        self.edit(lambda b: setattr(b['amount_case']['F12'], 'value', '9223372036854775808'))
        self.reject('IntegerEncoding')

    def test_unknown_sheet(self):
        self.edit(lambda b: b.create_sheet('undeclared'))
        self.reject('SheetSet')

    def test_domain_isolation(self):
        data = json.loads(self.manifest.read_text()); data['root'] = str(ROOT / 'config_excel')
        self.manifest.write_text(json.dumps(data))
        self.reject('ManifestIsolation')

    def test_path_escape(self):
        data = json.loads(self.manifest.read_text()); data['workbooks'][0]['file'] = '../foundation.xlsx'
        self.manifest.write_text(json.dumps(data))
        self.reject('ManifestPath')

    def test_atomic_pointer_failure_keeps_previous(self):
        output = self.directory / 'publish'
        before = compiler.compile_bundle(self.manifest)
        compiler.publish(before, output)
        pointer = (output / 'active.json').read_bytes()
        after = copy.deepcopy(before); after['testPublication'] = True
        replace = compiler.os.replace
        def fail(source, destination):
            if Path(destination).name == 'active.json': raise OSError('injected pointer failure')
            return replace(source, destination)
        with patch.object(compiler.os, 'replace', side_effect=fail):
            with self.assertRaises(OSError): compiler.publish(after, output)
        self.assertEqual(pointer, (output / 'active.json').read_bytes())

    def test_missing_foreign_key(self):
        # A minimal independent source using the shared equipment declaration.
        schema = json.loads(compiler.SCHEMA.read_text(encoding='utf-8'))
        columns = schema['tables']['equipment']['columns']
        def append(book):
            s = book.create_sheet('equipment'); s.append(columns); s.append(['fixture']*len(columns)); s.append(['literal']*len(columns))
            row = dict(name='fixture_weapon', des='independent', level=1, dmgtype=1, dmg=3, dmgMulti=0.2, cd=1,
                       res_1=999, cost_1=1, costMulti_1=0.2, para1=1)
            s.append([row.get(k) for k in columns])
        self.edit(append)
        data = json.loads(self.manifest.read_text()); data['workbooks'][0]['sheets'].append(dict(sheet='equipment',table='equipment'))
        self.manifest.write_text(json.dumps(data))
        self.reject('MissingReference')

    def add_research(self, unknown=False):
        schema = json.loads(compiler.SCHEMA.read_text(encoding='utf-8'))
        columns = schema['tables']['ai_tech_tree']['columns']
        def append(book):
            s = book.create_sheet('ai_tech_tree'); s.append(columns); s.append(['fixture']*len(columns)); s.append(['literal']*len(columns))
            for identity, requirements in [('a','b:1;c:1'), ('b',None), ('c','a:1')]:
                row = dict(id=identity, name='fixture', description='independent graph', tier='basic', base_seconds=1,
                           growth=1, effect='unknown_effect' if unknown else 'ai_count', factor=1,
                           unlock=None, requirements=requirements, legacy_unlock=None)
                s.append([row.get(k) for k in columns])
        self.edit(append)
        data = json.loads(self.manifest.read_text()); data['workbooks'][0]['sheets'].append(dict(sheet='ai_tech_tree',table='ai_tech_tree'))
        self.manifest.write_text(json.dumps(data))

    def test_cycle_on_second_semicolon_dependency(self):
        self.add_research()
        self.reject('DependencyCycle')

    def test_unknown_effect(self):
        self.add_research(unknown=True)
        self.reject('UnknownRule')

    def test_numeric_range(self):
        self.edit(lambda b: setattr(b['amount_case']['I4'], 'value', -1))
        self.reject('Range')

    def test_incomplete_production_manifest(self):
        data = json.loads((ROOT / 'foundation/config/production.json').read_text(encoding='utf-8'))
        data['root'] = str(ROOT / 'config_excel'); data['workbooks'].pop()
        self.manifest.write_text(json.dumps(data))
        self.reject('ManifestCoverage')

    def test_production_content_is_preserved(self):
        bundle = compiler.compile_bundle(ROOT / 'foundation/config/production.json')
        self.assertEqual(len(bundle['tables']['level']), 220)
        self.assertEqual({r['id'] for r in bundle['tables']['level']}, set(range(1,221)))
        self.assertEqual(len(bundle['tables']['mon']), 10461)
        self.assertEqual(len(bundle['tables']['monGroup']), 2330)
        self.assertEqual(sum(r['order'] is None for r in bundle['tables']['hyper_members']), 42)
        self.assertEqual(bundle['tableClasses']['hightech'], 'retired')
        self.assertNotIn('battle_design', bundle['tables'])
        self.assertNotIn('level_monGroup_generated', bundle['tables'])
        # Compare the entire set of literal field values, not only row counts.
        for source in json.loads((ROOT / 'foundation/config/production.json').read_text(encoding='utf-8'))['workbooks']:
            if source['classification'] not in ('active','retired'): continue
            book = openpyxl.load_workbook(ROOT / 'config_excel' / source['file'], read_only=True, data_only=True)
            for declaration in source['sheets']:
                values = list(book[declaration['sheet']].values); headers = values[0]
                rows = [{k:v for k,v in zip(headers,row) if k is not None} for row in values[3:] if row and row[0] is not None]
                self.assertEqual(sorted(map(compiler.canonical,rows)), sorted(map(compiler.canonical,bundle['tables'][declaration['table']])))
            book.close()


if __name__ == '__main__': unittest.main()
